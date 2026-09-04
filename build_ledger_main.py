#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
上海康肯销售订单管理系统台账生成器（源表单行样式）。

结构（每 JOB 一页）：
  第1行  受注管理情報
  第2~3行 合同头：担当者/取引先/PO NO./JOB NO/注文内容/支払条件/発送方式/出荷先/納入先
  第4行  分组辅助行：出荷状況 | 回収状況 | 其他情報
  第5行  列头；第6行起 一行一台设备

出荷状況：发货批次、设备型号、PO No、制造番号、机番、未税单价、出荷日、出荷台数、送货单回收
回収状況：按付款款类动态生成（有发票=开票日/状况/金额+回款日/金额 5列；无发票=回款日/金额 2列）；
          存在“全额”发票时增加 全额·开票日/状况/金额 3 列。
其他情報：是否无偿、验收状态、质保开始日、质保结束日、质保期、备注

合并规则：发货批次/出荷日/出荷台数按批整段合并；开票按发票覆盖合并；
          回款按回款覆盖整段合并（中间夹免费/未覆盖行也连成一段，与源表一致）。
"""
import json
import os
import re
from collections import Counter
from datetime import datetime, timedelta

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

EXPORT_DIR = '/tmp/export'
PROJECT_DIR = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(PROJECT_DIR, 'data', 'BKTC_Ledger.xlsx')
THIN = Side(style='thin', color='B7C3D0')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
F_TITLE = Font(bold=True, size=14, color='1F4E79')
F_HEAD = Font(bold=True, size=10, color='FFFFFF')
F_BODY = Font(size=10)
F_BOLD = Font(bold=True, size=10)
FILL_HEAD = PatternFill('solid', fgColor='1F4E79')
FILL_INFO = PatternFill('solid', fgColor='DDEBF7')
FILL_TOTAL = PatternFill('solid', fgColor='F2F2F2')
# 预警灯底色/字色（条件格式，跨平台：不依赖 emoji 字形）
FILL_WARN_RED = PatternFill('solid', fgColor='FBE2E2')
FILL_WARN_YEL = PatternFill('solid', fgColor='FFF2CC')
FILL_WARN_GRN = PatternFill('solid', fgColor='E2EFDA')
F_WARN_RED = Font(bold=True, size=10, color='C00000')
F_WARN_GRN = Font(bold=True, size=10, color='548235')
AL_C = Alignment(horizontal='center', vertical='center', wrap_text=True)
AL_L = Alignment(horizontal='left', vertical='center', wrap_text=True)
AL_R = Alignment(horizontal='right', vertical='center')

FMT_MONEY = '#,##0.00'
FMT_PRICE = '#,##0.000'
FMT_DATE = 'yyyy/mm/dd'
CANON = {'预付款': 0, '发货款': 1, '到货款': 2, '验收款': 3, '质保款': 4, '全额': 5}

DEFAULT_RULES = {
    '临近天数': 30,
    '金额容差': 0.01,
    '未验收待确认': True,
    '发票异常待确认': True,
}


def load_table(paths):
    fields, rows, rids = None, [], []
    for p in paths:
        d = json.load(open(p))['data']
        if fields is None:
            fields = d['fields']
        rows += d['data']
        rids += d['record_id_list']
    out = []
    for rid, row in zip(rids, rows):
        rec = {'记录ID': rid}
        for f, v in zip(fields, row):
            rec[f] = v
        out.append(rec)
    return out


def s(v):
    if v is None:
        return ''
    if isinstance(v, list):
        return s(v[0]) if v else ''
    return str(v).strip()


def fmt_date(v):
    if v is None or v == '':
        return ''
    if isinstance(v, (int, float)):
        ts = v / 1000 if v > 10_000_000_000 else v
        return datetime.fromtimestamp(ts).strftime('%Y/%m/%d')
    st = s(v)
    for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M', '%Y-%m-%d'):
        try:
            return datetime.strptime(st[:19], fmt).strftime('%Y/%m/%d')
        except ValueError:
            continue
    return st[:10].replace('-', '/')


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def coverage_set(v):
    return {x.strip() for x in re.split(r'[;；\r\n]+', s(v)) if x.strip()}


def batches_set(v):
    return {x.strip() for x in re.split(r'[;；\r\n]+', s(v)) if x.strip()}


def cov_contains(cov_text, ser):
    """覆盖串是否包含该制造番号（支持 → 调拨写法，如 23BS004-047→23BS006-068）。"""
    for el in coverage_set(cov_text):
        if ser == el:
            return True
        parts = re.split(r'[→⇒⟶➡➝]', el)
        if ser in parts:
            return True
    return False


def cov_match(x, ser, batch):
    """发票/回款覆盖匹配：有製造番号→番号+批次；无製造番号→按覆盖批次回退。"""
    cov = s(x['覆盖製造番号'])
    bs = batches_set(x['覆盖批次'])
    if cov and ser:
        return cov_contains(cov, ser) and (not bs or batch in bs)
    return bool(bs) and batch in bs


def matching_records(cands, ser, batch):
    """返回设备覆盖到的全部记录，保留番号优先的兜底语义。

    只要有显式製造番号匹配，就不再同时计入纯批次记录，
    避免历史的批次兜底行与精确分配行重复计金额。
    """
    explicit = [x for x in cands
                if s(x['覆盖製造番号']) and cov_match(x, ser, batch)]
    if explicit:
        return explicit
    return [x for x in cands
            if not s(x['覆盖製造番号']) and cov_match(x, ser, batch)]


def first_match(cands, ser, batch):
    """兼容旧调用；新的金额计算应使用 matching_records。"""
    matches = matching_records(cands, ser, batch)
    return matches[0] if matches else None


def to_dt(v):
    st = s(v)
    for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M', '%Y-%m-%d'):
        try:
            return datetime.strptime(st[:19], fmt)
        except ValueError:
            continue
    return None


def payment_term_days(term, default=0):
    """Return a non-negative integer term in days."""
    value = num(term.get('账期天数')) if term else None
    return max(0, int(value)) if value is not None else default


def warranty_term_settings(kind, term):
    """Return (uses warranty expiry, days after expiry).

    A missing trigger is treated as the warranty-expiry rule for legacy rows;
    an explicit non-warranty trigger keeps the existing invoice-date rule.
    """
    if s(kind) != '质保款':
        return False, payment_term_days(term)
    trigger = s(term.get('触发条件')) if term else ''
    if term is not None and trigger not in ('', '质保期满后'):
        return False, payment_term_days(term)
    return True, payment_term_days(term)


def warranty_due_date(kind, term, warranty_end):
    """Return the effective due datetime for a warranty-expiry term."""
    uses_warranty, days = warranty_term_settings(kind, term)
    if not uses_warranty or warranty_end is None:
        return None
    return warranty_end + timedelta(days=days)


def set_cell(ws, r, c, v, font=F_BODY, fill=None, align=AL_L, fmt=None,
             trusted_formula=False):
    if isinstance(v, str) and v.startswith(('=', '+', '-', '@')) and not trusted_formula:
        v = "'" + v
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font
    cell.border = BORDER
    cell.alignment = align
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    return cell


def natural_batch_key(b):
    parts = []
    for x in s(b).split('-'):
        parts.append((0, int(x)) if x.isdigit() else (1, x))
    return parts


def norm_order(v):
    """订单内容规范化：×→*，；和+ 改为单元格内换行。"""
    t = s(v)
    t = t.replace('＋', '+').replace('×', '*')
    t = t.replace('；', '\n')
    parts = [p.strip() for p in re.split(r'[+]', t)]
    return '\n'.join(p for p in parts if p)


def style_link(cell):
    cell.font = Font(color='0000FF', underline='single')


def first_run_rows(rows, keycol):
    """每个记录键在行序中的首个连续块首行（金额只在首块显示一次，避免列合计重复）。"""
    mask = set()
    seen = set()
    i = 0
    while i < len(rows):
        k = rows[i].get(keycol)
        if k is None:
            i += 1
            continue
        if k not in seen:
            mask.add(i)
            seen.add(k)
        while i < len(rows) and rows[i].get(keycol) == k:
            i += 1
    return mask


def write_job_sheet(ws, job, ctr, devs, tms, ships, shp_key, invs, pays, rules=None):
    devs = list(devs)  # 保持导出顺序 = 源表行序
    tol = {**DEFAULT_RULES, **(rules or {})}['金额容差']  # 行预警已回款判定容差,与未回收表同口径
    types = sorted({s(t['款类']) for t in tms} |
                   {s(x['款类']) for x in invs if s(x['款类']) != '全额'} |
                   {s(x['款类']) for x in pays},
                   key=lambda x: CANON.get(x, 99))
    full_invs = [x for x in invs if s(x['款类']) == '全额']
    inv_by = {t: [x for x in invs if s(x['款类']) == t] for t in types}
    pay_by = {t: [x for x in pays if s(x['款类']) == t] for t in types}
    term_by_kind = {s(t.get('款类')): t for t in tms}

    matched_sets = {}

    def match_summary(records, date_field, prefix):
        if not records:
            return None, None, None
        ids = [s(x['记录ID']) for x in records]
        key = ids[0] if len(ids) == 1 else f"{prefix}:" + ";".join(ids)
        matched_sets[key] = records
        latest = max(records, key=lambda x: to_dt(x.get(date_field)) or datetime.min)
        amount = sum(num(x.get('含税金额')) or 0 for x in records)
        return key, latest.get(date_field), amount

    rows = []
    for d in devs:
        ser = s(d['製造番号'])
        batch = s(d['发货批次'])
        ship = shp_key.get((job, batch), {})
        row = {
            '批次': batch, '批次日期': ship.get('出荷日'), '批台数': ship.get('台数'),
            '製造番号': ser, '機番': s(d['機番']), '设备型号': s(d['设备型号']),
            'PO No': s(d['PO No']), '未税单价': d['未税单价'],
            '是否无偿': s(d['是否无偿']), '验收状态': s(d['验收状态']),
            '质保开始日': d['质保开始日'], '质保结束日': d['质保结束日'],
            '质保期': s(d['质保期']), '送货单回收': s(d['送货单回收']),
            '备注': s(d['备注']),
            '_warranty_pending': False, '_warranty_rule': False,
        }
        dues = []
        free = s(d['是否无偿']) == '是'
        for t in types:
            matched_invoices = matching_records(inv_by[t], ser, batch)
            matched_payments = matching_records(pay_by[t], ser, batch)
            term = term_by_kind.get(t)
            uses_warranty, _ = warranty_term_settings(t, term)
            warranty_end = to_dt(d.get('质保结束日')) if uses_warranty and not free else None
            if uses_warranty and not free:
                row['_warranty_rule'] = True
            if uses_warranty and not free and warranty_end is None:
                row['_warranty_pending'] = True
            inv_key, inv_date, inv_amount = match_summary(
                matched_invoices, '开票日', f'invoice:{t}'
            )
            pay_key, pay_date, pay_amount = match_summary(
                matched_payments, '回款日', f'payment:{t}'
            )
            if matched_invoices:
                if uses_warranty:
                    effective_due = warranty_due_date(t, term, warranty_end)
                    dues.extend(effective_due for _ in matched_invoices)
                else:
                    dues.extend(to_dt(inv.get('应收回款日')) for inv in matched_invoices)
                row[t + '·开票日'] = inv_date
                row[t + '·开票金额'] = inv_amount
                row[t + '·开票状况'] = '已开票'
                row[t + '·开票键'] = inv_key
            else:
                row[t + '·开票日'] = None
                row[t + '·开票金额'] = None
                row[t + '·开票状况'] = '无偿不开票' if free else '未开票'
                row[t + '·开票键'] = None
            if not matched_invoices and full_invs:
                fallback_invoices = matching_records(full_invs, ser, batch)
                if uses_warranty:
                    effective_due = warranty_due_date(t, term, warranty_end)
                    dues.extend(effective_due for _ in fallback_invoices)
                else:
                    dues.extend(to_dt(inv.get('应收回款日')) for inv in fallback_invoices)
            if matched_payments:
                row[t + '·回款日'] = pay_date
                row[t + '·回款金额'] = pay_amount
                row[t + '·回款键'] = pay_key
            else:
                row[t + '·回款日'] = None
                row[t + '·回款金额'] = None
                row[t + '·回款键'] = None
        # 全额发票块
        matched_full = matching_records(full_invs, ser, batch)
        full_key, full_date, full_amount = match_summary(matched_full, '开票日', 'invoice:全额')
        if matched_full:
            row['全额·开票日'] = full_date
            row['全额·开票金额'] = full_amount
            row['全额·开票状况'] = '已开票'
            row['全额·开票键'] = full_key
        else:
            row['全额·开票日'] = None
            row['全额·开票金额'] = None
            row['全额·开票状况'] = '无偿不开票' if free else '未开票'
            row['全额·开票键'] = None
        row['_max_due'] = max((d for d in dues if d), default=None)
        rows.append(row)

    # 无设备但有开票/回款（如 26BS009 未制造）：生成一行汇总显示，保证金额可见且合计公式可对上
    if not rows and (invs or pays):
        syn = {'批次': '', '批次日期': None, '批台数': None, '製造番号': '', '機番': '',
               '设备型号': '', 'PO No': '', '未税单价': 0, '是否无偿': '否', '验收状态': '',
               '质保开始日': None, '质保结束日': None, '质保期': '', '送货单回收': '',
               '备注': '暂无设备/发货（源事实）；开票/回款仅作汇总显示',
               '_max_due': None, '_warranty_pending': False, '_warranty_rule': False}
        dues = []
        for t in types:
            type_invoices = inv_by[t]
            type_payments = pay_by[t]
            inv = max(type_invoices, key=lambda x: to_dt(x.get('开票日')) or datetime.min,
                      default=None)
            pay = max(type_payments, key=lambda x: to_dt(x.get('回款日')) or datetime.min,
                      default=None)
            if inv is not None:
                syn[t + '·开票日'] = inv['开票日']
                syn[t + '·开票金额'] = sum(num(x.get('含税金额')) or 0 for x in type_invoices)
                syn[t + '·开票状况'] = '已开票'
                syn[t + '·开票键'] = 'synthetic:' + ';'.join(s(x['记录ID']) for x in type_invoices)
                dues.extend(to_dt(x.get('应收回款日')) for x in type_invoices)
            else:
                syn[t + '·开票日'] = syn[t + '·开票金额'] = None
                syn[t + '·开票状况'] = ''
                syn[t + '·开票键'] = None
            if pay is not None:
                syn[t + '·回款日'] = pay['回款日']
                syn[t + '·回款金额'] = sum(num(x.get('含税金额')) or 0 for x in type_payments)
                syn[t + '·回款键'] = 'synthetic:' + ';'.join(s(x['记录ID']) for x in type_payments)
            else:
                syn[t + '·回款日'] = syn[t + '·回款金额'] = None
                syn[t + '·回款键'] = None
        if full_invs:
            invf = max(full_invs, key=lambda x: to_dt(x.get('开票日')) or datetime.min)
            syn['全额·开票日'] = invf['开票日']
            syn['全额·开票金额'] = sum(num(x.get('含税金额')) or 0 for x in full_invs)
            syn['全额·开票状况'] = '已开票'
            syn['全额·开票键'] = 'synthetic:' + ';'.join(s(x['记录ID']) for x in full_invs)
            dues.extend(to_dt(x.get('应收回款日')) for x in full_invs)
        else:
            syn['全额·开票日'] = syn['全额·开票金额'] = None
            syn['全额·开票状况'] = ''
            syn['全额·开票键'] = None
        syn['_max_due'] = max((d for d in dues if d), default=None)
        rows = [syn]

    # 回款整段合并：同一笔回款中间夹免费/未覆盖行时连成一段
    for t in types:
        filled = list(rows[i][t + '·回款键'] for i in range(len(rows)))
        for i in range(len(rows)):
            if filled[i] is not None:
                continue
            j = i
            while j < len(rows) and filled[j] is None:
                j += 1
            if i > 0 and j < len(rows) and filled[i - 1] == filled[j]:
                for k in range(i, j):
                    filled[k] = filled[j]
        for i in range(len(rows)):
            rows[i][t + '·回款键'] = filled[i]

    # 列定义
    cols = [
        ('发货批次', '批次', '批次', None, AL_L),
        ('设备型号', '设备型号', None, None, AL_L),
        ('PO No', 'PO No', None, None, AL_L),
        ('製造番号', '製造番号', None, None, AL_L),
        ('機番', '機番', None, None, AL_C),
        ('未税单价', '未税单价', None, FMT_PRICE, AL_R),
        ('出荷日', '批次日期', '批次', FMT_DATE, AL_C),
        ('出荷台数', '批台数', '批次', '0', AL_C),
        ('送货单回收', '送货单回收', None, None, AL_C),
    ]
    blocks = []  # (kind, type_or_None)
    if full_invs:
        blocks.append(('full', None))
    for t in types:
        blocks.append(('type', t, 5 if inv_by[t] else 2))
    for kind, t, *_ in blocks:
        if kind == 'full':
            cols += [
                ('全额·开票日', '全额·开票日', '全额·开票键', FMT_DATE, AL_C),
                ('全额·开票状况', '全额·开票状况', '全额·开票状况键', None, AL_C),
                ('全额·开票金额', '全额·开票金额', '全额·开票键', FMT_MONEY, AL_R),
            ]
        else:
            if len(_) and _[0] == 5:
                cols += [
                    (t + '·开票日', t + '·开票日', t + '·开票键', FMT_DATE, AL_C),
                    (t + '·开票状况', t + '·开票状况', t + '·开票状况键', None, AL_C),
                    (t + '·开票金额', t + '·开票金额', t + '·开票键', FMT_MONEY, AL_R),
                    (t + '·回款日', t + '·回款日', t + '·回款键', FMT_DATE, AL_C),
                    (t + '·回款金额', t + '·回款金额', t + '·回款键', FMT_MONEY, AL_R),
                ]
            else:
                cols += [
                    (t + '·回款日', t + '·回款日', t + '·回款键', FMT_DATE, AL_C),
                    (t + '·回款金额', t + '·回款金额', t + '·回款键', FMT_MONEY, AL_R),
                ]
    cols += [
        ('是否无偿', '是否无偿', None, None, AL_C),
        ('验收状态', '验收状态', None, None, AL_C),
        ('质保开始日', '质保开始日', None, FMT_DATE, AL_C),
        ('质保结束日', '质保结束日', None, FMT_DATE, AL_C),
        ('质保期', '质保期', None, None, AL_C),
        ('备注', '备注', None, None, AL_L),
    ]
    for r_ in rows:
        for t in types:
            r_[t + '·开票状况键'] = r_[t + '·开票键'] or ('none', r_[t + '·开票状况'])
        r_['全额·开票状况键'] = r_['全额·开票键'] or ('none', r_['全额·开票状况'])

    ncol = len(cols)
    # 第1行
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    set_cell(ws, 1, 1, '受注管理情報', font=F_TITLE, align=AL_C)
    for c in range(2, ncol + 1):
        ws.cell(row=1, column=c).border = BORDER
    # 第2~3行
    head2 = [(1, '担当者'), (2, '取引先'), (3, 'PO NO.'), (4, 'JOB NO'),
             (5, '注文内容'), (7, '支払条件'), (11, '発送方式'), (12, '出荷先'), (13, '納入先')]
    for c, lab in head2:
        set_cell(ws, 2, c, lab, font=F_BOLD, fill=FILL_INFO, align=AL_C)
    po_list = '\n'.join(sorted({s(d['PO No']) for d in devs if s(d['PO No'])}))
    val2 = [(1, s(ctr['担当者'])), (2, s(ctr['客户'])), (3, po_list),
            (4, s(ctr['JOB No'])), (5, norm_order(ctr['订单内容'])), (7, s(ctr['付款条件'])),
            (11, s(ctr['发货方式'])), (12, s(ctr['发货地点'])), (13, s(ctr['送货地点']))]
    for c, v in val2:
        set_cell(ws, 3, c, v, align=AL_L)
    ws.merge_cells(start_row=3, start_column=5, end_row=3, end_column=6)
    ws.merge_cells(start_row=3, start_column=7, end_row=3, end_column=10)
    order_lines = norm_order(ctr['订单内容']).count('\n') + 1
    ws.row_dimensions[3].height = max(16, 14 * (po_list.count('\n') + 1), 14 * order_lines)
    # 第4行分组
    n_ship = 9
    n_pay = len(cols) - n_ship - 6
    groups = [('出荷状況', 1, n_ship, 'DDEBF7'),
              ('回収状況', n_ship + 1, n_ship + n_pay, 'FCE4D6'),
              ('其他情報', n_ship + n_pay + 1, ncol, 'EDEDED')]
    for gname, c1, c2, color in groups:
        ws.merge_cells(start_row=4, start_column=c1, end_row=4, end_column=c2)
        set_cell(ws, 4, c1, gname, font=F_BOLD,
                 fill=PatternFill('solid', fgColor=color), align=AL_C)
        for c in range(c1 + 1, c2 + 1):
            ws.cell(row=4, column=c).fill = PatternFill('solid', fgColor=color)
            ws.cell(row=4, column=c).border = BORDER
    # 第5行列头
    for c, (h, *_rest) in enumerate(cols, 1):
        set_cell(ws, 5, c, h, font=F_HEAD, fill=FILL_HEAD, align=AL_C)
    ws.freeze_panes = 'A6'
    # 数据
    amt_masks = {}
    for h, key, mkey, fmt, align in cols:
        if h.endswith('·开票金额') or h.endswith('·回款金额'):
            amt_masks[h] = first_run_rows(rows, mkey)
    r = 6
    for idx, row in enumerate(rows):
        for c, (h, key, mkey, fmt, align) in enumerate(cols, 1):
            v = row.get(key)
            if h.endswith('·开票日') or h.endswith('·回款日'):
                v = fmt_date(v)
            elif h.endswith('·开票金额') or h.endswith('·回款金额'):
                v = num(v)
                if idx not in amt_masks[h]:
                    v = None
            elif key == '未税单价':
                v = num(v)
            elif key in ('批次日期', '质保开始日', '质保结束日'):
                v = fmt_date(v)
            set_cell(ws, r, c, v, align=align, fmt=fmt)
        r += 1
    # 纵向合并
    for c, (h, key, mkey, fmt, align) in enumerate(cols, 1):
        if not mkey:
            continue
        i = 0
        while i < len(rows):
            k = rows[i].get(mkey)
            if k is None:
                i += 1
                continue
            j = i
            while j + 1 < len(rows) and rows[j + 1].get(mkey) == k:
                j += 1
            if j > i:
                ws.merge_cells(start_row=6 + i, start_column=c,
                               end_row=6 + j, end_column=c)
            i = j + 1
    # 列宽
    widths = [12, 16, 20, 14, 11, 12, 12, 9, 14]
    for kind, t, *_ in blocks:
        if kind == 'full':
            widths += [12, 11, 13]
        else:
            widths += [12, 11, 13, 12, 13] if (_ and _[0] == 5) else [12, 13]
    widths += [9, 10, 12, 12, 8, 18]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    # ---- 隐藏辅助列（全部用公式引用可见列，改入金后自动联动）----
    col_letter = {h: get_column_letter(c) for c, (h, *_rest) in enumerate(cols, 1)}
    price_c = col_letter['未税单价']
    hc = ncol + 1
    batch_col = get_column_letter(hc)
    hidden_specs = [(hc, '批次行', 'batch', None)]
    hc += 1
    inv_row_cols, pay_row_cols, pay_row_map = [], [], {}
    for kind, t, *_ in blocks:
        if kind == 'full' or (_ and _[0] == 5):
            letter = get_column_letter(hc)
            inv_row_cols.append(letter)
            hidden_specs.append((hc, f'行开票金额·{"全额" if kind == "full" else t}',
                                 '开票', '全额' if kind == 'full' else t))
            hc += 1
        if kind == 'type':
            letter = get_column_letter(hc)
            pay_row_cols.append(letter)
            pay_row_map[t] = letter
            hidden_specs.append((hc, f'行回款金额·{t}', '回款', t))
            hc += 1
    due_col = get_column_letter(hc)
    hidden_specs.append((hc, '行应收回款日', 'due', None))
    hc += 1
    warranty_pending_col = get_column_letter(hc)
    hidden_specs.append((hc, '行质保待确认', 'warranty_pending', None))
    hc += 1
    warranty_rule_col = get_column_letter(hc)
    hidden_specs.append((hc, '行质保规则', 'warranty_rule', None))
    hc += 1
    warn_col = get_column_letter(hc)
    hidden_specs.append((hc, '行预警', 'warn', None))

    for cidx, hname, *_ in hidden_specs:
        set_cell(ws, 5, cidx, hname, font=Font(size=9, color='808080'),
                 fill=PatternFill('solid', fgColor='F5F5F5'), align=AL_C)
    # 每笔发票/回款覆盖的有偿设备未税总价（公式中作为分摊分母的静态常量）
    def cover_price(cands):
        out = {}
        for x in cands:
            tp = 0.0
            for d in devs:
                if s(d['是否无偿']) != '是' and cov_match(x, s(d['製造番号']), s(d['发货批次'])):
                    tp += num(d['未税单价']) or 0
            out[s(x['记录ID'])] = tp
        return out
    inv_price = {t: cover_price(inv_by[t]) for t in types}
    pay_price = {t: cover_price(pay_by[t]) for t in types}
    full_price = cover_price(full_invs)
    # 每笔发票/回款在可见金额列的首个金额块锚点行（后续分段为空，公式统一引用锚点）
    inv_anchor, pay_anchor = {}, {}
    for t in types + (['全额'] if full_invs else []):
        seen = {}
        for idx, row in enumerate(rows):
            k = row.get(t + '·开票键')
            if k and k not in seen:
                seen[k] = idx
                inv_anchor[(t, k)] = idx
    for t in types:
        seen = {}
        for idx, row in enumerate(rows):
            k = row.get(t + '·回款键')
            if k and k not in seen:
                seen[k] = idx
                pay_anchor[(t, k)] = idx

    for idx, row in enumerate(rows):
        r = 6 + idx
        for cidx, hname, kind, t in hidden_specs:
            if kind == 'batch':
                set_cell(ws, r, cidx, row.get('批次', ''), align=AL_C)
            elif kind == '开票':
                if t == '全额':
                    st_c = col_letter['全额·开票状况']
                    amt_c = col_letter['全额·开票金额']
                    denom = full_price.get(row.get('全额·开票键'), 0.0)
                    anchor = inv_anchor.get(('全额', row.get('全额·开票键')), idx)
                else:
                    st_c = col_letter[t + '·开票状况']
                    amt_c = col_letter[t + '·开票金额']
                    denom = inv_price[t].get(row.get(t + '·开票键'), 0.0)
                    anchor = inv_anchor.get((t, row.get(t + '·开票键')), idx)
                records = matched_sets.get(row.get('全额·开票键' if t == '全额'
                                                   else t + '·开票键'), [])
                if len(records) > 1:
                    prices = full_price if t == '全额' else inv_price[t]
                    allocated = sum((num(x.get('含税金额')) or 0)
                                    * (num(row.get('未税单价')) or 0)
                                    / prices.get(s(x['记录ID']), 1)
                                    for x in records if prices.get(s(x['记录ID']), 0))
                    set_cell(ws, r, cidx, allocated, align=AL_R, fmt='0.00')
                else:
                    formula = (f'=IF({st_c}{6 + anchor}="已开票",'
                               f'{amt_c}{6 + anchor}*{price_c}{r}/{denom},0)'
                               if denom else '=0')
                    set_cell(ws, r, cidx, formula, align=AL_R, fmt='0.00', trusted_formula=True)
            elif kind == '回款':
                d_c = col_letter[t + '·回款日']
                a_c = col_letter[t + '·回款金额']
                denom = pay_price[t].get(row.get(t + '·回款键'), 0.0)
                anchor = pay_anchor.get((t, row.get(t + '·回款键')), idx)
                records = matched_sets.get(row.get(t + '·回款键'), [])
                if len(records) > 1:
                    allocated = sum((num(x.get('含税金额')) or 0)
                                    * (num(row.get('未税单价')) or 0)
                                    / pay_price[t].get(s(x['记录ID']), 1)
                                    for x in records if pay_price[t].get(s(x['记录ID']), 0))
                    set_cell(ws, r, cidx, allocated, align=AL_R, fmt='0.00')
                else:
                    formula = (f'=IF({d_c}{6 + anchor}="",0,'
                               f'{a_c}{6 + anchor}*{price_c}{r}/{denom})'
                               if denom else '=0')
                    set_cell(ws, r, cidx, formula, align=AL_R, fmt='0.00', trusted_formula=True)
            elif kind == 'due':
                set_cell(ws, r, cidx, row.get('_max_due'), align=AL_C, fmt='yyyy/mm/dd')
            elif kind == 'warranty_pending':
                set_cell(ws, r, cidx, 1 if row.get('_warranty_pending') else 0,
                         align=AL_C)
            elif kind == 'warranty_rule':
                set_cell(ws, r, cidx, 1 if row.get('_warranty_rule') else 0,
                         align=AL_C)
            elif kind == 'warn':
                inv_expr = '+'.join(f'{c}{r}' for c in inv_row_cols) or '0'
                pay_expr = '+'.join(f'{c}{r}' for c in pay_row_cols) or '0'
                dc = f'{due_col}{r}'
                pc = f'{warranty_pending_col}{r}'
                rc = f'{warranty_rule_col}{r}'
                wc = f'{col_letter["质保结束日"]}{r}'
                formula = (f'=IF({inv_expr}=0,"",'
                           f'IF({pc},"待确认",'
                           f'IF(AND({rc}=1,{wc}<>"",TODAY()<={wc}),"未到期",'
                           f'IF({pay_expr}+{tol:g}>={inv_expr},"已回款",'
                           f'IF({dc}="","未到期",IF({dc}<TODAY(),"逾期","未到期"))))))')
                set_cell(ws, r, cidx, formula, align=AL_C, trusted_formula=True)
    for cidx, *_ in hidden_specs:
        ws.column_dimensions[get_column_letter(cidx)].hidden = True

    # 返回导航页超链接（数据最后一行之后）
    back_row = 6 + len(rows)
    ws.merge_cells(start_row=back_row, start_column=1, end_row=back_row, end_column=ncol)
    set_cell(ws, back_row, 1, f'=HYPERLINK("#\'导航页\'!A1","← 返回导航页")',
             font=F_BOLD, fill=PatternFill('solid', fgColor='EDEDED'), align=AL_C,
             trusted_formula=True)
    style_link(ws.cell(back_row, 1))

    info = {
        'inv_amount_cols': [get_column_letter(c) for c, (h, *_rest) in enumerate(cols, 1)
                            if h.endswith('·开票金额')],
        'pay_amount_cols': [get_column_letter(c) for c, (h, *_rest) in enumerate(cols, 1)
                            if h.endswith('·回款金额')],
        'pay_row_map': pay_row_map,
        'batch_col': batch_col,
        'warn_col': warn_col,
        'n_rows': len(rows),
    }
    return rows, cols, blocks, info


def compute_unpaid_rows(contracts, terms, shipments, invoices, payments, devices, rules=None):
    """计算未回收行（含应收常量、静态原因与预警豁免标记）。"""
    rules = {**DEFAULT_RULES, **(rules or {})}
    tol = rules['金额容差']
    days = rules['临近天数']
    today = datetime.now()
    rows = []
    jobs = sorted({s(c['JOB No']) for c in contracts})
    terms_by, ships_by, inv_by, pay_by, dev_by = {}, {}, {}, {}, {}
    for t in terms:
        terms_by.setdefault(s(t['JOB No']), []).append(t)
    for x in shipments:
        ships_by.setdefault(s(x['JOB No']), []).append(x)
    for x in invoices:
        inv_by.setdefault(s(x['JOB No']), []).append(x)
    for x in payments:
        pay_by.setdefault(s(x['JOB No']), []).append(x)
    for x in devices:
        dev_by.setdefault(s(x['JOB No']), []).append(x)
    cust = {s(c['JOB No']): s(c['客户']) for c in contracts}

    for job in jobs:
        tms = terms_by.get(job, [])
        invs = inv_by.get(job, [])
        pays = pay_by.get(job, [])
        devs = dev_by.get(job, [])
        types = sorted({s(t['款类']) for t in tms} |
                       {s(x['款类']) for x in invs if s(x['款类']) != '全额'} |
                       {s(x['款类']) for x in pays},
                       key=lambda x: CANON.get(x, 99))
        if not types:
            continue
        rates = {}
        for t in tms:
            rates[s(t['款类'])] = (num(t['比例%']) or 0) / 100.0
        term_by_kind = {s(t.get('款类')): t for t in tms}
        full_invs = [x for x in invs if s(x['款类']) == '全额']
        inv_by_t = {t: [x for x in invs if s(x['款类']) == t] for t in types}
        pay_by_t = {t: [x for x in pays if s(x['款类']) == t] for t in types}
        ship_batches = {s(x['发货批次']) for x in ships_by.get(job, [])}

        # 无设备但有回款/条款（如 26BS009）：按已付款比例反推合同总额
        if not devs and (pays or invs):
            estimates = []
            full_invoice_total = sum(num(x.get('含税金额')) or 0 for x in full_invs)
            if full_invoice_total:
                estimates.append(full_invoice_total)
            for t in types:
                rt = rates.get(t)
                if not rt:
                    continue
                invoiced_t = sum(num(x.get('含税金额')) or 0 for x in inv_by_t[t])
                paid_t = sum(num(x.get('含税金额')) or 0 for x in pay_by_t[t])
                if invoiced_t:
                    estimates.append(invoiced_t / rt)
                if paid_t:
                    estimates.append(paid_t / rt)
            total = max(estimates, default=None)
            if total:
                for t in types:
                    paid_t = sum(num(p['含税金额']) or 0 for p in pay_by_t[t])
                    unpaid = max(0.0, total * rates.get(t, 0.0) - paid_t)
                    if unpaid > tol:
                        rows.append({'客户': cust[job], 'JOB No': job, '批次': '（未制造）',
                                     '款类': t, '预警等级': '④待确认', '开票日期': '',
                                     '预定回收日期': '', '未回收金额': round(unpaid, 2),
                                     '未回收原因': '未发货；未开票；待确认',
                                     'recv': round(unpaid, 2), 'unacc': 0,
                                     'warn_exempt': 1, 'static': True, 'raw_batch': '',
                                     'inv_dt': None, 'due_dt': None,
                                     '_warranty_end': None, '_warranty_pending': False,
                                     '_warranty_rule': False})
            else:
                invoiced = sum(num(x.get('含税金额')) or 0 for x in invs)
                paid = sum(num(x.get('含税金额')) or 0 for x in pays)
                unpaid = max(0.0, invoiced - paid)
                rows.append({'客户': cust[job], 'JOB No': job, '批次': '（未制造）',
                             '款类': '（比例未配置）', '预警等级': '④待确认', '开票日期': '',
                             '预定回收日期': '', '未回收金额': round(unpaid, 2),
                                     '未回收原因': '未发货；付款比例未配置；待确认',
                                     'recv': round(unpaid, 2), 'unacc': 0,
                                     'warn_exempt': 1, 'static': True, 'raw_batch': '',
                                     'inv_dt': None, 'due_dt': None,
                                     '_warranty_end': None, '_warranty_pending': False,
                                     '_warranty_rule': False})
            continue

        # 每笔发票/回款覆盖的有偿设备未税总价（按价格占比分摊，兼容同批不同单价）
        def cover_price(cands, devs_):
            out = {}
            for x in cands:
                tp = 0.0
                for d in devs_:
                    if s(d['是否无偿']) != '是' and cov_match(x, s(d['製造番号']), s(d['发货批次'])):
                        tp += num(d['未税单价']) or 0
                out[s(x['记录ID'])] = tp
            return out
        inv_price = {t: cover_price(inv_by_t[t], devs) for t in types}
        full_inv_price = cover_price(full_invs, devs)
        pay_price = {t: cover_price(pay_by_t[t], devs) for t in types}

        agg = {}
        batch_unacc = {}
        batch_war = {}
        batch_war_missing = {}
        for d in devs:
            if s(d['是否无偿']) == '是':
                continue
            ser, batch = s(d['製造番号']), s(d['发货批次'])
            price = num(d['未税单价']) or 0
            if s(d['验收状态']) != '已验收':
                batch_unacc.setdefault(batch, 0)
                batch_unacc[batch] += 1
            wend = to_dt(d.get('质保结束日'))
            if not wend:
                batch_war_missing[batch] = True
            if wend and (batch not in batch_war or wend > batch_war[batch]):
                batch_war[batch] = wend
            for t in types:
                rt = rates.get(t)
                if not rt:
                    continue
                key = (batch, t)
                a = agg.setdefault(key, {'应收': 0.0, '已开': 0.0, '已回': 0.0,
                                         'inv_dates': [], 'dues': [], 'inv_bad': False})
                recv = price * 1.13 * rt
                a['应收'] += recv
                matched_invoices = matching_records(inv_by_t[t], ser, batch)
                if not matched_invoices and full_invs:
                    matched_invoices = matching_records(full_invs, ser, batch)
                for inv in matched_invoices:
                    rid = s(inv['记录ID'])
                    if s(inv['款类']) == '全额':
                        tp = full_inv_price.get(rid, 0.0)
                        a['已开'] += ((num(inv['含税金额']) or 0) * price / tp * rt
                                    if tp else 0.0)
                    else:
                        tp = inv_price[t].get(rid, 0.0)
                        a['已开'] += (num(inv['含税金额']) or 0) * price / tp if tp else 0.0
                    a['inv_dates'].append(to_dt(inv.get('开票日')))
                    a['dues'].append(to_dt(inv.get('应收回款日')))
                    if s(inv.get('状态')) != '已开':
                        a['inv_bad'] = True
                for pay in matching_records(pay_by_t[t], ser, batch):
                    rid = s(pay['记录ID'])
                    tp = pay_price[t].get(rid, 0.0)
                    a['已回'] += (num(pay['含税金额']) or 0) * price / tp if tp else 0.0

        # 批次级封顶：整批已回 ≥ 整批应收时（如人工整台付清记录），该批不再列未回收
        batch_recv, batch_paid = {}, {}
        for (batch, t), a in agg.items():
            batch_recv[batch] = batch_recv.get(batch, 0.0) + a['应收']
            batch_paid[batch] = batch_paid.get(batch, 0.0) + a['已回']
        skip_batches = {b for b in batch_recv if batch_paid.get(b, 0.0) >= batch_recv[b] - tol}

        for (batch, t), a in sorted(agg.items(),
                                    key=lambda kv: (cust[job], job, natural_batch_key(kv[0][0]),
                                                    CANON.get(kv[0][1], 99))):
            if batch in skip_batches:
                continue
            unpaid = max(0.0, a['应收'] - a['已回'])
            if unpaid <= tol:
                continue
            uninv = max(0.0, a['应收'] - a['已开'])
            due = max((d for d in a['dues'] if d), default=None)
            inv_date = max((d for d in a['inv_dates'] if d), default=None)
            term = term_by_kind.get(t)
            uses_warranty, _ = warranty_term_settings(t, term)
            warranty_end = batch_war.get(batch)
            warranty_missing = uses_warranty and (
                batch_war_missing.get(batch, False) or warranty_end is None
            )
            warranty_active = bool(
                uses_warranty and warranty_end and
                today.date() <= warranty_end.date()
            )
            if uses_warranty:
                due = warranty_due_date(t, term, warranty_end)
            reasons = []
            if batch not in ship_batches:
                reasons.append('未发货')
            if uninv > tol:
                reasons.append('未开票')
            if t == '验收款' and batch_unacc.get(batch, 0) > 0:
                reasons.append('未验收')
            if warranty_missing:
                reasons.append('质保期未设置')
            elif warranty_active:
                reasons.append('质保未到期')
            pend = (('未验收' in reasons) and rules['未验收待确认']) or \
                   (a['inv_bad'] and rules['发票异常待确认']) or due is None
            if warranty_missing:
                pend = True
            if pend and '待确认' not in reasons:
                reasons.append('待确认')
            elif warranty_active:
                pass
            elif due and due.date() < today.date():
                reasons.append('逾期')
            elif due and due <= today.replace(hour=0, minute=0, second=0,
                                              microsecond=0) + timedelta(days=days):
                reasons.append('临近')
            elif due:
                reasons.append('未到付款期')
            if (('未验收' in reasons) and rules['未验收待确认']) or \
               warranty_missing or due is None:
                warn = '④待确认'
            elif warranty_active:
                warn = '③未到期'
            elif due and due.date() < today.date():
                warn = '①已逾期'
            elif due and due <= today + timedelta(days=days):
                warn = '②临近'
            else:
                warn = '③未到期'
            rows.append({'客户': cust[job], 'JOB No': job, '批次': batch or '（无批次）',
                         '款类': t, '预警等级': warn,
                         '开票日期': fmt_date(inv_date) if inv_date else '',
                         '预定回收日期': fmt_date(due) if due else '',
                         '未回收金额': round(unpaid, 2),
                         '未回收原因': '；'.join(reasons),
                         'recv': a['应收'], 'unacc': 1 if '未验收' in reasons else 0,
                         'warn_exempt': 1 if pend else 0,
                         'static': False, 'raw_batch': batch,
                         'inv_dt': inv_date, 'due_dt': due,
                         '_warranty_end': warranty_end if uses_warranty else None,
                         '_warranty_pending': warranty_missing,
                         '_warranty_rule': uses_warranty})

    rows.sort(key=lambda r: (r['客户'], r['JOB No'], natural_batch_key(r['批次']),
                             CANON.get(r['款类'], 99)))
    return rows


def build_unpaid_sheet(wb, contracts, terms, shipments, invoices, payments, devices,
                       job_infos, rules=None):
    """未回收管理表（sheet 2）：客户合并 → JOB 合并 → 批次/款类；金额/预警/原因均为公式。"""
    rules = {**DEFAULT_RULES, **(rules or {})}
    tol = rules['金额容差']
    days = rules['临近天数']
    ws = wb.create_sheet('未回收管理表', 1)
    headers = ['客户', 'JOB No', '批次', '款类', '预警等级', '开票日期',
               '预定回收日期', '未回收金额', '未回收原因']
    ncol = len(headers)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    set_cell(ws, 1, 1, '未回收管理表', font=F_TITLE, align=AL_C)
    for c, h in enumerate(headers, 1):
        set_cell(ws, 2, c, h, font=F_HEAD, fill=FILL_HEAD, align=AL_C)
    ws.freeze_panes = 'A3'
    widths = [26, 10, 12, 10, 10, 12, 13, 14, 42]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    warranty_end_col = get_column_letter(ncol + 1)
    warranty_pending_col = get_column_letter(ncol + 2)
    warranty_rule_col = get_column_letter(ncol + 3)
    for c, h in ((ncol + 1, '辅助质保结束日'),
                 (ncol + 2, '辅助质保待确认'),
                 (ncol + 3, '辅助质保规则')):
        set_cell(ws, 2, c, h, font=Font(size=9, color='808080'),
                 fill=PatternFill('solid', fgColor='F5F5F5'), align=AL_C)
        ws.column_dimensions[get_column_letter(c)].hidden = True

    rows = compute_unpaid_rows(contracts, terms, shipments, invoices, payments, devices, rules)
    r = 3
    cust_start, job_start, prev_cust, prev_job = r, r, None, None
    warn_fill = {'①已逾期': PatternFill('solid', fgColor='FBE2E2'),
                 '②临近': PatternFill('solid', fgColor='FFF2CC')}
    for row in rows:
        job = row['JOB No']
        info = job_infos.get(job)
        if row['static'] or info is None:
            vals = [row['客户'], f'=HYPERLINK("#\'{job}\'!A1","{job}")',
                    row['批次'], row['款类'], row['预警等级'], row['开票日期'],
                    row['预定回收日期'], row['未回收金额'], row['未回收原因']]
        else:
            pay_c = info['pay_row_map'][row['款类']]
            batch_c = info['batch_col']
            last = 5 + info['n_rows']
            crit = f'"{row["raw_batch"]}"' if row['raw_batch'] else '""'
            amt_f = (f"=MAX(0,ROUND({row['recv']:.4f}-"
                     f"SUMIFS('{job}'!${pay_c}$6:${pay_c}${last},"
                     f"'{job}'!${batch_c}$6:${batch_c}${last},$C{r}),2))")
            warn_f = (f'=IF($H{r}<={tol:g},"已回清",'
                      f'IF(${warranty_pending_col}{r},"④待确认",'
                      f'IF(AND(${warranty_rule_col}{r}=1,${warranty_end_col}{r}<>"",'
                      f'TODAY()<=${warranty_end_col}{r}),"③未到期",'
                      f'IF({row["warn_exempt"]},"④待确认",'
                      f'IF($G{r}="","④待确认",'
                      f'IF($G{r}<TODAY(),"①已逾期",'
                      f'IF($G{r}<=TODAY()+{days},"②临近","③未到期"))))))')
            reason_f = f'=IF($H{r}<={tol:g},"","{row["未回收原因"]}")'
            vals = [row['客户'], f'=HYPERLINK("#\'{job}\'!A1","{job}")',
                    row['批次'], row['款类'], warn_f, row['inv_dt'],
                    row['due_dt'], amt_f, reason_f]
        for c, v in enumerate(vals, 1):
            fmt = None
            align = AL_L
            if c == 1:
                align = AL_C
            elif c in (2, 3, 4, 5):
                align = AL_C
            elif c in (6, 7):
                fmt = FMT_DATE
                align = AL_C
            elif c == 8:
                fmt = FMT_MONEY
                align = AL_R
            elif c == 9:
                align = AL_L
            fill = warn_fill.get(row['预警等级']) if c == 5 else None
            if row['static'] and c == 5:
                fill = warn_fill.get(row['预警等级'])
            trusted_formula = c == 2 or (not row['static'] and c in (5, 8, 9))
            set_cell(ws, r, c, v, align=align, fmt=fmt, fill=fill,
                     trusted_formula=trusted_formula)
        set_cell(ws, r, ncol + 1, row.get('_warranty_end'), align=AL_C, fmt=FMT_DATE)
        set_cell(ws, r, ncol + 2, 1 if row.get('_warranty_pending') else 0, align=AL_C)
        set_cell(ws, r, ncol + 3, 1 if row.get('_warranty_rule') else 0, align=AL_C)
        style_link(ws.cell(r, 2))
        if row['客户'] != prev_cust:
            if prev_cust is not None and cust_start < r - 1:
                ws.merge_cells(start_row=cust_start, start_column=1,
                               end_row=r - 1, end_column=1)
            cust_start = r
            prev_cust = row['客户']
        if row['JOB No'] != prev_job:
            if prev_job is not None and job_start < r - 1:
                ws.merge_cells(start_row=job_start, start_column=2,
                               end_row=r - 1, end_column=2)
            job_start = r
            prev_job = row['JOB No']
        r += 1
    if prev_cust is not None and cust_start < r - 1:
        ws.merge_cells(start_row=cust_start, start_column=1, end_row=r - 1, end_column=1)
    if prev_job is not None and job_start < r - 1:
        ws.merge_cells(start_row=job_start, start_column=2, end_row=r - 1, end_column=2)
    # 合计
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    set_cell(ws, r, 1, f'合计（{len(rows)} 条）', font=F_BOLD,
             fill=FILL_TOTAL, align=AL_R)
    for c in range(2, 8):
        ws.cell(row=r, column=c).fill = FILL_TOTAL
        ws.cell(row=r, column=c).border = BORDER
    set_cell(ws, r, 8, f'=SUM(H3:H{r - 1})', font=F_BOLD, fill=FILL_TOTAL,
             align=AL_R, fmt=FMT_MONEY, trusted_formula=True)
    set_cell(ws, r, 9, '', fill=FILL_TOTAL)
    return len(rows)


def build():
    contracts = load_table([EXPORT_DIR + '/contracts.json'])
    terms = load_table([EXPORT_DIR + '/terms.json'])
    shipments = load_table([EXPORT_DIR + '/shipments.json'])
    invoices = load_table([EXPORT_DIR + '/invoices.json'])
    payments = load_table([EXPORT_DIR + '/payments.json'])
    devices = []
    for p in sorted(os.listdir(EXPORT_DIR + '/devices')):
        devices += load_table([EXPORT_DIR + '/devices/' + p])
    return build_from_data(contracts, terms, devices, shipments, invoices, payments, OUT)


def build_from_data(contracts, terms, devices, shipments, invoices, payments, out=OUT,
                    rules=None):
    """从内存六表数据生成完整台账工作簿（含导航页、未回收管理表、各 JOB 页）。"""
    JOBS = sorted({s(c['JOB No']) for c in contracts})

    by_job = {s(c['JOB No']): c for c in contracts}
    terms_by, ships_by, inv_by, pay_by, dev_by = {}, {}, {}, {}, {}
    for t in terms:
        terms_by.setdefault(s(t['JOB No']), []).append(t)
    for x in shipments:
        ships_by.setdefault(s(x['JOB No']), []).append(x)
    for x in invoices:
        inv_by.setdefault(s(x['JOB No']), []).append(x)
    for x in payments:
        pay_by.setdefault(s(x['JOB No']), []).append(x)
    for x in devices:
        dev_by.setdefault(s(x['JOB No']), []).append(x)
    shp_key = {(s(x['JOB No']), s(x['发货批次'])): x for x in shipments}

    wb = Workbook()
    ws_nav = wb.active
    ws_nav.title = '导航页'
    ncol_nav = 13
    ws_nav.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol_nav)
    set_cell(ws_nav, 1, 1, '上海康肯销售订单管理系统 — 导航页', font=F_TITLE, align=AL_C)
    nav_headers = ['客户', 'JOB No', 'PO', '担当者', '订单内容', '总台数', '发货批次数',
                   '设备含税合计', '开票金额合计', '已回款合计', '回款率', '预警', '备注']
    for c, h in enumerate(nav_headers, 1):
        set_cell(ws_nav, 2, c, h, font=F_HEAD, fill=FILL_HEAD, align=AL_C)
    ws_nav.freeze_panes = 'A3'
    nav_widths = [24, 10, 26, 10, 36, 8, 10, 14, 14, 14, 8, 14, 20]
    for c, w in enumerate(nav_widths, 1):
        ws_nav.column_dimensions[chr(64 + c)].width = w

    job_infos = {}
    for job in JOBS:
        ws = wb.create_sheet(job)
        rows, cols, blocks, info = write_job_sheet(
            ws, job, by_job[job], dev_by.get(job, []), terms_by.get(job, []),
            ships_by.get(job, []), shp_key, inv_by.get(job, []), pay_by.get(job, []),
            rules=rules)
        job_infos[job] = info

    n_unpaid = build_unpaid_sheet(wb, contracts, terms, shipments, invoices, payments,
                                  devices, job_infos, rules=rules)

    nav_rows = []
    for job in JOBS:
        ctr = by_job[job]
        po_list = '\n'.join(sorted({s(d['PO No']) for d in dev_by.get(job, []) if s(d['PO No'])}))
        nav_rows.append((s(ctr['客户']), job, po_list, s(ctr['担当者']), norm_order(ctr['订单内容']),
                         s(ctr['备注'])))
    nav_rows.sort(key=lambda x: (x[0], x[1]))

    nav_row = 3
    cust_start = nav_row
    prev_cust = None
    for i, (cust, job, po, tanto, content, remark) in enumerate(nav_rows):
        info = job_infos[job]
        last_data = 5 + info['n_rows']
        inv_sum = '+'.join(f"SUM('{job}'!{c}6:{c}{last_data})" for c in info['inv_amount_cols']) or '0'
        pay_sum = '+'.join(f"SUM('{job}'!{c}6:{c}{last_data})" for c in info['pay_amount_cols']) or '0'
        warn_rng = f"'{job}'!${info['warn_col']}$6:${info['warn_col']}${last_data}"
        vals = [
            cust,
            f'=HYPERLINK("#\'{job}\'!A1","{job}")',
            po, tanto, content,
            f"=COUNTA('{job}'!D6:D{last_data})",
            f"=COUNTA('{job}'!A6:A{last_data})",
            f"=SUM('{job}'!F6:F{last_data})*1.13",
            f'={inv_sum}',
            f'={pay_sum}',
            f'=IF(I{nav_row}=0,"",J{nav_row}/I{nav_row})',
            f'=IF(COUNTIF({warn_rng},"*逾期*")>0,"有逾期",'
            f'IF(COUNTIF({warn_rng},"*未到期*")>0,"部分未回/关注",'
            f'IF(COUNTIF({warn_rng},"*已回款*")>0,"已回款","无开票")))',
            remark,
        ]
        for c, v in enumerate(vals, 1):
            fmt = None
            align = AL_L
            if c == 1:
                align = AL_C
            elif c in (6, 7):
                fmt = '0'
                align = AL_C
            elif c in (8, 9, 10):
                fmt = FMT_MONEY
                align = AL_R
            elif c == 11:
                fmt = '0.0%'
                align = AL_C
            elif c == 12:
                align = AL_C
            set_cell(ws_nav, nav_row, c, v, align=align, fmt=fmt,
                     trusted_formula=c in (2, 6, 7, 8, 9, 10, 11, 12))
        style_link(ws_nav.cell(nav_row, 2))
        ws_nav.row_dimensions[nav_row].height = max(
            16, 14 * (po.count('\n') + 1), 14 * (content.count('\n') + 1))
        if cust != prev_cust:
            if prev_cust is not None and cust_start < nav_row - 1:
                ws_nav.merge_cells(start_row=cust_start, start_column=1,
                                   end_row=nav_row - 1, end_column=1)
            cust_start = nav_row
            prev_cust = cust
        nav_row += 1
    if prev_cust is not None and cust_start < nav_row - 1:
        ws_nav.merge_cells(start_row=cust_start, start_column=1,
                           end_row=nav_row - 1, end_column=1)

    # 预警列（L）条件格式：底色+字色，Mac/Win 一致（替代 emoji 红绿黄灯）
    if nav_row > 3:
        warn_range = f"L3:L{nav_row - 1}"
        ws_nav.conditional_formatting.add(warn_range, FormulaRule(
            formula=['ISNUMBER(SEARCH("逾期",L3))'], fill=FILL_WARN_RED, font=F_WARN_RED))
        ws_nav.conditional_formatting.add(warn_range, FormulaRule(
            formula=['ISNUMBER(SEARCH("部分未回",L3))'], fill=FILL_WARN_YEL))
        ws_nav.conditional_formatting.add(warn_range, FormulaRule(
            formula=['ISNUMBER(SEARCH("已回款",L3))'], fill=FILL_WARN_GRN, font=F_WARN_GRN))

    wb.save(out)
    return out


def verify(path, rules=None):
    """生成后校验：行数、批次顺序、各金额列合计 vs 飞书、发票/回款覆盖。"""
    rules = {**DEFAULT_RULES, **(rules or {})}
    tol = rules['金额容差']
    base = {}
    for name, paths, key in [
        ('terms', [EXPORT_DIR + '/terms.json'], '款类'),
        ('invoices', [EXPORT_DIR + '/invoices.json'], '款类'),
        ('payments', [EXPORT_DIR + '/payments.json'], '款类'),
        ('shipments', [EXPORT_DIR + '/shipments.json'], '发货批次'),
    ]:
        recs = load_table(paths)
        d = {}
        for x in recs:
            d.setdefault(s(x['JOB No']), []).append(x)
        base[name] = d
    devs = []
    for p in sorted(os.listdir(EXPORT_DIR + '/devices')):
        devs += load_table([EXPORT_DIR + '/devices/' + p])
    dev_by = {}
    for x in devs:
        dev_by.setdefault(s(x['JOB No']), []).append(x)
    contracts = load_table([EXPORT_DIR + '/contracts.json'])
    terms_all = load_table([EXPORT_DIR + '/terms.json'])
    shipments_all = load_table([EXPORT_DIR + '/shipments.json'])
    invoices_all = load_table([EXPORT_DIR + '/invoices.json'])
    payments_all = load_table([EXPORT_DIR + '/payments.json'])
    JOBS = sorted({s(c['JOB No']) for c in contracts})

    wb = load_workbook(path, data_only=False)
    problems = []
    for job in JOBS:
        ws = wb[job]
        heads = [ws.cell(5, c).value for c in range(1, ws.max_column + 1)]
        n = ws.max_row - 6  # 最后一行是“返回导航页”
        dev_n = len(dev_by.get(job, []))
        exp_n = dev_n if dev_n else (1 if (base['invoices'].get(job) or base['payments'].get(job)) else 0)
        if n != exp_n:
            problems.append((job, '行数', n, exp_n))
        batches = []
        for r in range(6, ws.max_row + 1):
            b = ws.cell(r, 1).value
            if b and (not batches or batches[-1] != b):
                batches.append(b)
        nat = sorted(batches, key=natural_batch_key)
        if batches != nat:
            problems.append((job, '批次顺序', batches, nat))
        # 金额列合计
        sums = {}
        for c, h in enumerate(heads, 1):
            if h and h.endswith('·开票金额'):
                sums[('inv', h.split('·')[0])] = sum(
                    float(ws.cell(r, c).value) for r in range(6, ws.max_row + 1)
                    if isinstance(ws.cell(r, c).value, (int, float)))
            elif h and h.endswith('·回款金额'):
                sums[('pay', h.split('·')[0])] = sum(
                    float(ws.cell(r, c).value) for r in range(6, ws.max_row + 1)
                    if isinstance(ws.cell(r, c).value, (int, float)))
        exp_inv = {}
        for x in base['invoices'].get(job, []):
            t = s(x['款类'])
            exp_inv[t] = exp_inv.get(t, 0) + float(x['含税金额'])
        exp_pay = {}
        for x in base['payments'].get(job, []):
            t = s(x['款类'])
            exp_pay[t] = exp_pay.get(t, 0) + float(x['含税金额'])
        for (kind, t), got in sums.items():
            exp = exp_inv.get(t, 0.0) if kind == 'inv' else exp_pay.get(t, 0.0)
            if exp is None or abs(got - exp) > 0.02:
                problems.append((job, f'{kind}:{t}金额', got, exp))
        # 发票/回款每条至少出现一次
        for x in base['invoices'].get(job, []):
            t = s(x['款类'])
            d = fmt_date(x['开票日'])
            amt = float(x['含税金额'])
            col = heads.index(t + '·开票金额') + 1 if (t + '·开票金额') in heads else None
            if col and not any(ws.cell(r, col).value == amt and fmt_date(ws.cell(r, col - 2).value) == d
                               for r in range(6, ws.max_row + 1)):
                problems.append((job, '发票未出现', t, d, amt))
        for x in base['payments'].get(job, []):
            t = s(x['款类'])
            d = fmt_date(x['回款日'])
            amt = float(x['含税金额'])
            col = heads.index(t + '·回款金额') + 1 if (t + '·回款金额') in heads else None
            if col and not any(ws.cell(r, col).value == amt and fmt_date(ws.cell(r, col - 1).value) == d
                               for r in range(6, ws.max_row + 1)):
                problems.append((job, '回款未出现', t, d, amt))
        # PO 头换行（无分号）
        po = ws.cell(3, 3).value or ''
        if ';' in po:
            problems.append((job, 'PO含分号', po))
        # 行预警公式引用的列必须都在隐藏辅助列内
        hidden_letters = set()
        for c in range(ws.max_column - 12, ws.max_column + 1):
            h = ws.cell(5, c).value
            if h and str(h).startswith('行'):
                hidden_letters.add(get_column_letter(c))
        wf = ws.cell(6, ws.max_column).value or ''
        for letter in re.findall(r'([A-Z]{1,3})\d+', wf):
            if letter not in hidden_letters:
                problems.append((job, 'warn公式引用列', letter, wf))
    # 导航页：F~L 必须为公式；PO 无分号
    nav = wb['导航页']
    for r in range(3, 3 + len(JOBS)):
        for c in range(6, 13):
            v = nav.cell(r, c).value
            if not (isinstance(v, str) and v.startswith('=')):
                problems.append(('nav', r, c, v))
        po = str(nav.cell(r, 3).value or '')
        if ';' in po:
            problems.append(('nav-po', r, po))
    if nav.max_row != 2 + len(JOBS):
        problems.append(('nav行数', nav.max_row, 2 + len(JOBS)))
    # 未回收管理表
    if wb.sheetnames[1] != '未回收管理表':
        problems.append(('sheet2', wb.sheetnames[1]))
    up = wb['未回收管理表']
    up_heads = [up.cell(2, c).value for c in range(1, 10)]
    if up_heads != ['客户', 'JOB No', '批次', '款类', '预警等级', '开票日期',
                    '预定回收日期', '未回收金额', '未回收原因']:
        problems.append(('未回收表头', up_heads))
    anchor_b = {}
    for mr in up.merged_cells.ranges:
        if mr.min_col == 2:
            for r in range(mr.min_row, mr.max_row + 1):
                anchor_b[r] = mr.min_row
    exp_rows = compute_unpaid_rows(contracts, terms_all, shipments_all,
                                   invoices_all, payments_all, devs, rules)
    if up.max_row - 3 != len(exp_rows):
        problems.append(('未回收行数', up.max_row - 3, len(exp_rows)))
    for i, er in enumerate(exp_rows):
        r = 3 + i
        if str(up.cell(r, 3).value) != er['批次'] or str(up.cell(r, 4).value) != er['款类']:
            problems.append(('未回收行列', er['JOB No'], r))
        if er['static']:
            amt = up.cell(r, 8).value
            if not isinstance(amt, (int, float)) or abs(amt - er['未回收金额']) > 0.01:
                problems.append(('未回收静态金额', er['JOB No'], r, amt, er['未回收金额']))
        else:
            h = str(up.cell(r, 8).value or '')
            if not h.startswith('=MAX(0,ROUND(') or 'SUMIFS' not in h:
                problems.append(('未回收金额公式', er['JOB No'], r, h))
            else:
                m = re.search(r'ROUND\(([\d.]+)-SUMIFS', h)
                if m and abs(float(m.group(1)) - er['recv']) > 0.01:
                    problems.append(('未回收应收常量', er['JOB No'], r, m.group(1), er['recv']))
                jws = wb[er['JOB No']]
                pay_letter = next((get_column_letter(c) for c in range(1, jws.max_column + 1)
                                   if jws.cell(5, c).value == f'行回款金额·{er["款类"]}'), None)
                if pay_letter and f'${pay_letter}$6' not in h:
                    problems.append(('未回收SUMIFS列', er['JOB No'], r, pay_letter))
                crit = f'$C{r}' if er['raw_batch'] else '""'
                if crit not in h:
                    problems.append(('未回收批次条件', er['JOB No'], r, crit))
            e = str(up.cell(r, 5).value or '')
            if 'TODAY()' not in e or f'$H{r}' not in e:
                problems.append(('未回收预警公式', er['JOB No'], r, e))
            rr = str(up.cell(r, 9).value or '')
            if not rr.startswith(f'=IF($H{r}<={tol:g}'):
                problems.append(('未回收原因公式', er['JOB No'], r, rr))
    last = 2 + len(exp_rows)
    if str(up.cell(up.max_row, 8).value or '') != f'=SUM(H3:H{last})':
        problems.append(('未回收合计公式', up.cell(up.max_row, 8).value))
    # 每页最后一行返回导航超链接
    for job in JOBS:
        ws = wb[job]
        last = ws.max_row
        v = ws.cell(last, 1).value
        if not (isinstance(v, str) and v.startswith('=HYPERLINK') and '返回导航页' in v):
            problems.append((job, '返回导航链接缺失', last, v))
    print('verify problems:', len(problems))
    for p in problems:
        print('  ', p)
    return problems


if __name__ == '__main__':
    out = build()
    verify(out)
