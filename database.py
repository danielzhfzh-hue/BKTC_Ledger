#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""SQLite persistence for 上海康肯销售订单管理系统."""

from __future__ import annotations

import hashlib
import json
import os
import re
import sqlite3
import uuid
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import core


DATABASE_VERSION = 5
WORKBOOK_VERSION = 2
META_SHEET = "_BKTC_META"
HELP_SHEET = "使用说明"
ID_HEADER = "_记录ID"
DATABASE_EXTENSIONS = {".db", ".sqlite", ".sqlite3"}


class DatabaseValidationError(ValueError):
    def __init__(self, issues):
        self.issues = issues
        errors = [x["msg"] for x in issues if x.get("severity") == "error"]
        detail = "；".join(errors[:5])
        if len(errors) > 5:
            detail += f"；另有 {len(errors) - 5} 项"
        super().__init__("数据校验失败，未写入数据库：" + (detail or "存在未知错误"))


class SpreadsheetFormatError(ValueError):
    pass


class StaleImportError(RuntimeError):
    pass


class QuotationValidationError(ValueError):
    pass


@dataclass
class PreparedImport:
    database_path: str
    workbook_path: str
    database_id: str
    database_hash: str
    base_revision: int
    current_revision: int
    data: dict[str, list[dict[str, Any]]]
    rules: dict[str, Any]
    changes: list[dict[str, Any]]
    action_counts: dict[str, int]
    table_counts: dict[str, dict[str, int]]
    issues: list[dict[str, str]]
    file_signature: tuple[int, int]

    @property
    def stale(self):
        return self.base_revision != self.current_revision


def is_database_path(path):
    return Path(os.fspath(path)).suffix.lower() in DATABASE_EXTENSIONS


def database_path_for(path):
    """Return a sibling .db path for a legacy JSON path."""
    p = Path(os.fspath(path))
    if is_database_path(p):
        return str(p)
    name = p.name
    if name.endswith(".records.json"):
        name = name[:-len(".records.json")] + ".db"
    elif p.suffix.lower() == ".json":
        name = p.stem + ".db"
    else:
        name += ".db"
    return str(p.with_name(name))


def _q(name):
    return '"' + str(name).replace('"', '""') + '"'


def _sql_type(field_type):
    if field_type in (core.NUMBER,):
        return "REAL"
    if field_type in (core.INT,):
        return "INTEGER"
    return "TEXT"


def _connect(path):
    path = os.path.abspath(os.fspath(path))
    os.makedirs(os.path.dirname(path), exist_ok=True)
    conn = sqlite3.connect(path, timeout=15)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA journal_mode = WAL")
    conn.execute("PRAGMA synchronous = FULL")
    conn.execute("PRAGMA busy_timeout = 15000")
    _create_schema(conn)
    return conn


def _create_schema(conn):
    audit_was_present = conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name='_audit_event'"
    ).fetchone() is not None
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS "_meta" (
            "key" TEXT PRIMARY KEY,
            "value" TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS "_warning_rules" (
            "key" TEXT PRIMARY KEY,
            "value_json" TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS "_change_log" (
            "revision" INTEGER PRIMARY KEY,
            "created_at" TEXT NOT NULL,
            "reason" TEXT NOT NULL,
            "summary_json" TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS "_audit_event" (
            "event_id" TEXT PRIMARY KEY,
            "revision" INTEGER NOT NULL UNIQUE,
            "created_at" TEXT NOT NULL,
            "operator_name" TEXT NOT NULL,
            "source" TEXT NOT NULL,
            "actions_json" TEXT NOT NULL,
            "app_version" TEXT NOT NULL,
            "platform" TEXT NOT NULL,
            "hostname" TEXT NOT NULL,
            "change_count" INTEGER NOT NULL,
            "summary_json" TEXT NOT NULL,
            "previous_hash" TEXT NOT NULL,
            "event_hash" TEXT NOT NULL UNIQUE
        );
        CREATE TABLE IF NOT EXISTS "_audit_change" (
            "change_id" INTEGER PRIMARY KEY AUTOINCREMENT,
            "event_id" TEXT NOT NULL,
            "table_name" TEXT NOT NULL,
            "record_id" TEXT NOT NULL,
            "job_no" TEXT NOT NULL,
            "operation" TEXT NOT NULL,
            "field_name" TEXT NOT NULL,
            "old_value_json" TEXT,
            "new_value_json" TEXT,
            "origin" TEXT NOT NULL,
            FOREIGN KEY ("event_id") REFERENCES "_audit_event" ("event_id")
                ON DELETE RESTRICT
        );
        CREATE TABLE IF NOT EXISTS "quotation" (
            "quote_id" TEXT PRIMARY KEY,
            "quote_no" TEXT NOT NULL UNIQUE,
            "revision_label" TEXT NOT NULL DEFAULT '1.0',
            "quote_date" TEXT NOT NULL,
            "issuer_name" TEXT NOT NULL DEFAULT '',
            "issuer_address" TEXT NOT NULL DEFAULT '',
            "issuer_contact" TEXT NOT NULL DEFAULT '',
            "issuer_key" TEXT NOT NULL DEFAULT '',
            "customer" TEXT NOT NULL,
            "customer_address" TEXT NOT NULL DEFAULT '',
            "contact" TEXT NOT NULL DEFAULT '',
            "salesperson" TEXT NOT NULL DEFAULT '',
            "source_job_no" TEXT NOT NULL DEFAULT '',
            "currency" TEXT NOT NULL DEFAULT 'RMB',
            "tax_rate" REAL NOT NULL DEFAULT 13,
            "valid_until" TEXT NOT NULL DEFAULT '',
            "subject" TEXT NOT NULL DEFAULT '',
            "warranty" TEXT NOT NULL DEFAULT '',
            "incoterm" TEXT NOT NULL DEFAULT '',
            "ship_to" TEXT NOT NULL DEFAULT '',
            "description" TEXT NOT NULL DEFAULT '',
            "payment_terms" TEXT NOT NULL DEFAULT '',
            "payment_language" TEXT NOT NULL DEFAULT 'zh',
            "delivery_terms" TEXT NOT NULL DEFAULT '',
            "status" TEXT NOT NULL DEFAULT '草稿',
            "notes" TEXT NOT NULL DEFAULT '',
            "created_at" TEXT NOT NULL,
            "updated_at" TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS "quotation_item" (
            "item_id" TEXT PRIMARY KEY,
            "quote_id" TEXT NOT NULL,
            "line_no" INTEGER NOT NULL,
            "model" TEXT NOT NULL,
            "description" TEXT NOT NULL DEFAULT '',
            "quantity" REAL NOT NULL,
            "unit" TEXT NOT NULL DEFAULT '台',
            "unit_price" REAL NOT NULL,
            "amount" REAL NOT NULL,
            "remark" TEXT NOT NULL DEFAULT '',
            FOREIGN KEY ("quote_id") REFERENCES "quotation" ("quote_id")
                ON UPDATE CASCADE ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS "quotation_payment_item" (
            "payment_item_id" TEXT PRIMARY KEY,
            "quote_id" TEXT NOT NULL,
            "line_no" INTEGER NOT NULL,
            "kind" TEXT NOT NULL DEFAULT '',
            "ratio" REAL NOT NULL DEFAULT 0,
            "days" INTEGER NOT NULL DEFAULT 0,
            "trigger" TEXT NOT NULL DEFAULT '',
            "description_zh" TEXT NOT NULL DEFAULT '',
            "description_en" TEXT NOT NULL DEFAULT '',
            FOREIGN KEY ("quote_id") REFERENCES "quotation" ("quote_id")
                ON UPDATE CASCADE ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS "quotation_option" (
            "option_id" TEXT PRIMARY KEY,
            "kind" TEXT NOT NULL,
            "value" TEXT NOT NULL,
            "language" TEXT NOT NULL DEFAULT 'zh',
            "updated_at" TEXT NOT NULL,
            UNIQUE ("kind", "value", "language")
        );
        CREATE INDEX IF NOT EXISTS "idx_audit_event_created"
            ON "_audit_event" ("created_at");
        CREATE INDEX IF NOT EXISTS "idx_audit_change_event"
            ON "_audit_change" ("event_id", "change_id");
        CREATE INDEX IF NOT EXISTS "idx_audit_change_job"
            ON "_audit_change" ("job_no");
        CREATE INDEX IF NOT EXISTS "idx_quotation_customer"
            ON "quotation" ("customer", "quote_date");
        CREATE INDEX IF NOT EXISTS "idx_quotation_item_quote"
            ON "quotation_item" ("quote_id", "line_no");
        CREATE INDEX IF NOT EXISTS "idx_quotation_item_model"
            ON "quotation_item" ("model");
        CREATE INDEX IF NOT EXISTS "idx_quotation_payment_quote"
            ON "quotation_payment_item" ("quote_id", "line_no");
        CREATE INDEX IF NOT EXISTS "idx_quotation_option_recent"
            ON "quotation_option" ("kind", "language", "updated_at");
        CREATE TRIGGER IF NOT EXISTS "audit_event_no_update"
            BEFORE UPDATE ON "_audit_event"
            BEGIN SELECT RAISE(ABORT, 'audit events are append-only'); END;
        CREATE TRIGGER IF NOT EXISTS "audit_event_no_delete"
            BEFORE DELETE ON "_audit_event"
            BEGIN SELECT RAISE(ABORT, 'audit events are append-only'); END;
        CREATE TRIGGER IF NOT EXISTS "audit_change_no_update"
            BEFORE UPDATE ON "_audit_change"
            BEGIN SELECT RAISE(ABORT, 'audit changes are append-only'); END;
        CREATE TRIGGER IF NOT EXISTS "audit_change_no_delete"
            BEFORE DELETE ON "_audit_change"
            BEGIN SELECT RAISE(ABORT, 'audit changes are append-only'); END;
        """
    )
    # v4 databases already contain quotation, so add v5 fields explicitly.
    quotation_columns = {
        row[1] for row in conn.execute('PRAGMA table_info("quotation")')
    }
    for field, definition in (
        ("issuer_key", "TEXT NOT NULL DEFAULT ''"),
        ("payment_language", "TEXT NOT NULL DEFAULT 'zh'"),
    ):
        if field not in quotation_columns:
            conn.execute(f'ALTER TABLE "quotation" ADD COLUMN "{field}" {definition}')
    for table in core.TABLES:
        columns = [f'{_q("记录ID")} TEXT PRIMARY KEY']
        for field, field_type, *_ in core.SCHEMA[table]:
            constraints = []
            if field == "JOB No":
                constraints.append("NOT NULL")
                if table == "合同订单":
                    constraints.append("UNIQUE")
            if table == "发货批次" and field == "发货批次":
                constraints.append("NOT NULL")
            columns.append(
                f"{_q(field)} {_sql_type(field_type)} {' '.join(constraints)}".rstrip()
            )
        columns.extend([
            f'{_q("_extra_json")} TEXT NOT NULL DEFAULT "{{}}"',
            f'{_q("_row_order")} INTEGER NOT NULL',
        ])
        if table != "合同订单":
            columns.append(
                f'FOREIGN KEY ({_q("JOB No")}) REFERENCES {_q("合同订单")} '
                f'({_q("JOB No")}) ON UPDATE CASCADE ON DELETE CASCADE '
                "DEFERRABLE INITIALLY DEFERRED"
            )
        if table == "发货批次":
            columns.append(f'UNIQUE ({_q("JOB No")}, {_q("发货批次")})')
        if table == "设备台账":
            columns.append(
                f'FOREIGN KEY ({_q("JOB No")}, {_q("发货批次")}) '
                f'REFERENCES {_q("发货批次")} ({_q("JOB No")}, {_q("发货批次")}) '
                "ON UPDATE CASCADE ON DELETE RESTRICT DEFERRABLE INITIALLY DEFERRED"
            )
        conn.execute(f'CREATE TABLE IF NOT EXISTS {_q(table)} ({", ".join(columns)})')
        existing_columns = {
            row[1] for row in conn.execute(f'PRAGMA table_info({_q(table)})')
        }
        for field, field_type, *_ in core.SCHEMA[table]:
            if field not in existing_columns:
                conn.execute(
                    f'ALTER TABLE {_q(table)} ADD COLUMN {_q(field)} {_sql_type(field_type)}'
                )
        if table != "合同订单":
            conn.execute(
                f'CREATE INDEX IF NOT EXISTS {_q("idx_" + table + "_job")} '
                f'ON {_q(table)} ({_q("JOB No")})'
            )
    conn.execute(
        f'CREATE INDEX IF NOT EXISTS {_q("idx_设备台账_batch")} '
        f'ON {_q("设备台账")} ({_q("JOB No")}, {_q("发货批次")})'
    )
    conn.execute(
        "INSERT OR IGNORE INTO _meta(key, value) VALUES('database_id', ?)",
        (str(uuid.uuid4()),),
    )
    conn.execute("INSERT OR IGNORE INTO _meta(key, value) VALUES('revision', '0')")
    if not audit_was_present:
        now = datetime.now().isoformat(timespec="seconds")
        revision = _meta(conn, "revision", "0") or "0"
        conn.execute(
            "INSERT OR IGNORE INTO _meta(key, value) VALUES('audit_started_at', ?)",
            (now,),
        )
        conn.execute(
            "INSERT OR IGNORE INTO _meta(key, value) VALUES('audit_start_revision', ?)",
            (revision,),
        )
        conn.execute(
            "INSERT OR IGNORE INTO _meta(key, value) VALUES('audit_chain_head', '')"
        )
    conn.execute(
        "INSERT INTO _meta(key, value) VALUES('schema_version', ?) "
        "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (str(DATABASE_VERSION),),
    )
    conn.execute(f"PRAGMA user_version = {DATABASE_VERSION}")
    conn.commit()


def _meta(conn, key, default=""):
    row = conn.execute("SELECT value FROM _meta WHERE key = ?", (key,)).fetchone()
    return row[0] if row else default


def _set_meta(conn, key, value):
    conn.execute(
        "INSERT INTO _meta(key, value) VALUES(?, ?) "
        "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (key, str(value)),
    )


def _new_record_id():
    return "local-" + uuid.uuid4().hex


def _prepare_data(data):
    prepared = core.derive(data)
    for table in core.TABLES:
        seen = set()
        for rec in prepared[table]:
            rid = str(rec.get("记录ID") or "").strip()
            if not rid or rid in seen:
                rid = _new_record_id()
                rec["记录ID"] = rid
            seen.add(rid)
    issues = core.validate(prepared)
    if any(x.get("severity") == "error" for x in issues):
        raise DatabaseValidationError(issues)
    return prepared, issues


def _backup_database(path, keep=10):
    path = os.path.abspath(os.fspath(path))
    if not os.path.exists(path) or os.path.getsize(path) == 0:
        return None
    backup_dir = os.path.join(os.path.dirname(path), "备份")
    os.makedirs(backup_dir, exist_ok=True)
    stem = os.path.splitext(os.path.basename(path))[0]
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")[:-3]
    target = os.path.join(backup_dir, f"{stem}_database_{stamp}.db")
    source = _connect(path)
    destination = sqlite3.connect(target)
    try:
        source.backup(destination)
    finally:
        destination.close()
        source.close()
    backups = sorted(
        (os.path.join(backup_dir, f) for f in os.listdir(backup_dir)
         if f.startswith(stem + "_database_") and f.endswith(".db")),
        key=os.path.getmtime,
    )
    while len(backups) > keep:
        os.remove(backups.pop(0))
    return target


def _insert_rows(conn, table, records):
    fields = [f[0] for f in core.SCHEMA[table]]
    known = set(fields) | {"记录ID"}
    columns = ["记录ID", *fields, "_extra_json", "_row_order"]
    sql = (
        f'INSERT INTO {_q(table)} ({", ".join(_q(x) for x in columns)}) '
        f'VALUES ({", ".join("?" for _ in columns)})'
    )
    for order, rec in enumerate(records):
        extra = {k: v for k, v in rec.items() if k not in known}
        values = [rec["记录ID"]]
        for field in fields:
            value = rec.get(field)
            if table == "设备台账" and field == "发货批次" and value == "":
                value = None
            values.append(value)
        values.extend([json.dumps(extra, ensure_ascii=False, separators=(",", ":")), order])
        conn.execute(sql, values)


def _rules_from_connection(conn):
    rules = {}
    for row in conn.execute("SELECT key, value_json FROM _warning_rules"):
        try:
            rules[row[0]] = json.loads(row[1])
        except (TypeError, ValueError):
            continue
    return {**core.DEFAULT_RULES, **rules}


def _data_from_connection(conn):
    data = core.empty_data()
    for table in core.TABLES:
        fields = [f[0] for f in core.SCHEMA[table]]
        select = ["记录ID", *fields, "_extra_json"]
        rows = conn.execute(
            f'SELECT {", ".join(_q(x) for x in select)} FROM {_q(table)} '
            f'ORDER BY {_q("_row_order")}'
        )
        for row in rows:
            try:
                rec = json.loads(row["_extra_json"] or "{}")
            except (TypeError, ValueError):
                rec = {}
            rec["记录ID"] = row["记录ID"]
            for field in fields:
                rec[field] = row[field]
            data[table].append(rec)
    return core.normalize(data)


def _json_scalar(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)


def _audit_snapshot(table, record):
    derived = set(core.DERIVED.get(table, []))
    return {
        key: value for key, value in record.items()
        if key != "记录ID" and key not in derived
    }


def _job_rename_map(old_data, new_data):
    old_rows = {x.get("记录ID"): x for x in old_data.get("合同订单", [])}
    result = {}
    for rec in new_data.get("合同订单", []):
        old = old_rows.get(rec.get("记录ID"))
        before, after = (old or {}).get("JOB No"), rec.get("JOB No")
        if old and before != after:
            result[before] = after
    return result


def _audit_changes(old_data, new_data, old_rules, new_rules, requested_data=None):
    changes = []
    job_renames = _job_rename_map(old_data, new_data)
    requested_data = requested_data or new_data
    for table in core.TABLES:
        old_rows = {x.get("记录ID"): x for x in old_data.get(table, [])}
        new_rows = {x.get("记录ID"): x for x in new_data.get(table, [])}
        requested_rows = {x.get("记录ID"): x for x in requested_data.get(table, [])}
        for record_id in sorted(set(old_rows) | set(new_rows), key=str):
            old = old_rows.get(record_id)
            new = new_rows.get(record_id)
            if old is None:
                snapshot = _audit_snapshot(table, new)
                changes.append({
                    "table_name": table, "record_id": str(record_id or ""),
                    "job_no": str(new.get("JOB No") or ""), "operation": "insert",
                    "field_name": "", "old_value": None, "new_value": snapshot,
                    "origin": "user",
                })
                continue
            if new is None:
                snapshot = _audit_snapshot(table, old)
                changes.append({
                    "table_name": table, "record_id": str(record_id or ""),
                    "job_no": str(old.get("JOB No") or ""), "operation": "delete",
                    "field_name": "", "old_value": snapshot, "new_value": None,
                    "origin": "user",
                })
                continue
            derived = set(core.DERIVED.get(table, []))
            fields = sorted((set(old) | set(new)) - {"记录ID"} - derived)
            for field in fields:
                before, after = old.get(field), new.get(field)
                if before == after:
                    continue
                origin = "user"
                if (table != "合同订单" and field == "JOB No"
                        and job_renames.get(before) == after):
                    origin = "association_sync"
                elif (record_id in requested_rows
                      and requested_rows[record_id].get(field) == before
                      and requested_rows[record_id].get(field) != after):
                    origin = "system_recompute"
                changes.append({
                    "table_name": table, "record_id": str(record_id or ""),
                    "job_no": str(new.get("JOB No") or old.get("JOB No") or ""),
                    "operation": "update", "field_name": field,
                    "old_value": before, "new_value": after, "origin": origin,
                })
    old_rules = old_rules or {}
    new_rules = new_rules or {}
    for key in sorted(set(old_rules) | set(new_rules)):
        before, after = old_rules.get(key), new_rules.get(key)
        if before != after:
            changes.append({
                "table_name": "预警规则", "record_id": "warning_rules", "job_no": "",
                "operation": "update", "field_name": key,
                "old_value": before, "new_value": after, "origin": "user",
            })
    return changes


def _audit_summary(changes):
    summary = {}
    for change in changes:
        table = change["table_name"]
        operation = change["operation"]
        table_summary = summary.setdefault(table, {"insert": 0, "update": 0, "delete": 0})
        table_summary[operation] += 1
    return summary


def _append_audit_event(conn, revision, created_at, reason, changes, context=None):
    context = context if isinstance(context, dict) else {}
    event_id = "audit-" + uuid.uuid4().hex
    previous_hash = _meta(conn, "audit_chain_head", "")
    actions = context.get("actions") if isinstance(context.get("actions"), list) else []
    event = {
        "event_id": event_id,
        "revision": revision,
        "created_at": created_at,
        "operator_name": str(context.get("operator_name") or "unknown"),
        "source": str(context.get("source") or reason or "manual_save"),
        "actions": [str(x) for x in actions if str(x).strip()],
        "app_version": str(context.get("app_version") or ""),
        "platform": str(context.get("platform") or ""),
        "hostname": str(context.get("hostname") or ""),
        "summary": _audit_summary(changes),
    }
    payload = {
        "event": event,
        "changes": changes,
    }
    canonical = json.dumps(payload, ensure_ascii=False, sort_keys=True,
                           separators=(",", ":"), default=str)
    event_hash = hashlib.sha256((previous_hash + "\n" + canonical).encode("utf-8")).hexdigest()
    conn.execute(
        "INSERT INTO _audit_event(event_id, revision, created_at, operator_name, source, "
        "actions_json, app_version, platform, hostname, change_count, summary_json, "
        "previous_hash, event_hash) VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        (
            event_id, revision, created_at, event["operator_name"], event["source"],
            json.dumps(event["actions"], ensure_ascii=False), event["app_version"],
            event["platform"], event["hostname"], len(changes),
            json.dumps(event["summary"], ensure_ascii=False), previous_hash, event_hash,
        ),
    )
    for change in changes:
        conn.execute(
            "INSERT INTO _audit_change(event_id, table_name, record_id, job_no, operation, "
            "field_name, old_value_json, new_value_json, origin) VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (
                event_id, change["table_name"], change["record_id"], change["job_no"],
                change["operation"], change["field_name"],
                _json_scalar(change["old_value"]), _json_scalar(change["new_value"]),
                change["origin"],
            ),
        )
    _set_meta(conn, "audit_chain_head", event_hash)
    return event_id, event_hash


def save_database(path, data, rules=None, *, reason="manual_save", backup=True,
                  expected_revision=None, audit_context=None):
    path = os.path.abspath(os.fspath(path))
    requested_data = core.normalize(data)
    prepared, issues = _prepare_data(data)
    if rules is None:
        rules = get_rules(path) if os.path.exists(path) else dict(core.DEFAULT_RULES)
    early_revision = get_revision(path) if os.path.exists(path) else 0
    if expected_revision is not None and early_revision != expected_revision:
        raise StaleImportError(
            f"数据库已从修订 {expected_revision} 更新到 {early_revision}，请重新加载后再操作"
        )
    backup_path = _backup_database(path) if backup else None
    conn = _connect(path)
    try:
        conn.execute("BEGIN IMMEDIATE")
        current_revision = int(_meta(conn, "revision", "0") or 0)
        if expected_revision is not None and current_revision != expected_revision:
            raise StaleImportError(
                f"数据库已从修订 {expected_revision} 更新到 {current_revision}，请重新加载后再操作"
            )
        old_data = _data_from_connection(conn)
        old_rules = _rules_from_connection(conn)
        changes = _audit_changes(
            old_data, prepared, old_rules, rules, requested_data=requested_data
        )
        for table in ("回款记录", "开票记录", "设备台账", "付款条件", "发货批次", "合同订单"):
            conn.execute(f'DELETE FROM {_q(table)}')
        for table in ("合同订单", "发货批次", "付款条件", "设备台账", "开票记录", "回款记录"):
            _insert_rows(conn, table, prepared[table])
        conn.execute("DELETE FROM _warning_rules")
        for key, value in (rules or {}).items():
            conn.execute(
                "INSERT INTO _warning_rules(key, value_json) VALUES(?, ?)",
                (key, json.dumps(value, ensure_ascii=False)),
            )
        foreign_key_issues = conn.execute("PRAGMA foreign_key_check").fetchall()
        if foreign_key_issues:
            raise sqlite3.IntegrityError(f"外键检查失败：{foreign_key_issues[:3]}")
        revision = current_revision + 1
        now = datetime.now().isoformat(timespec="seconds")
        _set_meta(conn, "revision", revision)
        _set_meta(conn, "last_saved_at", now)
        summary = {table: len(prepared[table]) for table in core.TABLES}
        conn.execute(
            "INSERT INTO _change_log(revision, created_at, reason, summary_json) "
            "VALUES(?, ?, ?, ?)",
            (revision, now, reason, json.dumps(summary, ensure_ascii=False)),
        )
        audit_event_id = None
        suppressed = (
            reason in {"database_created", "json_migration", "portable_starter"}
            and conn.execute("SELECT COUNT(*) FROM _audit_event").fetchone()[0] == 0
        )
        if changes and not suppressed:
            audit_event_id, _ = _append_audit_event(
                conn, revision, now, reason, changes, audit_context
            )
        elif suppressed:
            _set_meta(conn, "audit_started_at", now)
            _set_meta(conn, "audit_start_revision", revision)
            _set_meta(conn, "audit_chain_head", "")
        conn.commit()
    except sqlite3.IntegrityError as exc:
        conn.rollback()
        raise DatabaseValidationError([
            {"severity": "error", "msg": f"关系约束检查失败：{exc}"}
        ]) from exc
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return {
        "ok": True,
        "path": path,
        "revision": revision,
        "backup": backup_path,
        "audit_event_id": audit_event_id,
        "audit_change_count": len(changes) if audit_event_id else 0,
        "data": load_database(path),
        "issues": issues,
    }


def load_database(path):
    path = os.path.abspath(os.fspath(path))
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    conn = _connect(path)
    try:
        integrity = conn.execute("PRAGMA integrity_check").fetchone()[0]
        if integrity != "ok":
            raise RuntimeError(f"数据库完整性检查失败：{integrity}")
        return _data_from_connection(conn)
    finally:
        conn.close()


def get_rules(path):
    if not os.path.exists(path):
        return dict(core.DEFAULT_RULES)
    conn = _connect(path)
    try:
        return _rules_from_connection(conn)
    finally:
        conn.close()


def get_revision(path):
    if not os.path.exists(path):
        return 0
    conn = _connect(path)
    try:
        return int(_meta(conn, "revision", "0") or 0)
    finally:
        conn.close()


def get_database_info(path):
    conn = _connect(path)
    try:
        audit_count = conn.execute("SELECT COUNT(*) FROM _audit_event").fetchone()[0]
        return {
            "path": os.path.abspath(os.fspath(path)),
            "database_id": _meta(conn, "database_id"),
            "revision": int(_meta(conn, "revision", "0") or 0),
            "last_saved_at": _meta(conn, "last_saved_at"),
            "integrity": conn.execute("PRAGMA integrity_check").fetchone()[0],
            "audit_started_at": _meta(conn, "audit_started_at"),
            "audit_start_revision": int(_meta(conn, "audit_start_revision", "0") or 0),
            "audit_event_count": audit_count,
        }
    finally:
        conn.close()


QUOTATION_FIELDS = (
    "quote_no", "revision_label", "quote_date", "issuer_name", "issuer_address",
    "issuer_contact", "issuer_key", "customer", "customer_address", "contact", "salesperson",
    "source_job_no", "currency", "tax_rate", "valid_until", "subject",
    "warranty", "incoterm", "ship_to", "description", "payment_terms", "payment_language",
    "delivery_terms", "status", "notes",
)
QUOTATION_ITEM_FIELDS = (
    "line_no", "model", "description", "quantity", "unit", "unit_price",
    "amount", "remark",
)
QUOTATION_PAYMENT_FIELDS = (
    "line_no", "kind", "ratio", "days", "trigger", "description_zh", "description_en",
)
ISSUER_PRESETS = {
    "beijing": {
        "key": "beijing",
        "label": "北京康肯环境保护设备有限公司",
        "name": "Beijing Kanken Environmental Protection Equipment Co.,Ltd",
        "address": "Rm No.A-1501, No.8,building No.129,eight Li Zhuang,Chaoyang District,Beijing,10025,China",
        "contact": "Tel : 021-61253711   Fax:021-61253710",
    },
    "japan": {
        "key": "japan",
        "label": "KANKEN TECHNO（日本）",
        "name": "KANKEN TECHNO CO.,LTD",
        "address": "30-2 Ota Kotari, Nagokakyo,Kyoto,Japan",
        "contact": "Tel : +81-75-955-8826   Fax:+81-75-955-8915",
    },
}
DEFAULT_WARRANTY = "设备验收后12个月或设备到货后15个月"
DEFAULT_INCOTERMS = ("CIP", "CIF", "DAP", "FOB", "FCA", "DDP")
PAYMENT_KIND_EN = {
    "预付款": "Advance payment", "发货款": "Shipment payment", "到货款": "Arrival payment",
    "验收款": "Acceptance payment", "质保款": "Warranty payment", "全额": "Full payment",
}
PAYMENT_TRIGGER_EN = {
    "开票后": "after invoicing", "货到签收后": "after delivery and receipt",
    "验收合格后": "after acceptance", "验收后": "after acceptance",
    "质保期满后": "after the warranty period expires", "合同生效后": "after the contract takes effect",
    "月结次月月底": "at the end of the month following the monthly settlement month",
}
QUOTATION_FIELD_LABELS = {
    "quote_no": "报价单号", "revision_label": "版本", "quote_date": "报价日期",
    "issuer_name": "报价方", "issuer_address": "报价方地址",
    "issuer_contact": "报价方联系方式", "issuer_key": "报价主体", "customer": "客户",
    "customer_address": "客户地址", "contact": "客户联系人",
    "salesperson": "担当者", "source_job_no": "来源 JOB", "currency": "币种",
    "tax_rate": "税率%", "valid_until": "有效期至", "subject": "报价主题",
    "warranty": "质保条款", "incoterm": "贸易条款", "ship_to": "交货地点",
    "description": "说明", "payment_terms": "付款条件", "payment_language": "付款条件语言",
    "delivery_terms": "交期", "status": "状态", "notes": "备注",
}
QUOTATION_ITEM_LABELS = {
    "line_no": "序号", "model": "型号/项目", "description": "项目说明",
    "quantity": "数量", "unit": "单位", "unit_price": "未税单价",
    "amount": "未税金额", "remark": "明细备注",
}
QUOTATION_PAYMENT_LABELS = {
    "line_no": "序号", "kind": "款类", "ratio": "比例%", "days": "账期天数",
    "trigger": "触发条件", "description_zh": "中文说明", "description_en": "英文说明",
}


def _quote_text(value):
    return "" if value is None else str(value).strip()


def _quote_decimal(value, label, *, minimum=None):
    try:
        result = Decimal(str(value if value not in (None, "") else 0))
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise QuotationValidationError(f"{label}必须是数字") from exc
    if not result.is_finite():
        raise QuotationValidationError(f"{label}必须是有限数字")
    if minimum is not None and result < Decimal(str(minimum)):
        raise QuotationValidationError(f"{label}不能小于 {minimum}")
    return result


def _validate_quote_date(value, label, required=False):
    value = _quote_text(value)
    if not value and not required:
        return ""
    if not value:
        raise QuotationValidationError(f"{label}不能为空")
    try:
        date.fromisoformat(value)
    except ValueError as exc:
        raise QuotationValidationError(f"{label}必须为 YYYY-MM-DD") from exc
    return value


def _payment_description_en(kind, ratio, days, trigger):
    kind_text = PAYMENT_KIND_EN.get(kind, kind or "Payment")
    trigger_text = PAYMENT_TRIGGER_EN.get(trigger, trigger or "as agreed")
    ratio_text = f"{ratio:g}%" if ratio is not None else ""
    days = int(days or 0)
    due_text = f"within {days} days " if days else ""
    return " ".join(x for x in (ratio_text, kind_text, due_text + trigger_text) if x).strip()


def _normalize_payment_items(raw_items):
    if not isinstance(raw_items, list):
        return []
    result = []
    seen = set()
    for index, raw in enumerate(raw_items, 1):
        raw = raw if isinstance(raw, dict) else {}
        kind = _quote_text(raw.get("kind") or raw.get("款类"))
        if not kind:
            continue
        ratio = _quote_decimal(raw.get("ratio", raw.get("比例%", 0)), f"第 {index} 条付款比例", minimum=0)
        days = _quote_decimal(raw.get("days", raw.get("账期天数", 0)), f"第 {index} 条账期天数", minimum=0)
        if days != days.to_integral_value():
            raise QuotationValidationError(f"第 {index} 条账期天数必须是整数")
        trigger = _quote_text(raw.get("trigger") or raw.get("触发条件"))
        desc_zh = _quote_text(raw.get("description_zh") or raw.get("desc") or raw.get("说明"))
        desc_en = _quote_text(raw.get("description_en")) or _payment_description_en(kind, float(ratio), int(days), trigger)
        payment_id = _quote_text(raw.get("payment_item_id"))
        if not payment_id or payment_id in seen:
            payment_id = "quote-payment-" + uuid.uuid4().hex
        seen.add(payment_id)
        result.append({
            "payment_item_id": payment_id, "line_no": index, "kind": kind,
            "ratio": float(ratio), "days": int(days), "trigger": trigger,
            "description_zh": desc_zh, "description_en": desc_en,
        })
    return result


def _suggest_quotation_number_from_connection(conn, quote_date=None):
    quote_date = _validate_quote_date(
        quote_date or date.today().isoformat(), "报价日期", required=True
    )
    day = date.fromisoformat(quote_date)
    base = f"{day:%y}SHMT{day:%m%d}"
    existing = {
        row[0] for row in conn.execute(
            'SELECT "quote_no" FROM "quotation" WHERE "quote_no" LIKE ?',
            (base + "%",),
        )
    }
    if base not in existing:
        return base
    sequence = 2
    while f"{base}-{sequence:02d}" in existing:
        sequence += 1
    return f"{base}-{sequence:02d}"


def suggest_quotation_number(path, quote_date=None):
    conn = _connect(path)
    try:
        return _suggest_quotation_number_from_connection(conn, quote_date)
    finally:
        conn.close()


def _normalize_quotation(conn, quotation):
    quotation = quotation if isinstance(quotation, dict) else {}
    is_new = not _quote_text(quotation.get("quote_id"))
    quote_id = _quote_text(quotation.get("quote_id")) or "quote-" + uuid.uuid4().hex
    quote_date = _validate_quote_date(
        quotation.get("quote_date") or date.today().isoformat(),
        "报价日期", required=True,
    )
    customer = _quote_text(quotation.get("customer"))
    if not customer:
        raise QuotationValidationError("客户不能为空")
    valid_until = _validate_quote_date(quotation.get("valid_until"), "有效期")
    tax_rate = _quote_decimal(quotation.get("tax_rate", 13), "税率", minimum=0)
    if tax_rate > 100:
        raise QuotationValidationError("税率不能大于 100")
    raw_items = quotation.get("items")
    if not isinstance(raw_items, list) or not raw_items:
        raise QuotationValidationError("至少需要一条报价明细")
    items = []
    seen_ids = set()
    for index, raw in enumerate(raw_items, 1):
        raw = raw if isinstance(raw, dict) else {}
        model = _quote_text(raw.get("model"))
        if not model:
            raise QuotationValidationError(f"第 {index} 条明细缺少型号/项目")
        quantity = _quote_decimal(raw.get("quantity"), f"第 {index} 条数量")
        if quantity <= 0:
            raise QuotationValidationError(f"第 {index} 条数量必须大于 0")
        unit_price = _quote_decimal(
            raw.get("unit_price"), f"第 {index} 条未税单价", minimum=0
        )
        item_id = _quote_text(raw.get("item_id"))
        if not item_id or item_id in seen_ids:
            item_id = "quote-item-" + uuid.uuid4().hex
        seen_ids.add(item_id)
        amount = (quantity * unit_price).quantize(Decimal("0.01"), ROUND_HALF_UP)
        items.append({
            "item_id": item_id,
            "line_no": index,
            "model": model,
            "description": _quote_text(raw.get("description")),
            "quantity": float(quantity),
            "unit": _quote_text(raw.get("unit")) or "台",
            "unit_price": float(unit_price),
            "amount": float(amount),
            "remark": _quote_text(raw.get("remark")),
        })
    subject = _quote_text(quotation.get("subject"))
    if not subject:
        subject = " / ".join(dict.fromkeys(item["model"] for item in items))
    issuer_key = _quote_text(quotation.get("issuer_key"))
    issuer = ISSUER_PRESETS.get(issuer_key)
    payment_items = _normalize_payment_items(quotation.get("payment_items"))
    payment_language = _quote_text(quotation.get("payment_language")) or "zh"
    if payment_language not in {"zh", "en"}:
        payment_language = "zh"
    result = {
        "quote_id": quote_id,
        "quote_no": _quote_text(quotation.get("quote_no"))
                    or _suggest_quotation_number_from_connection(conn, quote_date),
        "revision_label": _quote_text(quotation.get("revision_label")) or "1.0",
        "quote_date": quote_date,
        "issuer_name": _quote_text(quotation.get("issuer_name")) or (issuer or {}).get("name", ""),
        "issuer_address": _quote_text(quotation.get("issuer_address")) or (issuer or {}).get("address", ""),
        "issuer_contact": _quote_text(quotation.get("issuer_contact")) or (issuer or {}).get("contact", ""),
        "issuer_key": issuer_key,
        "customer": customer,
        "customer_address": _quote_text(quotation.get("customer_address")),
        "contact": _quote_text(quotation.get("contact")),
        "salesperson": _quote_text(quotation.get("salesperson")),
        "source_job_no": _quote_text(quotation.get("source_job_no")),
        "currency": _quote_text(quotation.get("currency")) or "RMB",
        "tax_rate": float(tax_rate),
        "valid_until": valid_until,
        "subject": subject,
        "warranty": _quote_text(quotation.get("warranty")) or (DEFAULT_WARRANTY if is_new else ""),
        "incoterm": _quote_text(quotation.get("incoterm")) or (DEFAULT_INCOTERMS[0] if is_new else ""),
        "ship_to": _quote_text(quotation.get("ship_to")),
        "description": _quote_text(quotation.get("description")),
        "payment_terms": _quote_text(quotation.get("payment_terms")),
        "payment_language": payment_language,
        "delivery_terms": _quote_text(quotation.get("delivery_terms")),
        "status": _quote_text(quotation.get("status")) or "草稿",
        "notes": _quote_text(quotation.get("notes")),
        "items": items,
        "payment_items": payment_items,
    }
    return result


def _quotation_from_connection(conn, quote_id):
    row = conn.execute(
        'SELECT * FROM "quotation" WHERE "quote_id" = ?', (quote_id,)
    ).fetchone()
    if row is None:
        return None
    result = {field: row[field] for field in ("quote_id", *QUOTATION_FIELDS)}
    result["created_at"] = row["created_at"]
    result["updated_at"] = row["updated_at"]
    result["items"] = [
        {field: item[field] for field in ("item_id", *QUOTATION_ITEM_FIELDS)}
        for item in conn.execute(
            'SELECT * FROM "quotation_item" WHERE "quote_id" = ? '
            'ORDER BY "line_no", "item_id"', (quote_id,),
        )
    ]
    result["payment_items"] = [
        {field: item[field] for field in ("payment_item_id", *QUOTATION_PAYMENT_FIELDS)}
        for item in conn.execute(
            'SELECT * FROM "quotation_payment_item" WHERE "quote_id" = ? '
            'ORDER BY "line_no", "payment_item_id"', (quote_id,),
        )
    ]
    subtotal = sum(Decimal(str(item["amount"] or 0)) for item in result["items"])
    subtotal = subtotal.quantize(Decimal("0.01"), ROUND_HALF_UP)
    tax_total = (subtotal * Decimal(str(result["tax_rate"] or 0)) / 100).quantize(
        Decimal("0.01"), ROUND_HALF_UP
    )
    result["subtotal"] = float(subtotal)
    result["tax_total"] = float(tax_total)
    result["grand_total"] = float(subtotal + tax_total)
    return result


def get_quotation(path, quote_id):
    conn = _connect(path)
    try:
        result = _quotation_from_connection(conn, quote_id)
        if result is None:
            raise KeyError(f"报价单不存在：{quote_id}")
        return result
    finally:
        conn.close()


def list_quotations(path, search=""):
    conn = _connect(path)
    try:
        params = []
        where = ""
        if _quote_text(search):
            where = (
                'WHERE q."quote_no" LIKE ? OR q."customer" LIKE ? '
                'OR q."subject" LIKE ? OR q."source_job_no" LIKE ? '
                'OR EXISTS (SELECT 1 FROM "quotation_item" qi WHERE qi."quote_id" = q."quote_id" AND qi."model" LIKE ?)'
            )
            needle = "%" + _quote_text(search) + "%"
            params = [needle] * 5
        ids = [row[0] for row in conn.execute(
            f'SELECT q."quote_id" FROM "quotation" q {where} '
            'ORDER BY q."quote_date" DESC, q."updated_at" DESC, q."quote_no" DESC',
            params,
        )]
        result = []
        for quote_id in ids:
            quote = _quotation_from_connection(conn, quote_id)
            result.append({
                "quote_id": quote["quote_id"], "quote_no": quote["quote_no"],
                "quote_date": quote["quote_date"], "customer": quote["customer"],
                "subject": quote["subject"], "source_job_no": quote["source_job_no"],
                "currency": quote["currency"], "status": quote["status"],
                "models": " / ".join(dict.fromkeys(
                    item["model"] for item in quote["items"] if item["model"]
                )),
                "item_count": len(quote["items"]), "subtotal": quote["subtotal"],
                "grand_total": quote["grand_total"], "updated_at": quote["updated_at"],
            })
        return result
    finally:
        conn.close()


def quotation_history(path, customer, model, exclude_quote_id=None):
    customer = _quote_text(customer)
    model = _quote_text(model)
    if not model:
        return []
    conn = _connect(path)
    try:
        sql = (
            'SELECT q."quote_id", q."quote_no", q."revision_label", q."quote_date", '
            'q."customer", q."currency", q."tax_rate", q."status", q."source_job_no", '
            'i."item_id", i."model", i."description", i."quantity", i."unit", '
            'i."unit_price", i."amount", i."remark" '
            'FROM "quotation" q JOIN "quotation_item" i ON i."quote_id" = q."quote_id" '
            'WHERE LOWER(TRIM(i."model")) = LOWER(TRIM(?))'
        )
        params = [model]
        if customer:
            sql += ' AND LOWER(TRIM(q."customer")) = LOWER(TRIM(?))'
            params.append(customer)
        if exclude_quote_id:
            sql += ' AND q."quote_id" <> ?'
            params.append(_quote_text(exclude_quote_id))
        sql += ' ORDER BY q."quote_date" DESC, q."updated_at" DESC, i."line_no"'
        return [dict(row) for row in conn.execute(sql, params)]
    finally:
        conn.close()


def _payment_option_from_term(term):
    kind = _quote_text(term.get("款类"))
    ratio = float(term.get("比例%") or 0)
    days = int(term.get("账期天数") or 0)
    trigger = _quote_text(term.get("触发条件"))
    desc_zh = _quote_text(term.get("说明"))
    if not desc_zh:
        desc_zh = f"{trigger + '，' if trigger else ''}{ratio:g}% {kind}"
    return {
        "kind": kind, "ratio": ratio, "days": days, "trigger": trigger,
        "description_zh": desc_zh,
        "description_en": _payment_description_en(kind, ratio, days, trigger),
    }


def quotation_options(path, customer="", language="zh"):
    customer = _quote_text(customer)
    language = "en" if _quote_text(language).lower().startswith("en") else "zh"
    conn = _connect(path)
    try:
        data = _data_from_connection(conn)
        exact = customer.casefold()
        payment_options = []
        seen = set()
        grouped = {}
        for term in data["付款条件"]:
            term_customer = _quote_text(term.get("客户"))
            option = _payment_option_from_term(term)
            if not option["kind"]:
                continue
            group_key = (term_customer.casefold(), _quote_text(term.get("JOB No")))
            grouped.setdefault(group_key, {"customer": term_customer, "job_no": _quote_text(term.get("JOB No")), "items": []})["items"].append(option)
        for group in grouped.values():
            items = group["items"]
            key = json.dumps(items, ensure_ascii=False, sort_keys=True)
            if key in seen:
                continue
            seen.add(key)
            payment_options.append({
                "customer": group["customer"], "job_no": group["job_no"], "items": items,
                "description_zh": "；".join(x["description_zh"] for x in items),
                "description_en": "; ".join(x["description_en"] for x in items),
                "kind": items[0]["kind"], "ratio": items[0]["ratio"], "days": items[0]["days"],
                "trigger": items[0]["trigger"],
                "priority": 0 if exact and group["customer"].casefold() == exact else 1,
            })
        for row in conn.execute(
            'SELECT q."customer", q."quote_no", p.* FROM "quotation_payment_item" p '
            'JOIN "quotation" q ON q."quote_id" = p."quote_id" '
            'ORDER BY q."quote_date" DESC, q."updated_at" DESC, p."line_no"'
        ):
            # Quotation child rows are grouped into one reusable payment preset.
            option = {field: row[field] for field in QUOTATION_PAYMENT_FIELDS if field != "line_no"}
            key = (row["customer"], row["quote_no"])
            group = next((x for x in payment_options if x.get("quote_no") == row["quote_no"]), None)
            if group is None:
                group = {"customer": row["customer"], "job_no": "", "quote_no": row["quote_no"], "items": [], "priority": 0 if exact and _quote_text(row["customer"]).casefold() == exact else 1}
                payment_options.append(group)
            group["items"].append(option)
        for option in payment_options:
            if option.get("items"):
                option["description_zh"] = "；".join(x.get("description_zh", "") for x in option["items"])
                option["description_en"] = "; ".join(x.get("description_en", "") for x in option["items"])
                first = option["items"][0]
                option.setdefault("kind", first.get("kind", "")); option.setdefault("ratio", first.get("ratio", 0)); option.setdefault("days", first.get("days", 0)); option.setdefault("trigger", first.get("trigger", ""))
        payment_options.sort(key=lambda x: (x.get("priority", 1), -len(_quote_text(x.get("job_no"))), _quote_text(x.get("job_no"))))
        recent_warranty = [row[0] for row in conn.execute(
            'SELECT value FROM "quotation_option" WHERE kind = ? AND language = ? '
            'ORDER BY updated_at DESC LIMIT 20', ("warranty", "zh")
        )]
        recent_incoterm = [row[0] for row in conn.execute(
            'SELECT value FROM "quotation_option" WHERE kind = ? AND language = ? '
            'ORDER BY updated_at DESC LIMIT 20', ("incoterm", "zh")
        )]
        warranty_options = list(dict.fromkeys([DEFAULT_WARRANTY, *recent_warranty]))
        incoterm_options = list(dict.fromkeys([*DEFAULT_INCOTERMS, *recent_incoterm]))
        return {
            "issuer_presets": list(ISSUER_PRESETS.values()),
            "default_issuer_key": "beijing",
            "default_warranty": DEFAULT_WARRANTY,
            "default_incoterm": DEFAULT_INCOTERMS[0],
            "default_payment_language": language,
            "warranty_options": warranty_options,
            "incoterm_options": incoterm_options,
            "payment_options": payment_options,
            "payment_language": language,
        }
    finally:
        conn.close()


def delete_quotation(path, quote_id, *, expected_revision=None, audit_context=None, backup=True):
    path = os.path.abspath(os.fspath(path))
    early_revision = get_revision(path) if os.path.exists(path) else 0
    if expected_revision is not None and early_revision != expected_revision:
        raise StaleImportError(f"数据库已从修订 {expected_revision} 更新到 {early_revision}，请重新加载后再删除报价")
    backup_path = _backup_database(path) if backup else None
    conn = _connect(path)
    try:
        conn.execute("BEGIN IMMEDIATE")
        current_revision = int(_meta(conn, "revision", "0") or 0)
        if expected_revision is not None and current_revision != expected_revision:
            raise StaleImportError(f"数据库已从修订 {expected_revision} 更新到 {current_revision}，请重新加载后再删除报价")
        old = _quotation_from_connection(conn, _quote_text(quote_id))
        if old is None:
            raise KeyError(f"报价单不存在：{quote_id}")
        conn.execute('DELETE FROM "quotation" WHERE "quote_id" = ?', (old["quote_id"],))
        revision = current_revision + 1
        now = datetime.now().isoformat(timespec="seconds")
        _set_meta(conn, "revision", revision)
        _set_meta(conn, "last_saved_at", now)
        conn.execute(
            "INSERT INTO _change_log(revision, created_at, reason, summary_json) VALUES(?, ?, ?, ?)",
            (revision, now, "quotation_delete", json.dumps({"报价单": 1, "报价明细": len(old["items"])}, ensure_ascii=False)),
        )
        # Record an immutable deletion snapshot for the header and every item.
        changes = [{"table_name": "报价单", "record_id": old["quote_id"], "job_no": old.get("source_job_no", ""), "operation": "delete", "field_name": "", "old_value": _quotation_audit_snapshot(old), "new_value": None, "origin": "user"}]
        changes.extend({"table_name": "报价明细", "record_id": item["item_id"], "job_no": old.get("source_job_no", ""), "operation": "delete", "field_name": "", "old_value": _quotation_item_audit_snapshot(item), "new_value": None, "origin": "user"} for item in old["items"])
        changes.extend({"table_name": "报价付款条件", "record_id": item["payment_item_id"], "job_no": old.get("source_job_no", ""), "operation": "delete", "field_name": "", "old_value": _quotation_payment_audit_snapshot(item), "new_value": None, "origin": "user"} for item in old["payment_items"])
        audit_event_id, _ = _append_audit_event(conn, revision, now, "quotation_delete", changes, audit_context)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return {"ok": True, "quote_id": old["quote_id"], "revision": revision, "backup": backup_path, "audit_event_id": audit_event_id, "audit_change_count": len(changes)}


def _quotation_audit_snapshot(quote):
    return {
        QUOTATION_FIELD_LABELS[field]: quote.get(field)
        for field in QUOTATION_FIELDS
    }


def _quotation_item_audit_snapshot(item):
    return {
        QUOTATION_ITEM_LABELS[field]: item.get(field)
        for field in QUOTATION_ITEM_FIELDS
    }


def _quotation_payment_audit_snapshot(item):
    return {
        QUOTATION_PAYMENT_LABELS[field]: item.get(field)
        for field in QUOTATION_PAYMENT_FIELDS
    }


def _quotation_audit_changes(old, new):
    changes = []
    job_no = new.get("source_job_no") or (old or {}).get("source_job_no") or ""
    quote_id = new["quote_id"]
    if old is None:
        changes.append({
            "table_name": "报价单", "record_id": quote_id, "job_no": job_no,
            "operation": "insert", "field_name": "", "old_value": None,
            "new_value": _quotation_audit_snapshot(new), "origin": "user",
        })
    else:
        for field in QUOTATION_FIELDS:
            if old.get(field) != new.get(field):
                changes.append({
                    "table_name": "报价单", "record_id": quote_id, "job_no": job_no,
                    "operation": "update", "field_name": QUOTATION_FIELD_LABELS[field],
                    "old_value": old.get(field), "new_value": new.get(field),
                    "origin": "user",
                })
    old_items = {item["item_id"]: item for item in (old or {}).get("items", [])}
    new_items = {item["item_id"]: item for item in new.get("items", [])}
    for item_id in sorted(set(old_items) | set(new_items)):
        before, after = old_items.get(item_id), new_items.get(item_id)
        if before is None:
            changes.append({
                "table_name": "报价明细", "record_id": item_id, "job_no": job_no,
                "operation": "insert", "field_name": "", "old_value": None,
                "new_value": _quotation_item_audit_snapshot(after), "origin": "user",
            })
        elif after is None:
            changes.append({
                "table_name": "报价明细", "record_id": item_id, "job_no": job_no,
                "operation": "delete", "field_name": "",
                "old_value": _quotation_item_audit_snapshot(before), "new_value": None,
                "origin": "user",
            })
        else:
            for field in QUOTATION_ITEM_FIELDS:
                if before.get(field) != after.get(field):
                    changes.append({
                        "table_name": "报价明细", "record_id": item_id,
                        "job_no": job_no, "operation": "update",
                        "field_name": QUOTATION_ITEM_LABELS[field],
                        "old_value": before.get(field), "new_value": after.get(field),
                        "origin": "user",
                    })
    old_payments = {item["payment_item_id"]: item for item in (old or {}).get("payment_items", [])}
    new_payments = {item["payment_item_id"]: item for item in new.get("payment_items", [])}
    for item_id in sorted(set(old_payments) | set(new_payments)):
        before, after = old_payments.get(item_id), new_payments.get(item_id)
        if before is None:
            changes.append({"table_name": "报价付款条件", "record_id": item_id, "job_no": job_no, "operation": "insert", "field_name": "", "old_value": None, "new_value": _quotation_payment_audit_snapshot(after), "origin": "user"})
        elif after is None:
            changes.append({"table_name": "报价付款条件", "record_id": item_id, "job_no": job_no, "operation": "delete", "field_name": "", "old_value": _quotation_payment_audit_snapshot(before), "new_value": None, "origin": "user"})
        else:
            for field in QUOTATION_PAYMENT_FIELDS:
                if before.get(field) != after.get(field):
                    changes.append({"table_name": "报价付款条件", "record_id": item_id, "job_no": job_no, "operation": "update", "field_name": QUOTATION_PAYMENT_LABELS[field], "old_value": before.get(field), "new_value": after.get(field), "origin": "user"})
    return changes


def save_quotation(path, quotation, *, expected_revision=None, audit_context=None,
                   backup=True):
    path = os.path.abspath(os.fspath(path))
    early_revision = get_revision(path) if os.path.exists(path) else 0
    if expected_revision is not None and early_revision != expected_revision:
        raise StaleImportError(
            f"数据库已从修订 {expected_revision} 更新到 {early_revision}，请重新加载后再保存报价"
        )
    backup_path = _backup_database(path) if backup else None
    conn = _connect(path)
    try:
        conn.execute("BEGIN IMMEDIATE")
        current_revision = int(_meta(conn, "revision", "0") or 0)
        if expected_revision is not None and current_revision != expected_revision:
            raise StaleImportError(
                f"数据库已从修订 {expected_revision} 更新到 {current_revision}，请重新加载后再保存报价"
            )
        prepared = _normalize_quotation(conn, quotation)
        old = _quotation_from_connection(conn, prepared["quote_id"])
        duplicate = conn.execute(
            'SELECT "quote_id" FROM "quotation" WHERE "quote_no" = ? AND "quote_id" <> ?',
            (prepared["quote_no"], prepared["quote_id"]),
        ).fetchone()
        if duplicate:
            raise QuotationValidationError(f"报价单号已存在：{prepared['quote_no']}")
        now = datetime.now().isoformat(timespec="seconds")
        prepared["created_at"] = old["created_at"] if old else now
        prepared["updated_at"] = now
        changes = _quotation_audit_changes(old, prepared)
        header_columns = ["quote_id", *QUOTATION_FIELDS, "created_at", "updated_at"]
        updates = [field for field in header_columns if field != "quote_id"]
        conn.execute(
            f'INSERT INTO "quotation" ({", ".join(_q(x) for x in header_columns)}) '
            f'VALUES ({", ".join("?" for _ in header_columns)}) '
            f'ON CONFLICT("quote_id") DO UPDATE SET '
            + ", ".join(f'{_q(field)}=excluded.{_q(field)}' for field in updates),
            [prepared[field] for field in header_columns],
        )
        conn.execute('DELETE FROM "quotation_item" WHERE "quote_id" = ?',
                     (prepared["quote_id"],))
        item_columns = ["item_id", "quote_id", *QUOTATION_ITEM_FIELDS]
        item_sql = (
            f'INSERT INTO "quotation_item" ({", ".join(_q(x) for x in item_columns)}) '
            f'VALUES ({", ".join("?" for _ in item_columns)})'
        )
        for item in prepared["items"]:
            values = [item["item_id"], prepared["quote_id"]]
            values.extend(item[field] for field in QUOTATION_ITEM_FIELDS)
            conn.execute(item_sql, values)
        conn.execute('DELETE FROM "quotation_payment_item" WHERE "quote_id" = ?',
                     (prepared["quote_id"],))
        payment_columns = ["payment_item_id", "quote_id", *QUOTATION_PAYMENT_FIELDS]
        payment_sql = (
            f'INSERT INTO "quotation_payment_item" ({", ".join(_q(x) for x in payment_columns)}) '
            f'VALUES ({", ".join("?" for _ in payment_columns)})'
        )
        for item in prepared["payment_items"]:
            values = [item["payment_item_id"], prepared["quote_id"]]
            values.extend(item[field] for field in QUOTATION_PAYMENT_FIELDS)
            conn.execute(payment_sql, values)
        for kind, value in (("warranty", prepared["warranty"]), ("incoterm", prepared["incoterm"])):
            if value:
                conn.execute(
                    'INSERT INTO "quotation_option" (option_id, kind, value, language, updated_at) '
                    'VALUES (?, ?, ?, ?, ?) ON CONFLICT(kind, value, language) DO UPDATE SET updated_at=excluded.updated_at',
                    (f"option-{uuid.uuid4().hex}", kind, value, "zh", now),
                )
        foreign_key_issues = conn.execute("PRAGMA foreign_key_check").fetchall()
        if foreign_key_issues:
            raise sqlite3.IntegrityError(f"外键检查失败：{foreign_key_issues[:3]}")
        revision = current_revision + 1
        _set_meta(conn, "revision", revision)
        _set_meta(conn, "last_saved_at", now)
        summary = {"报价单": 1, "报价明细": len(prepared["items"])}
        conn.execute(
            "INSERT INTO _change_log(revision, created_at, reason, summary_json) "
            "VALUES(?, ?, ?, ?)",
            (revision, now, "quotation_save", json.dumps(summary, ensure_ascii=False)),
        )
        audit_event_id = None
        if changes:
            audit_event_id, _ = _append_audit_event(
                conn, revision, now, "quotation_save", changes, audit_context
            )
        conn.commit()
        saved = _quotation_from_connection(conn, prepared["quote_id"])
    except sqlite3.IntegrityError as exc:
        conn.rollback()
        raise QuotationValidationError(f"报价数据关系检查失败：{exc}") from exc
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return {
        "ok": True, "revision": revision, "backup": backup_path,
        "audit_event_id": audit_event_id,
        "audit_change_count": len(changes) if audit_event_id else 0,
        "quotation": saved,
    }


def quotation_defaults(path, job_no):
    job_no = _quote_text(job_no)
    data = core.derive(load_database(path))
    contract = next(
        (row for row in data["合同订单"] if _quote_text(row.get("JOB No")) == job_no),
        None,
    )
    if contract is None:
        raise QuotationValidationError(f"未找到订单：{job_no}")
    grouped = {}
    for device in data["设备台账"]:
        if _quote_text(device.get("JOB No")) != job_no:
            continue
        model = _quote_text(device.get("设备型号"))
        if not model:
            continue
        price = float(device.get("未税单价") or 0)
        key = (model, price)
        grouped[key] = grouped.get(key, 0) + 1
    items = [{
        "model": model, "description": "", "quantity": quantity,
        "unit": "台", "unit_price": price,
    } for (model, price), quantity in grouped.items()]
    if not items:
        contract_models = [
            value.strip() for value in re.split(
                r"[;；、\n]+", _quote_text(contract.get("设备型号"))
            ) if value.strip()
        ]
        for model in dict.fromkeys(contract_models):
            items.append({
                "model": model, "description": "", "quantity": 1,
                "unit": "台", "unit_price": 0,
            })
    subject = " / ".join(dict.fromkeys(item["model"] for item in items))
    return {
        "quote_date": date.today().isoformat(),
        "quote_no": suggest_quotation_number(path),
        "revision_label": "1.0",
        "customer": _quote_text(contract.get("客户")),
        "salesperson": _quote_text(contract.get("担当者")),
        "source_job_no": job_no,
        "currency": _quote_text(contract.get("币种")) or "RMB",
        "subject": subject or _quote_text(contract.get("设备型号")),
        "ship_to": _quote_text(contract.get("送货地点")),
        "description": _quote_text(contract.get("订单内容")),
        "payment_terms": _quote_text(contract.get("付款条件")),
        "tax_rate": 13,
        "status": "草稿",
        "items": items,
    }


def _decoded_json(value):
    try:
        return json.loads(value)
    except (TypeError, ValueError):
        return None


def _audit_changes_for_event(conn, event_id):
    rows = conn.execute(
        "SELECT table_name, record_id, job_no, operation, field_name, old_value_json, "
        "new_value_json, origin FROM _audit_change WHERE event_id=? ORDER BY change_id",
        (event_id,),
    )
    return [{
        "table_name": row["table_name"],
        "record_id": row["record_id"],
        "job_no": row["job_no"],
        "operation": row["operation"],
        "field_name": row["field_name"],
        "old_value": _decoded_json(row["old_value_json"]),
        "new_value": _decoded_json(row["new_value_json"]),
        "origin": row["origin"],
    } for row in rows]


def _audit_event_from_row(row, changes):
    return {
        "event_id": row["event_id"],
        "revision": row["revision"],
        "created_at": row["created_at"],
        "operator_name": row["operator_name"],
        "source": row["source"],
        "actions": _decoded_json(row["actions_json"]) or [],
        "app_version": row["app_version"],
        "platform": row["platform"],
        "hostname": row["hostname"],
        "summary": _decoded_json(row["summary_json"]) or {},
        "change_count": row["change_count"],
        "previous_hash": row["previous_hash"],
        "event_hash": row["event_hash"],
        "changes": changes,
    }


def verify_audit_chain(path):
    conn = _connect(path)
    try:
        previous_hash = ""
        checked = 0
        for row in conn.execute("SELECT * FROM _audit_event ORDER BY revision, event_id"):
            changes = _audit_changes_for_event(conn, row["event_id"])
            event = _audit_event_from_row(row, changes)
            payload_event = {
                key: event[key] for key in (
                    "event_id", "revision", "created_at", "operator_name", "source",
                    "actions", "app_version", "platform", "hostname", "summary"
                )
            }
            canonical = json.dumps(
                {"event": payload_event, "changes": changes}, ensure_ascii=False,
                sort_keys=True, separators=(",", ":"), default=str,
            )
            expected = hashlib.sha256(
                (previous_hash + "\n" + canonical).encode("utf-8")
            ).hexdigest()
            if row["previous_hash"] != previous_hash:
                return {"ok": False, "checked_events": checked,
                        "error": f"修订 {row['revision']} 的前序哈希不匹配"}
            if row["event_hash"] != expected:
                return {"ok": False, "checked_events": checked,
                        "error": f"修订 {row['revision']} 的审计内容已被改变"}
            previous_hash = expected
            checked += 1
        head = _meta(conn, "audit_chain_head", "")
        if head != previous_hash:
            return {"ok": False, "checked_events": checked,
                    "error": "审计链尾标记不匹配，可能有记录被删除"}
        return {
            "ok": True,
            "checked_events": checked,
            "head": head,
            "started_at": _meta(conn, "audit_started_at"),
            "start_revision": int(_meta(conn, "audit_start_revision", "0") or 0),
        }
    finally:
        conn.close()


def get_audit_events(path, filters=None, limit=500):
    filters = filters if isinstance(filters, dict) else {}
    conn = _connect(path)
    try:
        rows = conn.execute(
            "SELECT * FROM _audit_event ORDER BY revision DESC, event_id DESC LIMIT ?",
            (max(1, min(int(limit or 500), 5000)),),
        ).fetchall()
        result = []
        for row in rows:
            changes = _audit_changes_for_event(conn, row["event_id"])
            event = _audit_event_from_row(row, changes)
            if filters.get("operator_name") and filters["operator_name"] not in event["operator_name"]:
                continue
            if filters.get("source") and filters["source"] != event["source"]:
                continue
            if filters.get("date_from") and event["created_at"][:10] < filters["date_from"]:
                continue
            if filters.get("date_to") and event["created_at"][:10] > filters["date_to"]:
                continue
            table_name = filters.get("table_name")
            job_no = filters.get("job_no")
            operation = filters.get("operation")
            if table_name and not any(x["table_name"] == table_name for x in changes):
                continue
            if job_no and not any(job_no in x["job_no"] for x in changes):
                continue
            if operation and not any(x["operation"] == operation for x in changes):
                continue
            result.append(event)
        return {
            "events": result,
            "chain": verify_audit_chain(path),
            "total_returned": len(result),
        }
    finally:
        conn.close()


def export_audit_xlsx(path, events):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "审计记录"
    headers = [
        "时间", "修订", "操作人", "来源", "操作动作", "业务表", "JOB No",
        "操作类型", "字段", "旧值", "新值", "变化来源", "记录ID", "事件哈希",
    ]
    sheet.append(headers)
    for event in events:
        for change in event.get("changes", []):
            sheet.append([
                event.get("created_at", ""), event.get("revision", ""),
                event.get("operator_name", ""), event.get("source", ""),
                "、".join(event.get("actions") or []), change.get("table_name", ""),
                change.get("job_no", ""), change.get("operation", ""),
                change.get("field_name", ""),
                _json_scalar(change.get("old_value")),
                _json_scalar(change.get("new_value")), change.get("origin", ""),
                change.get("record_id", ""), event.get("event_hash", ""),
            ])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="004494")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [20, 9, 16, 18, 24, 14, 16, 12, 18, 40, 40, 14, 28, 24]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    workbook.save(path)
    return path


def migrate_json_to_database(json_path, database_path):
    json_path = os.path.abspath(os.fspath(json_path))
    database_path = os.path.abspath(os.fspath(database_path))
    if os.path.exists(database_path):
        return {
            "ok": True,
            "path": database_path,
            "revision": get_revision(database_path),
            "data": load_database(database_path),
            "migrated": False,
        }
    data = core.load_store(json_path)
    rules = core.get_rules(json_path)
    result = save_database(
        database_path, data, rules, reason="json_migration", backup=False
    )
    result["migrated"] = True
    result["source"] = json_path
    return result


def _canonical_hash(data):
    payload = json.dumps(data, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _safe_excel_text(value):
    text = str(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _excel_value(value, field_type):
    if value is None or value == "":
        return None
    if field_type == core.DATE:
        try:
            return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()
        except ValueError:
            return _safe_excel_text(value)
    if field_type == core.NUMBER:
        return float(value)
    if field_type == core.INT:
        return int(float(value))
    return _safe_excel_text(value)


def export_editable_workbook(database_path, workbook_path):
    database_path = os.path.abspath(os.fspath(database_path))
    workbook_path = os.path.abspath(os.fspath(workbook_path))
    data = load_database(database_path)
    info = get_database_info(database_path)
    wb = Workbook()
    help_ws = wb.active
    help_ws.title = HELP_SHEET
    help_ws.sheet_view.showGridLines = False
    help_ws["A1"] = "上海康肯销售订单管理系统 · 外部编辑副本"
    help_ws["A1"].font = Font(name="Noto Sans SC", size=20, bold=True, color="004494")
    instructions = [
        "使用方法",
        "1. 在六个业务工作表中修改数据；可新增整行，也可删除整行。",
        "2. 不要重命名工作表或表头；隐藏的 _记录ID 列用于可靠识别原记录。",
        "3. 自动计算字段未导出，重新导入后由上海康肯销售订单管理系统统一计算。",
        "4. 回到应用点击“导入变更”，先核对差异，再勾选二次确认后写入。",
        "5. 若数据库在导出后已被保存过，旧副本会被拒绝，请重新导出，避免覆盖新数据。",
        "安全保障：导入使用事务、外键校验和写入前数据库备份。",
        f"数据库：{os.path.basename(database_path)}",
        f"导出修订：{info['revision']}",
        f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
    ]
    for row, text in enumerate(instructions, 3):
        help_ws.cell(row, 1, text)
        help_ws.cell(row, 1).font = Font(
            name="Noto Sans SC", size=11, bold=(row == 3), color="1C1C1C"
        )
        help_ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    help_ws.column_dimensions["A"].width = 96
    help_ws.row_dimensions[1].height = 34

    header_fill = PatternFill("solid", fgColor="004494")
    accent_fill = PatternFill("solid", fgColor="13BCBC")
    stripe_fill = PatternFill("solid", fgColor="F4F8FB")
    for table_index, table in enumerate(core.TABLES):
        ws = wb.create_sheet(table)
        ws.sheet_properties.tabColor = "13BCBC" if table_index % 2 else "004494"
        editable = [f for f in core.SCHEMA[table] if f[0] not in core.DERIVED.get(table, [])]
        headers = [ID_HEADER, *[f[0] for f in editable]]
        ws.append(headers)
        for cell in ws[1]:
            cell.fill = accent_fill if cell.column == 1 else header_fill
            cell.font = Font(name="Noto Sans SC", bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 25
        for rec in data[table]:
            ws.append([
                rec.get("记录ID"),
                *[_excel_value(rec.get(field), field_type)
                  for field, field_type, *_ in editable],
            ])
        ws.column_dimensions["A"].hidden = True
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(2, ws.max_row)}"
        for row in range(2, ws.max_row + 1):
            if row % 2 == 0:
                for cell in ws[row]:
                    cell.fill = stripe_fill
            for cell in ws[row]:
                cell.font = Font(name="Noto Sans SC", size=10, color="1C1C1C")
        for col, (field, field_type, *options) in enumerate(editable, 2):
            letter = get_column_letter(col)
            ws.column_dimensions[letter].width = min(34, max(12, len(field) * 2 + 4))
            if field_type == core.DATE:
                for cell in ws[letter][1:]:
                    cell.number_format = "yyyy-mm-dd"
            if field_type in (core.NUMBER, core.INT):
                for cell in ws[letter][1:]:
                    cell.number_format = "0.###"
            choices = options[0] if options else []
            if choices:
                formula = '"' + ",".join(str(x) for x in choices if x != "") + '"'
                if len(formula) <= 255:
                    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
                    dv.error = "请从列表选择，或清空单元格"
                    dv.errorTitle = "值不在允许列表"
                    ws.add_data_validation(dv)
                    dv.add(f"{letter}2:{letter}1048576")
        ws.conditional_formatting.add(
            f"B2:{get_column_letter(len(headers))}{max(2, ws.max_row)}",
            FormulaRule(formula=["MOD(ROW(),2)=0"], fill=stripe_fill),
        )

    meta = wb.create_sheet(META_SHEET)
    metadata = {
        "format_version": WORKBOOK_VERSION,
        "database_id": info["database_id"],
        "base_revision": info["revision"],
        "base_hash": _canonical_hash(data),
        "exported_at": datetime.now().isoformat(timespec="seconds"),
    }
    for key, value in metadata.items():
        meta.append([key, value])
    meta.sheet_state = "veryHidden"
    os.makedirs(os.path.dirname(workbook_path), exist_ok=True)
    wb.save(workbook_path)
    return {
        "ok": True,
        "path": workbook_path,
        "revision": info["revision"],
        "rows": sum(len(data[t]) for t in core.TABLES),
    }


def _workbook_metadata(wb):
    if META_SHEET not in wb.sheetnames:
        raise SpreadsheetFormatError("这不是上海康肯销售订单管理系统导出的可编辑工作簿（缺少元数据）")
    ws = wb[META_SHEET]
    return {str(row[0].value): row[1].value for row in ws.iter_rows(min_col=1, max_col=2)
            if row[0].value not in (None, "")}


def _cell_value(cell, field, field_type, table, row_number):
    if cell.data_type == "f":
        raise SpreadsheetFormatError(
            f"{table} 第 {row_number} 行“{field}”包含公式；导入只接受值，请粘贴为值后重试"
        )
    value = cell.value
    if value is None or value == "":
        return None if field_type in (core.NUMBER, core.INT) else ""
    if field_type == core.DATE:
        if isinstance(value, (datetime, date)):
            return value.strftime("%Y-%m-%d")
        text = str(value).strip()
        normalized = core._norm(text, core.DATE)
        try:
            datetime.strptime(normalized, "%Y-%m-%d")
        except ValueError as exc:
            raise SpreadsheetFormatError(
                f"{table} 第 {row_number} 行“{field}”不是有效日期：{value!r}"
            ) from exc
        return normalized
    if field_type in (core.NUMBER, core.INT):
        try:
            number = float(str(value).replace(",", ""))
        except (TypeError, ValueError) as exc:
            raise SpreadsheetFormatError(
                f"{table} 第 {row_number} 行“{field}”不是数字：{value!r}"
            ) from exc
        return int(round(number)) if field_type == core.INT else number
    text = str(value)
    if len(text) > 1 and text[0] == "'" and text[1] in "=+-@":
        text = text[1:]
    return text


def _record_label(table, rec):
    job = str(rec.get("JOB No") or "（无 JOB）")
    secondaries = {
        "合同订单": "客户",
        "付款条件": "款类",
        "设备台账": "製造番号",
        "发货批次": "发货批次",
        "开票记录": "款类",
        "回款记录": "款类",
    }
    secondary = str(rec.get(secondaries[table]) or "").strip()
    return job + (" / " + secondary if secondary else "")


def _diff_data(before, after):
    changes = []
    action_counts = {"新增": 0, "修改": 0, "删除": 0}
    table_counts = {t: {"新增": 0, "修改": 0, "删除": 0} for t in core.TABLES}
    for table in core.TABLES:
        old = {str(r.get("记录ID")): r for r in before[table]}
        new = {str(r.get("记录ID")): r for r in after[table]}
        field_names = [f[0] for f in core.SCHEMA[table]]
        for rid in new.keys() - old.keys():
            fields = [
                {"field": f, "before": "", "after": new[rid].get(f)}
                for f in field_names if new[rid].get(f) not in (None, "")
            ]
            changes.append({"table": table, "action": "新增", "record_id": rid,
                            "label": _record_label(table, new[rid]), "fields": fields})
            action_counts["新增"] += 1
            table_counts[table]["新增"] += 1
        for rid in old.keys() - new.keys():
            changes.append({"table": table, "action": "删除", "record_id": rid,
                            "label": _record_label(table, old[rid]), "fields": []})
            action_counts["删除"] += 1
            table_counts[table]["删除"] += 1
        for rid in old.keys() & new.keys():
            fields = []
            for field in field_names:
                before_value = old[rid].get(field)
                after_value = new[rid].get(field)
                if before_value != after_value:
                    fields.append({"field": field, "before": before_value,
                                   "after": after_value,
                                   "derived": field in core.DERIVED.get(table, [])})
            if fields:
                changes.append({"table": table, "action": "修改", "record_id": rid,
                                "label": _record_label(table, new[rid]), "fields": fields})
                action_counts["修改"] += 1
                table_counts[table]["修改"] += 1
    order = {"删除": 0, "修改": 1, "新增": 2}
    changes.sort(key=lambda x: (core.TABLES.index(x["table"]), order[x["action"]], x["label"]))
    return changes, action_counts, table_counts


def _replace_coverage_token(value, old, new):
    """替换覆盖串中的完整键，同时兼容调拨番号的箭头写法。"""
    if not old or old == new:
        return value
    changed = False
    result = []
    for item in core.coverage_values(value):
        pieces = re.split(r"([→⇒⟶➡➝])", item)
        for index in range(0, len(pieces), 2):
            if pieces[index] == old:
                pieces[index] = new
                changed = True
        result.append("".join(pieces))
    return ";".join(result) if changed else value


def _propagate_import_relational_edits(current, incoming):
    """用稳定记录 ID 识别 Excel 改名，同步仍保留旧键的关联行。"""
    current_by_id = {
        table: {core.s(rec.get("记录ID")): rec for rec in current[table]
                if core.s(rec.get("记录ID"))}
        for table in core.TABLES
    }
    incoming_by_id = {
        table: {core.s(rec.get("记录ID")): rec for rec in incoming[table]
                if core.s(rec.get("记录ID"))}
        for table in core.TABLES
    }

    job_renames = {}
    for rid, old in current_by_id["合同订单"].items():
        new = incoming_by_id["合同订单"].get(rid)
        if new and core.s(old.get("JOB No")) != core.s(new.get("JOB No")):
            job_renames[core.s(old.get("JOB No"))] = core.s(new.get("JOB No"))
    for table in ("付款条件", "设备台账", "发货批次", "开票记录", "回款记录"):
        for rec in incoming[table]:
            old_job = core.s(rec.get("JOB No"))
            if old_job in job_renames:
                rec["JOB No"] = job_renames[old_job]

    for rid, old in current_by_id["发货批次"].items():
        new = incoming_by_id["发货批次"].get(rid)
        if not new:
            continue
        old_batch = core.s(old.get("发货批次"))
        new_batch = core.s(new.get("发货批次"))
        if not old_batch or old_batch == new_batch:
            continue
        job = core.s(new.get("JOB No"))
        for dev in incoming["设备台账"]:
            if (core.s(dev.get("JOB No")) == job
                    and core.s(dev.get("发货批次")) == old_batch):
                dev["发货批次"] = new_batch
        for table in ("开票记录", "回款记录"):
            for rec in incoming[table]:
                if core.s(rec.get("JOB No")) == job:
                    rec["覆盖批次"] = _replace_coverage_token(
                        rec.get("覆盖批次"), old_batch, new_batch
                    )

    for rid, old in current_by_id["设备台账"].items():
        new = incoming_by_id["设备台账"].get(rid)
        if not new:
            continue
        old_serial = core.s(old.get("製造番号"))
        new_serial = core.s(new.get("製造番号"))
        if not old_serial or old_serial == new_serial:
            continue
        job = core.s(new.get("JOB No"))
        for table in ("开票记录", "回款记录"):
            for rec in incoming[table]:
                if core.s(rec.get("JOB No")) == job:
                    rec["覆盖製造番号"] = _replace_coverage_token(
                        rec.get("覆盖製造番号"), old_serial, new_serial
                    )

    for rid, old in current_by_id["付款条件"].items():
        new = incoming_by_id["付款条件"].get(rid)
        if not new:
            continue
        old_kind = core.s(old.get("款类"))
        new_kind = core.s(new.get("款类"))
        if not old_kind or old_kind == new_kind:
            continue
        job = core.s(new.get("JOB No"))
        if any(rec is not new and core.s(rec.get("JOB No")) == job
               and core.s(rec.get("款类")) == old_kind
               for rec in incoming["付款条件"]):
            continue
        for table in ("开票记录", "回款记录"):
            for rec in incoming[table]:
                if (core.s(rec.get("JOB No")) == job
                        and core.s(rec.get("款类")) == old_kind):
                    rec["款类"] = new_kind

    return incoming


def prepare_editable_import(database_path, workbook_path):
    database_path = os.path.abspath(os.fspath(database_path))
    workbook_path = os.path.abspath(os.fspath(workbook_path))
    if not os.path.exists(workbook_path):
        raise FileNotFoundError(workbook_path)
    current = load_database(database_path)
    info = get_database_info(database_path)
    wb = load_workbook(workbook_path, data_only=False)
    metadata = _workbook_metadata(wb)
    if int(metadata.get("format_version") or 0) != WORKBOOK_VERSION:
        raise SpreadsheetFormatError("工作簿格式版本不兼容，请从当前应用重新导出")
    if str(metadata.get("database_id") or "") != info["database_id"]:
        raise SpreadsheetFormatError("工作簿来自另一个数据库，不能导入当前数据库")
    try:
        base_revision = int(metadata.get("base_revision"))
    except (TypeError, ValueError) as exc:
        raise SpreadsheetFormatError("工作簿缺少有效的数据库修订号") from exc
    base_hash = str(metadata.get("base_hash") or "")
    current_hash = _canonical_hash(current)
    if len(base_hash) != 64:
        raise SpreadsheetFormatError("工作簿缺少有效的导出数据摘要，请重新导出")
    if base_revision == info["revision"] and base_hash != current_hash:
        raise StaleImportError("数据库内容已被外部修改，请重新导出编辑副本")

    incoming = core.empty_data()
    for table in core.TABLES:
        if table not in wb.sheetnames:
            raise SpreadsheetFormatError(f"缺少工作表：{table}")
        ws = wb[table]
        editable = [f for f in core.SCHEMA[table] if f[0] not in core.DERIVED.get(table, [])]
        expected_headers = [ID_HEADER, *[f[0] for f in editable]]
        headers = [cell.value for cell in ws[1]]
        actual = [x for x in headers if x not in (None, "")]
        if len(actual) != len(set(actual)):
            raise SpreadsheetFormatError(f"{table} 表头存在重复列")
        if set(actual) != set(expected_headers):
            missing = [x for x in expected_headers if x not in actual]
            extra = [x for x in actual if x not in expected_headers]
            details = []
            if missing:
                details.append("缺少 " + "、".join(missing))
            if extra:
                details.append("多出 " + "、".join(str(x) for x in extra))
            raise SpreadsheetFormatError(f"{table} 表头已改变：" + "；".join(details))
        col_by_header = {value: index + 1 for index, value in enumerate(headers)
                         if value not in (None, "")}
        current_by_id = {str(r.get("记录ID") or ""): r for r in current[table]}
        seen_ids = set()
        for row_number in range(2, ws.max_row + 1):
            cells = [ws.cell(row_number, col_by_header[h]) for h in expected_headers]
            if not any(cell.value not in (None, "") for cell in cells):
                continue
            id_cell = cells[0]
            if id_cell.data_type == "f":
                raise SpreadsheetFormatError(f"{table} 第 {row_number} 行记录ID不能是公式")
            rid = str(id_cell.value or "").strip()
            if rid in seen_ids:
                rid = ""  # 复制整行时隐藏 ID 也会被复制；后续重复项按新增记录处理。
            elif rid:
                seen_ids.add(rid)
            if rid and rid not in current_by_id:
                raise SpreadsheetFormatError(
                    f"{table} 第 {row_number} 行记录ID不属于当前数据库，请勿修改隐藏的 _记录ID 列"
                )
            rec = dict(current_by_id[rid]) if rid else {}
            rec["记录ID"] = rid
            for field, field_type, *_ in editable:
                cell = ws.cell(row_number, col_by_header[field])
                rec[field] = _cell_value(cell, field, field_type, table, row_number)
            incoming[table].append(rec)

    incoming = _propagate_import_relational_edits(current, incoming)
    prepared_data, issues = _prepare_data(incoming)
    changes, action_counts, table_counts = _diff_data(current, prepared_data)
    stat = os.stat(workbook_path)
    return PreparedImport(
        database_path=database_path,
        workbook_path=workbook_path,
        database_id=info["database_id"],
        database_hash=current_hash,
        base_revision=base_revision,
        current_revision=info["revision"],
        data=prepared_data,
        rules=get_rules(database_path),
        changes=changes,
        action_counts=action_counts,
        table_counts=table_counts,
        issues=issues,
        file_signature=(stat.st_size, stat.st_mtime_ns),
    )


def apply_editable_import(database_path, prepared, audit_context=None):
    database_path = os.path.abspath(os.fspath(database_path))
    if database_path != prepared.database_path:
        raise StaleImportError("预览对应的数据库已切换，请重新选择文件并预览")
    current_info = get_database_info(database_path)
    if prepared.stale or current_info["revision"] != prepared.current_revision:
        raise StaleImportError("数据库在工作簿导出或差异预览后已发生变化，请重新导出/预览")
    if current_info["database_id"] != prepared.database_id:
        raise StaleImportError("当前数据库与差异预览不匹配")
    if _canonical_hash(load_database(database_path)) != prepared.database_hash:
        raise StaleImportError("数据库内容在差异预览后又被修改，请重新预览")
    stat = os.stat(prepared.workbook_path)
    if (stat.st_size, stat.st_mtime_ns) != prepared.file_signature:
        raise StaleImportError("Excel 文件在差异预览后又被修改，请重新预览")
    return save_database(
        database_path,
        prepared.data,
        prepared.rules,
        reason="xlsx_import",
        backup=True,
        expected_revision=prepared.current_revision,
        audit_context=audit_context,
    )
