#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate the Japanese 装置搬入依頼書 used by the shipment workflow.

The supplied workbook is an old BIFF ``.xls`` file with embedded stamp objects
and unreliable pagination.  This module rebuilds the printable form as a clean
A4 ``.xlsx``.  The first worksheet uses the supplied two-column equipment
layout: slots 1--30 are shown as fifteen paired rows, and later blocks repeat
that layout on a new printed page before the request details continue.
"""

from __future__ import annotations

import os
import tempfile
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break


THIN = Side(style="thin", color="000000")
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LABEL_FILL = PatternFill("solid", fgColor="E0E0E0")
FONT = Font(name="ＭＳ Ｐ明朝", size=12, color="000000")
SMALL_FONT = Font(name="ＭＳ Ｐ明朝", size=10, color="000000")
TITLE_FONT = Font(name="ＭＳ Ｐ明朝", size=18, color="000000")
DETAIL_FONT = Font(name="ＭＳ Ｐ明朝", size=10, color="000000")

NOTE_LINE_WIDTH = 40
FORM_LAST_COL = 16
DEVICE_ROWS_PER_BLOCK = 15
DEVICE_SLOTS_PER_BLOCK = DEVICE_ROWS_PER_BLOCK * 2


def _text(value):
    raw = "" if value is None else str(value)
    return "'" + raw if raw.startswith(("=", "+", "-", "@")) else raw


def _date_value(value):
    if isinstance(value, (datetime, date)):
        return value.date() if isinstance(value, datetime) else value
    try:
        return date.fromisoformat(str(value))
    except (TypeError, ValueError):
        return _text(value)


def _date_label(value):
    parsed = _date_value(value)
    if isinstance(parsed, date):
        return f"{parsed.year}/{parsed.month}/{parsed.day}"
    return parsed


def _set(ws, cell, value="", *, font=None, alignment=None, border=None,
         fill=None, number_format=None):
    target = ws[cell]
    target.value = _date_value(value) if number_format == "date" else (
        _text(value) if isinstance(value, str) else value
    )
    target.font = font or FONT
    target.alignment = alignment or Alignment(vertical="center")
    if border:
        target.border = border
    if fill:
        target.fill = fill
    if number_format == "date":
        target.number_format = "[$-804]yyyy\\年m\\月d\\日"
    elif number_format:
        target.number_format = number_format
    return target


def _merge(ws, start_row, start_col, end_row, end_col, value="", **style):
    ws.merge_cells(start_row=start_row, start_column=start_col,
                   end_row=end_row, end_column=end_col)
    border = style.get("border")
    fill = style.get("fill")
    for row in ws.iter_rows(min_row=start_row, max_row=end_row,
                            min_col=start_col, max_col=end_col):
        for cell in row:
            if border:
                cell.border = border
            if fill:
                cell.fill = fill
            cell.font = style.get("font") or FONT
            cell.alignment = style.get("alignment") or Alignment(vertical="center")
    return _set(ws, f"{chr(64 + start_col)}{start_row}", value, **style)


def _row_style(ws, row, start=1, end=17, *, height=21, border=GRID,
               fill=None, font=None):
    ws.row_dimensions[row].height = height
    for col in range(start, end + 1):
        cell = ws.cell(row, col)
        cell.font = font or FONT
        cell.alignment = Alignment(vertical="center")
        if border:
            cell.border = border
        if fill:
            cell.fill = fill


def _field_label(ws, row, number, label, end_row=None):
    end_row = end_row or row
    _merge(ws, row, 1, end_row, 1, number, font=FONT, fill=LABEL_FILL,
           border=GRID, alignment=Alignment(horizontal="center", vertical="center"))
    _merge(ws, row, 2, end_row, 2, label, font=FONT, fill=LABEL_FILL,
           border=GRID, alignment=Alignment(horizontal="center", vertical="center"))


def shipment_defaults(data, shipment_id=None, job_no=None, batch=None,
                      issue_date=None):
    """Build form defaults from an authoritative database-shaped data object."""
    shipments = data.get("发货批次", [])
    shipment = None
    if shipment_id:
        shipment = next((r for r in shipments
                         if str(r.get("记录ID", "")) == str(shipment_id)), None)
    if shipment is None and job_no is not None and batch is not None:
        shipment = next((r for r in shipments
                         if str(r.get("JOB No", "")) == str(job_no)
                         and str(r.get("发货批次", "")) == str(batch)), None)
    if shipment is None:
        raise ValueError("找不到所选发货批次，请刷新数据后重试")

    job = str(shipment.get("JOB No") or "")
    batch_name = str(shipment.get("发货批次") or "")
    devices = [r for r in data.get("设备台账", [])
               if str(r.get("JOB No") or "") == job
               and str(r.get("发货批次") or "") == batch_name]
    if not devices:
        raise ValueError(f"批次 {job} / {batch_name} 没有关联设备，无法生成搬入依頼書")
    contract = next((r for r in data.get("合同订单", [])
                     if str(r.get("JOB No") or "") == job), {})
    customer = str(contract.get("客户") or "")
    address = str(contract.get("送货地点") or "")
    ship_from = str(contract.get("发货地点") or "")
    models = sorted({str(r.get("设备型号") or "") for r in devices if r.get("设备型号")})
    machine_numbers = [str(r.get("機番") or "") for r in devices]
    serials = [str(r.get("製造番号") or "") for r in devices]
    note_lines = []
    if ship_from:
        note_lines.append(f"※出荷元：{ship_from}")
    if shipment.get("出荷日"):
        note_lines.append(f"※出荷日：{shipment.get('出荷日')}")
    return {
        "shipment_id": str(shipment.get("记录ID") or ""),
        "job_no": job,
        "batch": batch_name,
        "customer": customer,
        "address": address,
        "ship_from": ship_from,
        "ship_date": str(shipment.get("出荷日") or ""),
        "devices": [
            {"model": str(r.get("设备型号") or ""),
             "machine_no": str(r.get("機番") or ""),
             "serial_no": str(r.get("製造番号") or ""),
             "po_no": str(r.get("PO No") or "")}
            for r in devices
        ],
        "models": models,
        "has_customer_id": all(machine_numbers),
        "issue_date": str(issue_date or date.today().isoformat()),
        "move_in_date": "",
        "move_in_time": "",
        "contact": "",
        "vehicle_type": "",
        "vehicle_tonnage": "",
        "vehicle_note": "",
        "department": "本社工場　/　物流購買部　/　技術設計部　/　技術サービス部",
        "other_department": "平湖康肯",
        "system_notes": "\n".join(note_lines),
        "notes": "",
        "manager": "",
        "issuer": "",
    }


def _wrap_note_lines(values, width=NOTE_LINE_WIDTH):
    """Wrap note lines to a conservative width for the merged A4 note cells."""
    lines = []
    for value in values:
        text = "" if value is None else str(value)
        if not text:
            lines.append("")
            continue
        while len(text) > width:
            lines.append(text[:width])
            text = text[width:]
        lines.append(text)
    return lines


def _write_device_headers(ws, row):
    """Write one 1--30-style equipment header row."""
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    headers = ((1, 1, "No."), (2, 2, "PO No."), (3, 4, "機種"),
               (5, 6, "機番"), (7, 8, "制番"), (9, 9, "No."),
               (10, 10, "PO No."), (11, 12, "機種"), (13, 14, "機番"),
               (15, 16, "制番"))
    _row_style(ws, row, start=1, end=FORM_LAST_COL, height=24, border=GRID,
               fill=LABEL_FILL, font=DETAIL_FONT)
    for start_col, end_col, label in headers:
        _merge(ws, row, start_col, row, end_col, label,
               font=DETAIL_FONT, fill=LABEL_FILL, border=GRID, alignment=center)


def _write_device_side(ws, row, device, number, columns):
    """Write the non-grouped fields for one side of a paired device row."""
    no_col, po_col, model_start, model_end, machine_start, machine_end, serial_start, serial_end = columns
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    po_center = Alignment(horizontal="center", vertical="center",
                          shrink_to_fit=True)
    if device is None:
        _set(ws, f"{chr(64 + no_col)}{row}", "", font=DETAIL_FONT,
             fill=LABEL_FILL, border=GRID, alignment=center)
        po = model = machine = serial = ""
    else:
        _set(ws, f"{chr(64 + no_col)}{row}", number, font=DETAIL_FONT,
             fill=LABEL_FILL, border=GRID, alignment=center)
        po = device.get("po_no", "")
        model = device.get("model", "")
        machine = device.get("machine_no", "")
        serial = device.get("serial_no", "")
    _set(ws, f"{chr(64 + po_col)}{row}", po, font=DETAIL_FONT, border=GRID,
         alignment=po_center)
    _merge(ws, row, model_start, row, model_end, model, font=DETAIL_FONT,
           border=GRID, alignment=center)
    _merge(ws, row, machine_start, row, machine_end, machine,
           font=DETAIL_FONT, border=GRID, alignment=center)
    _merge(ws, row, serial_start, row, serial_end, serial,
           font=DETAIL_FONT, border=GRID, alignment=center)


def _merge_device_groups(ws, first_row, devices, columns):
    """Merge adjacent PO/model values within one side of a 30-slot block."""
    po_col, model_start, model_end = columns
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    po_center = Alignment(horizontal="center", vertical="center",
                          shrink_to_fit=True)
    for key, start_col, end_col in (
        ("po_no", po_col, po_col),
        ("model", model_start, model_end),
    ):
        group_start = 0
        for group_end in range(1, len(devices) + 1):
            at_end = group_end == len(devices)
            same_as_next = (not at_end
                            and devices[group_end].get(key, "")
                            == devices[group_start].get(key, ""))
            if at_end or not same_as_next:
                _merge(ws, first_row + group_start, start_col,
                       first_row + group_end - 1, end_col,
                       devices[group_start].get(key, ""),
                       font=DETAIL_FONT, border=GRID,
                       alignment=po_center if key == "po_no" else center)
                group_start = group_end


def _write_device_table(ws, header_row, devices):
    """Write paired 1--30 blocks and return the last row plus page breaks."""
    left_columns = (1, 2, 3, 4, 5, 6, 7, 8)
    right_columns = (9, 10, 11, 12, 13, 14, 15, 16)
    end_row = header_row
    page_breaks = []
    for block_start in range(0, len(devices), DEVICE_SLOTS_PER_BLOCK):
        if block_start:
            # The next block starts on a new page, so its own header is only
            # printed when another equipment block actually exists.
            page_breaks.append(end_row)
            header_row = end_row + 1
        _write_device_headers(ws, header_row)
        block = devices[block_start:block_start + DEVICE_SLOTS_PER_BLOCK]
        row_count = max(1, (len(block) + 1) // 2)
        first_data_row = header_row + 1
        for row_offset in range(row_count):
            row = first_data_row + row_offset
            _row_style(ws, row, start=1, end=FORM_LAST_COL, height=22,
                       border=GRID, font=DETAIL_FONT)
            left_index = row_offset * 2
            right_index = left_index + 1
            left_device = block[left_index] if left_index < len(block) else None
            right_device = block[right_index] if right_index < len(block) else None
            _write_device_side(ws, row, left_device, block_start + left_index + 1,
                               left_columns)
            _write_device_side(ws, row, right_device, block_start + right_index + 1,
                               right_columns)
        _merge_device_groups(ws, first_data_row,
                             [block[i] for i in range(0, len(block), 2)],
                             (2, 3, 4))
        right_devices = [block[i] for i in range(1, len(block), 2)]
        if right_devices:
            _merge_device_groups(ws, first_data_row, right_devices,
                                 (10, 11, 12))
        end_row = first_data_row + row_count - 1
    return end_row, page_breaks


def build_workbook(form):
    """Build the normalized KE-403-style workbook from a validated form."""
    wb = Workbook()
    ws = wb.active
    ws.title = "搬入依頼書"
    ws.sheet_view.showGridLines = False
    widths = {"A": 7.642857, "B": 17.142857, "C": 7.642857,
              "D": 7.642857, "E": 7.642857, "F": 7.642857,
              "G": 7.642857, "H": 7.642857, "I": 7.642857,
              "J": 17.107143, "K": 7.589286, "L": 7.642857,
              "M": 7.642857, "N": 7.642857, "O": 7.642857,
              "P": 8.43}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row, height in {1: 20.25, 2: 20.25, 3: 20.25, 4: 13.5, 5: 25.5,
                        6: 14.25, 7: 20.25, 8: 20.25, 9: 20.25,
                        10: 9, 11: 19.5}.items():
        ws.row_dimensions[row].height = height

    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)

    _set(ws, "A1", _date_label(form.get("issue_date", "")), font=FONT,
         alignment=left)
    _merge(ws, 1, 15, 1, 16, "[KE-403-BKTC]", font=SMALL_FONT,
           alignment=right_align())
    _merge(ws, 2, 1, 2, FORM_LAST_COL, form.get("department", ""),
           font=FONT, alignment=left)
    other = form.get("other_department", "")
    _merge(ws, 3, 1, 3, FORM_LAST_COL,
           f"他（　　　{other}　　　）" if other else "他（　　　　　　　　　）",
           font=FONT, alignment=left)
    _merge(ws, 5, 1, 5, FORM_LAST_COL, "装 置 搬 入 依 頼 書",
           font=TITLE_FONT, alignment=center)
    for cell, label, value in (("M6", "所属長", form.get("manager", "")),
                               ("O6", "発行者", form.get("issuer", ""))):
        _set(ws, cell, label, font=FONT, border=GRID, alignment=center)
        col = cell[0]
        start_col = ord(col) - 64
        end_col = start_col + 1
        _merge(ws, 7, start_col, 9, end_col, value, font=FONT,
               border=GRID, alignment=center)

    devices = form.get("devices") or []
    _field_label(ws, 11, 1, "設備明細")
    _set(ws, "C11", len(devices), font=FONT, fill=LABEL_FILL, border=GRID,
         alignment=center)
    _set(ws, "D11", "台", font=FONT, fill=LABEL_FILL, border=GRID,
         alignment=center)
    device_end_row, device_page_breaks = _write_device_table(ws, 12, devices)

    # Keep the request details directly below the equipment blocks.  A second
    # block starts on a new printed page, and carries its own header instead of
    # using global print-title rows that would repeat on notes-only pages.
    row = device_end_row + 1
    _field_label(ws, row, 2, "客先装置ID No.")
    _merge(ws, row, 3, row, 4, ("☑ 有り" if form.get("has_customer_id") else "☐ 有り"),
           font=FONT, border=GRID, alignment=center)
    _merge(ws, row, 5, row, 6, "ID No.",
           font=SMALL_FONT, border=GRID, alignment=center)
    _merge(ws, row, 7, row, 15, "", font=FONT, border=GRID, alignment=left)
    _set(ws, f"P{row}", ("☐ 無し" if form.get("has_customer_id") else "☑ 無し"),
           font=FONT, border=GRID, alignment=center)

    row += 1
    _field_label(ws, row, 3, "搬入先名")
    _merge(ws, row, 3, row, FORM_LAST_COL, form.get("customer", ""),
           font=FONT, border=GRID, alignment=left)
    row += 1
    _field_label(ws, row, 4, "搬入先住所")
    _merge(ws, row, 3, row, FORM_LAST_COL, form.get("address", ""),
           font=FONT, border=GRID, alignment=left)
    row += 1
    _field_label(ws, row, 5, "搬入連絡先")
    _merge(ws, row, 3, row, 4, "氏名", font=SMALL_FONT, border=GRID,
           alignment=center)
    _merge(ws, row, 5, row, 8, "", font=FONT, border=GRID, alignment=left)
    _set(ws, f"I{row}", "連絡先", font=SMALL_FONT, border=GRID,
         alignment=center)
    _merge(ws, row, 10, row, FORM_LAST_COL, form.get("contact", ""),
           font=FONT, border=GRID, alignment=left)
    row += 1
    ws.row_dimensions[row].height = 24
    _field_label(ws, row, 6, "搬入日時")
    _merge(ws, row, 3, row, 4, "搬入日", font=SMALL_FONT, border=GRID, alignment=center)
    _merge(ws, row, 5, row, 8, form.get("move_in_date", ""), font=FONT, border=GRID,
           alignment=center, number_format="date")
    _set(ws, f"I{row}", "時刻", font=SMALL_FONT, border=GRID, alignment=center)
    _merge(ws, row, 10, row, FORM_LAST_COL, form.get("move_in_time", ""),
           font=FONT, border=GRID, alignment=center)
    row += 1
    _field_label(ws, row, 7, "搬入車両", row + 3)
    vehicles = ["平ボディー", "ユニック車", "パワーゲート車", "その他"]
    chosen = str(form.get("vehicle_type") or "")
    for vehicle_row, label in enumerate(vehicles, row):
        ws.row_dimensions[vehicle_row].height = 17.6
        _row_style(ws, vehicle_row, start=3, end=FORM_LAST_COL, height=17.6,
                   border=GRID)
        _merge(ws, vehicle_row, 3, vehicle_row, 7,
               ("☑ " if chosen == label else "☐ ") + label,
               font=FONT, border=GRID, alignment=left)
        if vehicle_row < row + 3:
            _set(ws, f"H{vehicle_row}", "t", font=FONT, border=GRID,
                 alignment=center)
            _merge(ws, vehicle_row, 9, vehicle_row, FORM_LAST_COL,
                   str(form.get("vehicle_tonnage") or "").strip(),
                   font=SMALL_FONT, border=GRID, alignment=center)
        else:
            _set(ws, f"H{vehicle_row}", "", font=FONT, border=GRID,
                 alignment=center)
            _merge(ws, vehicle_row, 9, vehicle_row, FORM_LAST_COL,
                   form.get("vehicle_note", ""), font=SMALL_FONT, border=GRID,
                   alignment=left)
    row += 4

    system_notes = str(form.get("system_notes") or "").splitlines()
    if not system_notes:
        # Keep direct callers using the pre-v1.9 form shape readable.
        if form.get("ship_from"):
            system_notes.append(f"※出荷元：{form.get('ship_from')}")
        if form.get("ship_date"):
            system_notes.append(f"※出荷日：{form.get('ship_date')}")
    notes = _wrap_note_lines(system_notes + str(form.get("notes") or "").splitlines())
    note_start = row
    note_rows = max(8, len(notes), 1)
    note_end = note_start + note_rows - 1
    _field_label(ws, note_start, 8, "特記事項", note_end)
    for note_row in range(note_start, note_end + 1):
        _row_style(ws, note_row, start=3, end=FORM_LAST_COL, height=21,
                   border=GRID)
        _merge(ws, note_row, 3, note_row, FORM_LAST_COL,
               notes[note_row - note_start] if note_row - note_start < len(notes) else "",
               font=FONT, border=GRID, alignment=wrap)

    last_row = note_end
    ws.print_area = f"A1:P{last_row}"
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.25, bottom=0.25,
                                  header=0.1, footer=0.1)
    ws.sheet_properties.pageSetUpPr.autoPageBreaks = True
    for break_row in device_page_breaks:
        ws.row_breaks.append(Break(id=break_row))
    ws.freeze_panes = "A13"
    ws.oddFooter.center.text = "搬入依頼書　第 &P 頁"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.properties.title = f"搬入依頼書 {form.get('job_no', '')} {form.get('batch', '')}"
    wb.properties.creator = "上海康肯销售订单管理系统"

    return wb


def right_align():
    return Alignment(horizontal="right", vertical="center")


def validate_form(form):
    errors = []
    if not str(form.get("job_no") or "").strip():
        errors.append("缺少 JOB No")
    if not str(form.get("batch") or "").strip():
        errors.append("缺少发货批次")
    if not form.get("devices"):
        errors.append("批次至少需要一台设备")
    if not str(form.get("customer") or "").strip():
        errors.append("搬入先名不能为空")
    if not str(form.get("address") or "").strip():
        errors.append("搬入先地址不能为空")
    if not str(form.get("move_in_date") or "").strip():
        errors.append("搬入日期不能为空")
    return errors


def export_xlsx(form, path):
    errors = validate_form(form)
    if errors:
        raise ValueError("；".join(errors))
    path = os.path.abspath(os.fspath(path))
    os.makedirs(os.path.dirname(path), exist_ok=True)
    workbook = build_workbook(form)
    fd, temporary = tempfile.mkstemp(prefix=".move-in-", suffix=".xlsx",
                                      dir=os.path.dirname(path))
    os.close(fd)
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.remove(temporary)
    return path
