#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate the standard quotation workbook used by the desktop application."""

from __future__ import annotations

import os
import tempfile
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins


THIN = Side(style="thin", color="000000")
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LABEL_FILL = PatternFill("solid", fgColor="EAF1F7")
BASE_FONT = Font(name="Arial", size=10, color="1C1C1C")
TITLE_FONT = Font(name="Arial", size=20, bold=True, color="1C1C1C")
COMPANY_FONT = Font(name="Arial", size=12, bold=True, color="1C1C1C")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="1C1C1C")


def _safe_text(value):
    text = "" if value is None else str(value)
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def _set(ws, cell, value="", *, font=None, alignment=None, border=None, fill=None,
         number_format=None, formula=False):
    target = ws[cell]
    target.value = value if formula else (_safe_text(value) if isinstance(value, str) else value)
    target.font = font or BASE_FONT
    target.alignment = alignment or Alignment(vertical="center")
    if border:
        target.border = border
    if fill:
        target.fill = fill
    if number_format:
        target.number_format = number_format
    return target


def _merge_set(ws, start_row, start_col, end_row, end_col, value="", **style):
    ws.merge_cells(
        start_row=start_row, start_column=start_col,
        end_row=end_row, end_column=end_col,
    )
    for row in ws.iter_rows(
        min_row=start_row, max_row=end_row, min_col=start_col, max_col=end_col
    ):
        for cell in row:
            cell.border = style.get("border") or Border()
            cell.fill = style.get("fill") or PatternFill(fill_type=None)
    return _set(ws, f"{get_column_letter(start_col)}{start_row}", value, **style)


def _style_row(ws, row, start_col=1, end_col=9, *, border=GRID, fill=None,
               font=None, height=22):
    ws.row_dimensions[row].height = height
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row, col)
        cell.border = border
        cell.font = font or BASE_FONT
        cell.alignment = Alignment(vertical="center")
        if fill:
            cell.fill = fill


def _date_value(raw):
    try:
        return date.fromisoformat(str(raw))
    except (TypeError, ValueError):
        return _safe_text(raw)


def build_quotation_workbook(quote):
    """Build a one-sheet quotation following the two supplied quotation examples."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Quotation"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = None
    widths = {"A": 8, "B": 22, "C": 14, "D": 16, "E": 10,
              "F": 10, "G": 18, "H": 18, "I": 34}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    centered = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")

    _merge_set(ws, 1, 1, 1, 9, quote.get("issuer_name") or "KANKEN TECHNO CO.,LTD",
               font=COMPANY_FONT, alignment=centered)
    _merge_set(ws, 2, 1, 2, 9, quote.get("issuer_address", ""), alignment=centered)
    issuer_contact = quote.get("issuer_contact", "")
    if not issuer_contact and quote.get("salesperson"):
        issuer_contact = f"Sales: {quote['salesperson']}"
    _merge_set(ws, 3, 1, 3, 9, issuer_contact, alignment=centered)
    for row in (4, 5, 6):
        _style_row(ws, row, 1, 3, border=GRID, height=21)
    _set(ws, "A4", "QT No.", font=HEADER_FONT, border=GRID)
    _merge_set(ws, 4, 2, 4, 3, quote.get("quote_no", ""),
               border=GRID, alignment=centered)
    _set(ws, "A5", "Rev.", font=HEADER_FONT, border=GRID)
    _merge_set(ws, 5, 2, 5, 3, quote.get("revision_label", "1.0"),
               border=GRID, alignment=centered)
    _set(ws, "A6", "Date", font=HEADER_FONT, border=GRID)
    date_cell = _merge_set(ws, 6, 2, 6, 3, _date_value(quote.get("quote_date", "")),
                           border=GRID, alignment=centered)
    date_cell.number_format = "yyyy-mm-dd"

    ws.row_dimensions[7].height = 32
    _merge_set(ws, 7, 1, 7, 9, "Q u o t a t i o n",
               font=TITLE_FONT, alignment=centered)
    ws.row_dimensions[8].height = 10
    _merge_set(ws, 9, 1, 9, 9, quote.get("customer", ""),
               font=HEADER_FONT, alignment=left)
    _merge_set(ws, 10, 1, 10, 9, quote.get("customer_address", ""), alignment=left)
    contact = quote.get("contact", "")
    _merge_set(ws, 11, 1, 11, 9, f"Attn: {contact}" if contact else "", alignment=left)
    _merge_set(ws, 12, 1, 12, 9, "Dear Sir or Madam,", alignment=left)
    _merge_set(ws, 14, 1, 14, 9, "We are pleased to offer you as follows :", alignment=left)
    _merge_set(ws, 15, 1, 15, 9, quote.get("subject", ""),
               font=HEADER_FONT, alignment=left)

    header_row = 17
    _style_row(ws, header_row, fill=LABEL_FILL, font=HEADER_FONT, height=24)
    _merge_set(ws, header_row, 2, header_row, 4, "Item", font=HEADER_FONT,
               alignment=centered, border=GRID, fill=LABEL_FILL)
    headers = {1: "No.", 5: "Qty.", 6: "Unit", 7: "Price", 8: "Amount", 9: "Remark"}
    for col, label in headers.items():
        _set(ws, f"{get_column_letter(col)}{header_row}", label,
             font=HEADER_FONT, alignment=centered, border=GRID, fill=LABEL_FILL)

    item_start = header_row + 1
    row = item_start
    amount_rows = []
    for number, item in enumerate(quote.get("items") or [], 1):
        _style_row(ws, row, height=24)
        _merge_set(ws, row, 2, row, 4, item.get("model", ""),
                   alignment=left, border=GRID)
        _set(ws, f"A{row}", number, alignment=centered, border=GRID)
        _set(ws, f"E{row}", item.get("quantity", 0), alignment=centered,
             border=GRID, number_format="0.###")
        _set(ws, f"F{row}", item.get("unit", "台"), alignment=centered, border=GRID)
        _set(ws, f"G{row}", item.get("unit_price", 0), alignment=right,
             border=GRID, number_format="#,##0.00")
        _set(ws, f"H{row}", f"=E{row}*G{row}", alignment=right,
             border=GRID, number_format="#,##0.00", formula=True)
        _set(ws, f"I{row}", item.get("remark", ""), alignment=wrap, border=GRID)
        amount_rows.append(row)
        row += 1
        if item.get("description"):
            _style_row(ws, row, height=23)
            _merge_set(ws, row, 2, row, 4, item.get("description", ""),
                       alignment=wrap, border=GRID)
            for col in (1, 5, 6, 7, 8, 9):
                _set(ws, f"{get_column_letter(col)}{row}", "", border=GRID)
            row += 1

    total_row = row
    _style_row(ws, total_row, height=24, font=HEADER_FONT)
    _merge_set(ws, total_row, 2, total_row, 4, "TOTAL", font=HEADER_FONT,
               alignment=left, border=GRID)
    sum_formula = "+".join(f"H{value}" for value in amount_rows) or "0"
    _set(ws, f"H{total_row}", f"={sum_formula}", font=HEADER_FONT,
         alignment=right, border=GRID, number_format="#,##0.00", formula=True)
    _set(ws, f"I{total_row}",
         f"VAT {quote.get('tax_rate', 0):g}% excluded" if quote.get("tax_rate") else "",
         alignment=wrap, border=GRID)

    terms_row = total_row + 1
    _style_row(ws, terms_row, height=22)
    _merge_set(ws, terms_row, 2, terms_row, 4, quote.get("incoterm", ""),
               alignment=left, border=GRID)
    currency_row = terms_row + 1
    _style_row(ws, currency_row, height=22)
    _merge_set(ws, currency_row, 5, currency_row, 7, quote.get("currency", "RMB"),
               alignment=centered, border=GRID)
    _set(ws, f"H{currency_row}", f"=H{total_row}", alignment=right,
         border=GRID, number_format="#,##0.00", formula=True)

    footer_row = currency_row + 2
    footer_specs = []
    if quote.get("warranty"):
        footer_specs.append(("Warranty", quote.get("warranty"), "", ""))
    footer_specs.extend([
        ("Ship to", quote.get("ship_to", ""), "Date of delivery",
         quote.get("delivery_terms", "")),
        ("Description", quote.get("description", ""), "Terms of payment",
         quote.get("payment_terms", "")),
        ("Validity of this Quotation",
         f"By {quote.get('valid_until')}" if quote.get("valid_until") else "",
         "Remark", quote.get("notes", "")),
    ])
    for left_label, left_value, right_label, right_value in footer_specs:
        _style_row(ws, footer_row, height=34 if "\n" in str(right_value) else 24)
        _merge_set(ws, footer_row, 1, footer_row, 2, left_label,
                   font=HEADER_FONT, alignment=left, border=GRID, fill=LABEL_FILL)
        _merge_set(ws, footer_row, 3, footer_row, 5, left_value,
                   alignment=wrap, border=GRID)
        _merge_set(ws, footer_row, 6, footer_row, 7, right_label,
                   font=HEADER_FONT, alignment=left, border=GRID, fill=LABEL_FILL)
        _merge_set(ws, footer_row, 8, footer_row, 9, right_value,
                   alignment=wrap, border=GRID)
        footer_row += 1

    signature_row = footer_row + 2
    _merge_set(ws, signature_row, 6, signature_row, 9,
               quote.get("issuer_name") or "KANKEN TECHNO CO.,LTD",
               font=HEADER_FONT, alignment=centered,
               border=Border(top=THIN))

    last_row = signature_row + 1
    ws.print_area = f"A1:I{last_row}"
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(
        left=0.35, right=0.35, top=0.45, bottom=0.45, header=0.15, footer=0.15
    )
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.properties.title = f"Quotation {quote.get('quote_no', '')}"
    wb.properties.creator = "上海康肯销售订单管理系统"
    return wb


def export_quotation_xlsx(quote, path):
    path = os.path.abspath(os.fspath(path))
    os.makedirs(os.path.dirname(path), exist_ok=True)
    workbook = build_quotation_workbook(quote)
    fd, temporary = tempfile.mkstemp(
        prefix=".quotation-", suffix=".xlsx", dir=os.path.dirname(path)
    )
    os.close(fd)
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.remove(temporary)
    return path
