import unittest

from openpyxl import Workbook

import build_ledger_main as ledger


class BuildEngineReliabilityTests(unittest.TestCase):
    def test_coverage_parser_accepts_chinese_semicolon_and_line_breaks(self):
        self.assertEqual(ledger.coverage_set("A；B\nC"), {"A", "B", "C"})

    def test_user_text_that_looks_like_formula_is_escaped(self):
        ws = Workbook().active

        ledger.set_cell(ws, 1, 1, '=HYPERLINK("https://bad.example")')
        ledger.set_cell(ws, 2, 1, "=SUM(A1:A1)", trusted_formula=True)

        self.assertEqual(ws.cell(1, 1).value, "'=HYPERLINK(\"https://bad.example\")")
        self.assertEqual(ws.cell(2, 1).value, "=SUM(A1:A1)")

    def test_no_device_job_aggregates_multiple_invoices_and_payments(self):
        wb = Workbook()
        ws = wb.active
        contract = {
            "担当者": "担当", "客户": "客户", "JOB No": "26BS001",
            "订单内容": "设备", "付款条件": "预付", "发货方式": "",
            "发货地点": "", "送货地点": "",
        }
        terms = [{"款类": "预付款"}]
        invoices = [
            {"记录ID": "i1", "款类": "预付款", "开票日": "2026-08-01",
             "含税金额": 100, "应收回款日": "2026-08-10"},
            {"记录ID": "i2", "款类": "预付款", "开票日": "2026-08-03",
             "含税金额": 200, "应收回款日": "2026-08-12"},
        ]
        payments = [
            {"记录ID": "p1", "款类": "预付款", "回款日": "2026-08-05", "含税金额": 30},
            {"记录ID": "p2", "款类": "预付款", "回款日": "2026-08-06", "含税金额": 50},
        ]

        ledger.write_job_sheet(
            ws, "26BS001", contract, [], terms, [], {}, invoices, payments
        )

        headers = {ws.cell(5, col).value: col for col in range(1, ws.max_column + 1)}
        self.assertEqual(ws.cell(6, headers["预付款·开票金额"]).value, 300)
        self.assertEqual(ws.cell(6, headers["预付款·回款金额"]).value, 80)
        self.assertEqual(ws.cell(6, headers["预付款·开票日"]).value, "2026/08/03")
        self.assertEqual(ws.cell(6, headers["预付款·回款日"]).value, "2026/08/06")

    def test_no_device_job_without_payment_ratio_is_visible_for_confirmation(self):
        contracts = [{"JOB No": "26BS001", "客户": "客户"}]
        invoices = [{
            "记录ID": "i1", "JOB No": "26BS001", "款类": "预付款",
            "含税金额": 100, "开票日": "2026-08-01", "应收回款日": "2026-08-10",
        }]
        payments = [{
            "记录ID": "p1", "JOB No": "26BS001", "款类": "预付款",
            "含税金额": 20, "回款日": "2026-08-05",
        }]

        rows = ledger.compute_unpaid_rows(contracts, [], [], invoices, payments, [])

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["预警等级"], "④待确认")
        self.assertEqual(rows[0]["未回收金额"], 80)
        self.assertIn("付款比例未配置", rows[0]["未回收原因"])

    def test_no_device_job_uses_cumulative_payments_when_estimating_total(self):
        contracts = [{"JOB No": "26BS001", "客户": "客户"}]
        terms = [
            {"JOB No": "26BS001", "款类": "预付款", "比例%": 50},
            {"JOB No": "26BS001", "款类": "验收款", "比例%": 50},
        ]
        payments = [
            {"JOB No": "26BS001", "款类": "预付款", "含税金额": 30},
            {"JOB No": "26BS001", "款类": "预付款", "含税金额": 50},
        ]

        rows = ledger.compute_unpaid_rows(contracts, terms, [], [], payments, [])

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["款类"], "验收款")
        self.assertEqual(rows[0]["未回收金额"], 80)

    def test_multiple_payments_for_same_device_are_accumulated(self):
        contracts = [{"JOB No": "26BS001", "客户": "客户"}]
        terms = [{"JOB No": "26BS001", "款类": "发货款", "比例%": 100}]
        shipments = [{"JOB No": "26BS001", "发货批次": "1"}]
        devices = [{
            "JOB No": "26BS001", "製造番号": "26BS001-001", "发货批次": "1",
            "未税单价": 100, "是否无偿": "否", "验收状态": "已验收",
        }]
        invoices = [{
            "记录ID": "i1", "JOB No": "26BS001", "款类": "发货款",
            "开票日": "2026-08-01", "状态": "已开", "含税金额": 113,
            "覆盖批次": "1", "覆盖製造番号": "26BS001-001",
            "应收回款日": "2026-08-10",
        }]
        payments = [
            {"\u8bb0\u5f55ID": "p1", "JOB No": "26BS001", "\u6b3e\u7c7b": "\u53d1\u8d27\u6b3e",
             "\u56de\u6b3e\u65e5": "2026-08-05", "\u542b\u7a0e\u91d1\u989d": 30,
             "\u8986\u76d6\u6279\u6b21": "1", "\u8986\u76d6\u88fd\u9020\u756a\u53f7": "26BS001-001"},
            {"\u8bb0\u5f55ID": "p2", "JOB No": "26BS001", "\u6b3e\u7c7b": "\u53d1\u8d27\u6b3e",
             "\u56de\u6b3e\u65e5": "2026-08-06", "\u542b\u7a0e\u91d1\u989d": 40,
             "\u8986\u76d6\u6279\u6b21": "1", "\u8986\u76d6\u88fd\u9020\u756a\u53f7": "26BS001-001"},
        ]

        rows = ledger.compute_unpaid_rows(
            contracts, terms, shipments, invoices, payments, devices
        )

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["未回收金额"], 43)

    def test_multiple_invoices_for_same_device_are_accumulated(self):
        contracts = [{"JOB No": "26BS001", "客户": "客户"}]
        terms = [{"JOB No": "26BS001", "款类": "发货款", "比例%": 100}]
        shipments = [{"JOB No": "26BS001", "发货批次": "1"}]
        devices = [{
            "JOB No": "26BS001", "製造番号": "26BS001-001", "发货批次": "1",
            "未税单价": 100, "是否无偿": "否", "验收状态": "已验收",
        }]
        invoices = [
            {"\u8bb0\u5f55ID": "i1", "JOB No": "26BS001", "\u6b3e\u7c7b": "\u53d1\u8d27\u6b3e",
             "\u5f00\u7968\u65e5": "2026-08-01", "\u72b6\u6001": "\u5df2\u5f00", "\u542b\u7a0e\u91d1\u989d": 50,
             "\u8986\u76d6\u6279\u6b21": "1", "\u8986\u76d6\u88fd\u9020\u756a\u53f7": "26BS001-001",
             "\u5e94\u6536\u56de\u6b3e\u65e5": "2026-08-10"},
            {"\u8bb0\u5f55ID": "i2", "JOB No": "26BS001", "\u6b3e\u7c7b": "\u53d1\u8d27\u6b3e",
             "\u5f00\u7968\u65e5": "2026-08-02", "\u72b6\u6001": "\u5df2\u5f00", "\u542b\u7a0e\u91d1\u989d": 63,
             "\u8986\u76d6\u6279\u6b21": "1", "\u8986\u76d6\u88fd\u9020\u756a\u53f7": "26BS001-001",
             "\u5e94\u6536\u56de\u6b3e\u65e5": "2026-08-10"},
        ]

        rows = ledger.compute_unpaid_rows(
            contracts, terms, shipments, invoices, [], devices
        )

        self.assertEqual(len(rows), 1)
        self.assertNotIn("未开票", rows[0]["未回收原因"])

    def _warranty_rows(self, warranty_end):
        contracts = [{"JOB No": "26BS001", "客户": "客户"}]
        terms = [{"JOB No": "26BS001", "款类": "质保款", "比例%": 10,
                  "账期天数": 30, "触发条件": "质保期满后"}]
        shipments = [{"JOB No": "26BS001", "发货批次": "1"}]
        devices = [{
            "JOB No": "26BS001", "製造番号": "26BS001-001", "发货批次": "1",
            "未税单价": 100, "是否无偿": "否", "验收状态": "已验收",
            "质保结束日": warranty_end,
        }]
        invoices = [{
            "记录ID": "full-1", "JOB No": "26BS001", "款类": "全额",
            "开票日": "2020-01-01", "状态": "已开", "含税金额": 56.5,
            "覆盖批次": "1", "覆盖製造番号": "26BS001-001",
            "应收回款日": "2020-01-01",
        }]
        return ledger.compute_unpaid_rows(
            contracts, terms, shipments, invoices, [], devices
        )

    def test_warranty_unexpired_is_not_marked_overdue(self):
        rows = self._warranty_rows("2099-12-31")

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["预定回收日期"], "2100/01/30")
        self.assertEqual(rows[0]["预警等级"], "③未到期")
        self.assertIn("质保未到期", rows[0]["未回收原因"])
        self.assertNotIn("逾期", rows[0]["未回收原因"])

    def test_warranty_expired_uses_warranty_end_plus_term_days(self):
        rows = self._warranty_rows("2020-01-01")

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["预定回收日期"], "2020/01/31")
        self.assertEqual(rows[0]["预警等级"], "①已逾期")
        self.assertIn("逾期", rows[0]["未回收原因"])
        self.assertNotIn("质保未到期", rows[0]["未回收原因"])

    def test_missing_warranty_end_requires_confirmation(self):
        rows = self._warranty_rows("")

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["预定回收日期"], "")
        self.assertEqual(rows[0]["预警等级"], "④待确认")
        self.assertIn("质保期未设置", rows[0]["未回收原因"])
        self.assertNotIn("逾期", rows[0]["未回收原因"])

    def test_job_sheet_warranty_due_does_not_fall_back_to_full_invoice_date(self):
        ws = Workbook().active
        contract = {"担当者": "担当", "客户": "客户", "JOB No": "26BS001",
                    "订单内容": "设备", "付款条件": "", "发货方式": "",
                    "发货地点": "", "送货地点": ""}
        devices = [{
            "JOB No": "26BS001", "製造番号": "26BS001-001", "发货批次": "1",
            "设备型号": "KT1000", "未税单价": 100, "是否无偿": "否", "機番": "A1",
            "PO No": "", "验收状态": "已验收", "质保开始日": "2098-01-01",
            "质保结束日": "2099-12-31", "质保期": "2年", "送货单回收": "", "备注": "",
        }]
        terms = [{"JOB No": "26BS001", "款类": "质保款", "比例%": 10,
                  "账期天数": 30, "触发条件": "质保期满后"}]
        invoices = [{
            "记录ID": "full-1", "款类": "全额", "开票日": "2020-01-01",
            "含税金额": 56.5, "覆盖批次": "1", "覆盖製造番号": "26BS001-001",
            "应收回款日": "2020-01-01",
        }]

        ledger.write_job_sheet(ws, "26BS001", contract, devices, terms, [],
                               {("26BS001", "1"): {"出荷日": "2020-01-01", "台数": 1}},
                               invoices, [])

        headers = {ws.cell(5, col).value: col for col in range(1, ws.max_column + 1)}
        due = ws.cell(6, headers["行应收回款日"]).value
        self.assertEqual(due.strftime("%Y/%m/%d"), "2100/01/30")
        self.assertIn("行质保待确认", headers)
        self.assertIn("行质保规则", headers)

    def test_explicit_serial_matches_override_batch_fallback(self):
        fallback = {"记录ID": "batch", "覆盖批次": "1", "覆盖製造番号": ""}
        explicit = {"记录ID": "serial", "覆盖批次": "1",
                    "覆盖製造番号": "26BS001-001"}

        matches = ledger.matching_records(
            [fallback, explicit], "26BS001-001", "1"
        )

        self.assertEqual(matches, [explicit])

    def test_job_sheet_shows_all_matching_invoice_and_payment_amounts(self):
        ws = Workbook().active
        contract = {
            "担当者": "担当", "客户": "客户", "JOB No": "26BS001",
            "订单内容": "设备", "付款条件": "", "发货方式": "",
            "发货地点": "", "送货地点": "",
        }
        devices = [{
            "JOB No": "26BS001", "製造番号": "26BS001-001", "发货批次": "1",
            "设备型号": "KT1000", "未税单价": 100, "是否无偿": "否",
            "機番": "A1", "PO No": "", "验收状态": "", "质保开始日": "",
            "质保结束日": "", "质保期": "", "送货单回收": "", "备注": "",
        }]
        invoices = [
            {"\u8bb0\u5f55ID": "i1", "\u6b3e\u7c7b": "\u53d1\u8d27\u6b3e", "\u5f00\u7968\u65e5": "2026-08-01",
             "\u542b\u7a0e\u91d1\u989d": 50, "\u8986\u76d6\u6279\u6b21": "1", "\u8986\u76d6\u88fd\u9020\u756a\u53f7": "26BS001-001"},
            {"\u8bb0\u5f55ID": "i2", "\u6b3e\u7c7b": "\u53d1\u8d27\u6b3e", "\u5f00\u7968\u65e5": "2026-08-02",
             "\u542b\u7a0e\u91d1\u989d": 63, "\u8986\u76d6\u6279\u6b21": "1", "\u8986\u76d6\u88fd\u9020\u756a\u53f7": "26BS001-001"},
        ]
        payments = [
            {"\u8bb0\u5f55ID": "p1", "\u6b3e\u7c7b": "\u53d1\u8d27\u6b3e", "\u56de\u6b3e\u65e5": "2026-08-05",
             "\u542b\u7a0e\u91d1\u989d": 30, "\u8986\u76d6\u6279\u6b21": "1", "\u8986\u76d6\u88fd\u9020\u756a\u53f7": "26BS001-001"},
            {"\u8bb0\u5f55ID": "p2", "\u6b3e\u7c7b": "\u53d1\u8d27\u6b3e", "\u56de\u6b3e\u65e5": "2026-08-06",
             "\u542b\u7a0e\u91d1\u989d": 40, "\u8986\u76d6\u6279\u6b21": "1", "\u8986\u76d6\u88fd\u9020\u756a\u53f7": "26BS001-001"},
        ]

        ledger.write_job_sheet(
            ws, "26BS001", contract, devices, [{"款类": "发货款"}], [], {},
            invoices, payments,
        )

        headers = {ws.cell(5, col).value: col for col in range(1, ws.max_column + 1)}
        self.assertEqual(ws.cell(6, headers["发货款·开票金额"]).value, 113)
        self.assertEqual(ws.cell(6, headers["发货款·回款金额"]).value, 70)


if __name__ == "__main__":
    unittest.main()
