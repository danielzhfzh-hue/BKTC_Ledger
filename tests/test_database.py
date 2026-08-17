import json
import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path

from openpyxl import load_workbook

import core
import database


def sample_data():
    data = core.empty_data()
    data["合同订单"] = [{
        "记录ID": "contract-1",
        "JOB No": "26BS001",
        "客户": "原客户",
        "币种": "RMB",
        "关联合同": ["remote-link"],
    }]
    data["付款条件"] = [{
        "记录ID": "term-1",
        "JOB No": "26BS001",
        "款类": "预付款",
        "比例%": 100,
        "账期天数": 0,
        "说明": "全额预付",
    }]
    data["发货批次"] = [{
        "记录ID": "shipment-1",
        "JOB No": "26BS001",
        "发货批次": "1",
        "出荷日": "2026-08-01",
    }]
    data["设备台账"] = [{
        "记录ID": "device-1",
        "JOB No": "26BS001",
        "设备型号": "KT1000",
        "製造番号": "26BS001-001",
        "機番": "A1",
        "未税单价": 100,
        "是否无偿": "否",
        "发货批次": "1",
    }]
    return data


class DatabaseTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.db = self.root / "ledger.db"
        self.rules = {
            "临近天数": 30,
            "金额容差": 0.01,
            "未验收待确认": True,
            "发票异常待确认": True,
        }

    def tearDown(self):
        self.tmp.cleanup()

    def test_json_migration_is_relational_and_preserves_unknown_fields(self):
        source = self.root / "ledger.records.json"
        payload = {"tables": core.TABLES, "version": 1, "预警规则": self.rules}
        payload.update(sample_data())
        source.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

        result = database.migrate_json_to_database(source, self.db)

        self.assertEqual(result["revision"], 1)
        loaded = database.load_database(self.db)
        self.assertEqual(loaded["合同订单"][0]["关联合同"], ["remote-link"])
        self.assertEqual(database.get_rules(self.db), self.rules)
        with closing(sqlite3.connect(self.db)) as conn:
            conn.row_factory = sqlite3.Row
            tables = {row[0] for row in conn.execute(
                "SELECT name FROM sqlite_master WHERE type='table'"
            )}
            self.assertTrue(set(core.TABLES).issubset(tables))
            self.assertEqual(conn.execute("PRAGMA foreign_key_check").fetchall(), [])

    def test_invalid_orphan_is_rejected_without_changing_database(self):
        database.save_database(self.db, sample_data(), self.rules, backup=False)
        original_revision = database.get_revision(self.db)
        broken = sample_data()
        broken["设备台账"].append({
            "记录ID": "device-orphan",
            "JOB No": "26BS999",
            "设备型号": "X",
            "製造番号": "X-1",
            "发货批次": "1",
        })

        with self.assertRaises(database.DatabaseValidationError):
            database.save_database(self.db, broken, self.rules, backup=False)

        self.assertEqual(database.get_revision(self.db), original_revision)
        self.assertEqual(len(database.load_database(self.db)["设备台账"]), 1)

    def test_editable_workbook_previews_and_applies_add_update_delete(self):
        database.save_database(self.db, sample_data(), self.rules, backup=False)
        book = self.root / "editable.xlsx"
        database.export_editable_workbook(self.db, book)
        wb = load_workbook(book)

        contracts = wb["合同订单"]
        headers = {cell.value: cell.column for cell in contracts[1]}
        contracts.cell(2, headers["客户"], "新客户")

        devices = wb["设备台账"]
        device_headers = {cell.value: cell.column for cell in devices[1]}
        devices.delete_rows(2, 1)
        new_row = devices.max_row + 1
        devices.cell(new_row, device_headers["JOB No"], "26BS001")
        devices.cell(new_row, device_headers["设备型号"], "KT2000")
        devices.cell(new_row, device_headers["製造番号"], "26BS001-002")
        devices.cell(new_row, device_headers["機番"], "A2")
        devices.cell(new_row, device_headers["未税单价"], 200)
        devices.cell(new_row, device_headers["是否无偿"], "否")
        devices.cell(new_row, device_headers["发货批次"], "1")
        wb.save(book)

        prepared = database.prepare_editable_import(self.db, book)

        self.assertFalse(prepared.stale)
        self.assertEqual(prepared.action_counts["新增"], 1)
        self.assertEqual(prepared.action_counts["删除"], 1)
        self.assertGreaterEqual(prepared.action_counts["修改"], 1)
        result = database.apply_editable_import(self.db, prepared)
        self.assertEqual(result["revision"], 2)
        self.assertTrue(Path(result["backup"]).is_file())
        loaded = database.load_database(self.db)
        self.assertEqual(loaded["合同订单"][0]["客户"], "新客户")
        self.assertEqual(loaded["合同订单"][0]["关联合同"], ["remote-link"])
        self.assertEqual(loaded["设备台账"][0]["製造番号"], "26BS001-002")

    def test_stale_workbook_cannot_be_applied(self):
        database.save_database(self.db, sample_data(), self.rules, backup=False)
        book = self.root / "editable.xlsx"
        database.export_editable_workbook(self.db, book)
        newer = sample_data()
        newer["合同订单"][0]["客户"] = "数据库中的新值"
        database.save_database(self.db, newer, self.rules, backup=False)

        prepared = database.prepare_editable_import(self.db, book)

        self.assertTrue(prepared.stale)
        with self.assertRaises(database.StaleImportError):
            database.apply_editable_import(self.db, prepared)

    def test_unrevisioned_database_change_after_preview_is_rejected(self):
        database.save_database(self.db, sample_data(), self.rules, backup=False)
        book = self.root / "editable.xlsx"
        database.export_editable_workbook(self.db, book)
        prepared = database.prepare_editable_import(self.db, book)
        with closing(sqlite3.connect(self.db)) as conn:
            conn.execute('UPDATE "合同订单" SET "客户" = ? WHERE "JOB No" = ?',
                         ("外部直接修改", "26BS001"))
            conn.commit()

        with self.assertRaises(database.StaleImportError):
            database.apply_editable_import(self.db, prepared)

    def test_formula_cells_are_rejected(self):
        database.save_database(self.db, sample_data(), self.rules, backup=False)
        book = self.root / "editable.xlsx"
        database.export_editable_workbook(self.db, book)
        wb = load_workbook(book)
        ws = wb["合同订单"]
        headers = {cell.value: cell.column for cell in ws[1]}
        ws.cell(2, headers["客户"], "=1+1")
        wb.save(book)

        with self.assertRaises(database.SpreadsheetFormatError):
            database.prepare_editable_import(self.db, book)

    def test_existing_database_is_migrated_when_schema_adds_customer_columns(self):
        database.save_database(self.db, sample_data(), self.rules, backup=False)
        with closing(sqlite3.connect(self.db)) as conn:
            conn.execute('ALTER TABLE "付款条件" DROP COLUMN "客户"')
            conn.commit()

        loaded = database.load_database(self.db)

        self.assertIn("客户", loaded["付款条件"][0])

    def test_editable_import_propagates_relational_key_renames(self):
        data = sample_data()
        data["开票记录"] = [{
            "记录ID": "invoice-1", "JOB No": "26BS001", "款类": "预付款",
            "开票日": "2026-08-01", "状态": "已开", "含税金额": 113,
            "覆盖批次": "1", "覆盖製造番号": "26BS001-001",
            "账期天数": 10,
        }]
        data["回款记录"] = [{
            "记录ID": "payment-1", "JOB No": "26BS001", "款类": "预付款",
            "回款日": "2026-08-05", "含税金额": 50,
            "覆盖批次": "1", "覆盖製造番号": "26BS001-001",
        }]
        database.save_database(self.db, data, self.rules, backup=False)
        book = self.root / "editable.xlsx"
        database.export_editable_workbook(self.db, book)
        wb = load_workbook(book)

        edits = {
            "合同订单": {"JOB No": "26BS002"},
            "付款条件": {"款类": "发货款"},
            "发货批次": {"发货批次": "2"},
            "设备台账": {"製造番号": "26BS002-001"},
        }
        for table, values in edits.items():
            ws = wb[table]
            headers = {cell.value: cell.column for cell in ws[1]}
            for field, value in values.items():
                ws.cell(2, headers[field], value)
        wb.save(book)

        prepared = database.prepare_editable_import(self.db, book)

        for table in core.TABLES:
            self.assertTrue(all(r["JOB No"] == "26BS002" for r in prepared.data[table]))
        self.assertEqual(prepared.data["设备台账"][0]["发货批次"], "2")
        for table in ("开票记录", "回款记录"):
            self.assertEqual(prepared.data[table][0]["款类"], "发货款")
            self.assertEqual(prepared.data[table][0]["覆盖批次"], "2")
            self.assertEqual(prepared.data[table][0]["覆盖製造番号"], "26BS002-001")


class CoreReliabilityTests(unittest.TestCase):
    def test_validation_catches_duplicate_contract_blank_child_and_invalid_date(self):
        data = sample_data()
        data["合同订单"].append({"记录ID": "contract-2", "JOB No": "26BS001"})
        data["回款记录"].append({"记录ID": "payment-1", "JOB No": ""})
        data["设备台账"][0]["质保开始日"] = "2026-99-99"

        messages = [issue["msg"] for issue in core.validate(data)]

        self.assertTrue(any("合同 JOB No 重复" in msg for msg in messages))
        self.assertTrue(any("缺少 JOB No" in msg for msg in messages))
        self.assertTrue(any("不是有效日期" in msg for msg in messages))

    def test_validation_rejects_invalid_select_and_coverage_references(self):
        data = sample_data()
        data["开票记录"] = [{
            "记录ID": "invoice-1", "JOB No": "26BS001", "款类": "未知款类",
            "覆盖批次": "2", "覆盖製造番号": "26BS001-999",
        }]

        messages = [issue["msg"] for issue in core.validate(data)]

        self.assertTrue(any("值不在允许列表" in msg for msg in messages))
        self.assertTrue(any("不存在的批次" in msg for msg in messages))
        self.assertTrue(any("不存在的製造番号" in msg for msg in messages))

    def test_blank_payment_placeholder_does_not_fail_ratio_validation(self):
        data = sample_data()
        data["付款条件"] = [{"记录ID": "placeholder", "JOB No": "26BS001"}]

        issues = core.validate(data)

        self.assertFalse(any("付款条件比例合计" in issue["msg"] for issue in issues))

    def test_meaningful_payment_with_devices_requires_coverage(self):
        data = sample_data()
        data["回款记录"] = [{
            "记录ID": "payment-1", "JOB No": "26BS001", "款类": "预付款",
            "回款日": "2026-08-05", "含税金额": 100,
            "覆盖批次": "", "覆盖製造番号": "",
        }]

        issues = core.validate(core.derive(data))

        self.assertTrue(any(
            issue["severity"] == "error" and "回款必须选择覆盖批次或设备" in issue["msg"]
            for issue in issues
        ))

    def test_blank_payment_placeholder_and_no_device_payment_allow_blank_coverage(self):
        placeholder = sample_data()
        placeholder["回款记录"] = [{"\u8bb0\u5f55ID": "payment-1", "JOB No": "26BS001"}]
        self.assertFalse(any(
            "回款必须选择覆盖批次或设备" in issue["msg"]
            for issue in core.validate(core.derive(placeholder))
        ))

        no_device = sample_data()
        no_device["设备台账"] = []
        no_device["发货批次"] = []
        no_device["回款记录"] = [{
            "记录ID": "payment-1", "JOB No": "26BS001", "款类": "预付款",
            "回款日": "2026-08-05", "含税金额": 100,
        }]
        self.assertFalse(any(
            "回款必须选择覆盖批次或设备" in issue["msg"]
            for issue in core.validate(core.derive(no_device))
        ))

    def test_customer_and_payment_status_fields_are_derived(self):
        data = sample_data()
        data["开票记录"] = [{
            "记录ID": "invoice-1", "JOB No": "26BS001", "款类": "预付款",
            "开票日": "2026-08-01", "状态": "已开", "含税金额": 113,
            "覆盖批次": "1", "覆盖製造番号": "26BS001-001",
            "账期天数": 10,
        }]
        data["回款记录"] = [{
            "记录ID": "payment-1", "JOB No": "26BS001", "款类": "预付款",
            "回款日": "2026-08-15", "含税金额": 50,
            "覆盖批次": "1", "覆盖製造番号": "26BS001-001",
        }]

        derived = core.derive(data)

        for table in ("付款条件", "设备台账", "发货批次", "开票记录", "回款记录"):
            self.assertEqual(derived[table][0]["客户"], "原客户")
        self.assertEqual(derived["开票记录"][0]["回款状态"], "部分回款")
        payment = derived["回款记录"][0]
        self.assertEqual(payment["对应应收回款日"], "2026-08-11")
        self.assertEqual(payment["是否超期"], "是")
        self.assertEqual(payment["超期天数"], 4)

    def test_overlapping_invoices_use_cumulative_amount_for_payment_status(self):
        data = sample_data()
        data["开票记录"] = [
            {
                "记录ID": "invoice-1", "JOB No": "26BS001", "款类": "预付款",
                "开票日": "2026-08-01", "状态": "已开", "含税金额": 50,
                "覆盖批次": "1", "覆盖製造番号": "26BS001-001",
            },
            {
                "记录ID": "invoice-2", "JOB No": "26BS001", "款类": "预付款",
                "开票日": "2026-08-02", "状态": "已开", "含税金额": 63,
                "覆盖批次": "1", "覆盖製造番号": "26BS001-001",
            },
        ]
        data["回款记录"] = [{
            "记录ID": "payment-1", "JOB No": "26BS001", "款类": "预付款",
            "回款日": "2026-08-05", "含税金额": 70,
            "覆盖批次": "1", "覆盖製造番号": "26BS001-001",
        }]

        derived = core.derive(data)

        self.assertEqual(
            [row["回款状态"] for row in derived["开票记录"]],
            ["部分回款", "部分回款"],
        )

    def test_device_batch_change_recomputes_transaction_coverage_batch(self):
        data = sample_data()
        data["发货批次"].append({
            "记录ID": "shipment-2", "JOB No": "26BS001", "发货批次": "2",
        })
        data["设备台账"][0]["发货批次"] = "2"
        data["回款记录"] = [{
            "记录ID": "payment-1", "JOB No": "26BS001", "款类": "预付款",
            "回款日": "2026-08-05", "含税金额": 100,
            "覆盖批次": "1", "覆盖製造番号": "26BS001-001",
        }]

        derived = core.derive(data)

        self.assertEqual(derived["回款记录"][0]["覆盖批次"], "2")

    def test_transfer_serial_keeps_ambiguous_coverage_batches(self):
        data = sample_data()
        data["发货批次"].append({
            "记录ID": "shipment-2", "JOB No": "26BS001", "发货批次": "2",
        })
        data["设备台账"][0].update({
            "製造番号": "26BS001-001→26BS002-001", "发货批次": "1",
        })
        data["设备台账"].append({
            "记录ID": "device-2", "JOB No": "26BS001", "製造番号": "26BS001-001",
            "设备型号": "KT1000", "未税单价": 100, "是否无偿": "否",
            "发货批次": "2",
        })
        data["回款记录"] = [{
            "记录ID": "payment-1", "JOB No": "26BS001", "款类": "预付款",
            "回款日": "2026-08-05", "含税金额": 200,
            "覆盖批次": "1;2", "覆盖製造番号": "26BS001-001→26BS002-001",
        }]

        derived = core.derive(data)

        self.assertEqual(derived["回款记录"][0]["覆盖批次"], "1;2")

    def test_excel_text_formula_is_escaped(self):
        self.assertEqual(core._exp_val("=HYPERLINK(\"bad\")", core.TEXT),
                         "'=HYPERLINK(\"bad\")")


if __name__ == "__main__":
    unittest.main()
