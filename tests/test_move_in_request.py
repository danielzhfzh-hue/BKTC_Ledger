import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

import app
import core
import move_in_request


def sample_data(device_count=2):
    data = core.empty_data()
    data["合同订单"] = [{
        "记录ID": "contract-1", "JOB No": "26BS006", "客户": "搬入客户",
        "送货地点": "上海市嘉定区测试路 1 号", "发货地点": "平湖康肯",
    }]
    data["发货批次"] = [
        {"记录ID": "ship-1", "JOB No": "26BS006", "发货批次": "1", "出荷日": "2026-08-28"},
        {"记录ID": "ship-2", "JOB No": "26BS006", "发货批次": "2", "出荷日": "2026-09-01"},
    ]
    data["设备台账"] = [{
        "记录ID": f"dev-{i}", "JOB No": "26BS006", "发货批次": "1",
        "设备型号": "KT1000FI" if i == 0 else "KT1000EPS-C300",
        "機番": f"2580{i}", "製造番号": f"26BS006-{i + 1:02d}",
        "PO No": "PO-1" if i == 0 else "PO-2",
    } for i in range(device_count)]
    return data


class MoveInRequestTests(unittest.TestCase):
    def test_defaults_use_exact_job_and_batch_and_do_not_confuse_ship_date(self):
        data = sample_data()
        defaults = move_in_request.shipment_defaults(data, shipment_id="ship-1")

        self.assertEqual(defaults["batch"], "1")
        self.assertEqual(defaults["address"], "上海市嘉定区测试路 1 号")
        self.assertEqual(defaults["ship_date"], "2026-08-28")
        self.assertEqual(defaults["move_in_date"], "")
        self.assertEqual(len(defaults["devices"]), 2)
        self.assertIn("出荷元", defaults["system_notes"])
        self.assertEqual(defaults["notes"], "")
        self.assertEqual(defaults["devices"][0]["po_no"], "PO-1")

    def test_missing_devices_is_rejected(self):
        data = sample_data()
        data["设备台账"] = []
        with self.assertRaisesRegex(ValueError, "没有关联设备"):
            move_in_request.shipment_defaults(data, shipment_id="ship-1")

    def test_workbook_uses_template_paired_equipment_layout(self):
        data = sample_data(device_count=6)
        form = move_in_request.shipment_defaults(data, shipment_id="ship-1")
        form.update({"move_in_date": "2026-08-31", "contact": "测试联系人 13800000000"})

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "搬入依頼書.xlsx"
            move_in_request.export_xlsx(form, path)
            workbook = load_workbook(path, data_only=False)
            sheet = workbook["搬入依頼書"]
            self.assertEqual(sheet.print_area, "'搬入依頼書'!$A$1:$P$32")
            self.assertEqual(sheet.page_setup.orientation, "portrait")
            self.assertEqual(sheet.page_setup.fitToWidth, 1)
            self.assertEqual(sheet.page_setup.fitToHeight, 0)
            self.assertEqual(workbook.sheetnames, ["搬入依頼書"])
            self.assertIsNone(sheet.print_title_rows)
            self.assertEqual(sheet["C11"].value, 6)
            self.assertEqual(sheet["D11"].value, "台")
            self.assertEqual(sheet["A12"].value, "No.")
            self.assertEqual(sheet["B12"].value, "PO No.")
            self.assertEqual(sheet["I12"].value, "No.")
            self.assertEqual(sheet["J12"].value, "PO No.")
            self.assertEqual(sheet["C13"].value, "KT1000FI")
            self.assertEqual(sheet["E13"].value, "25800")
            self.assertEqual(sheet["G13"].value, "26BS006-01")
            self.assertEqual(sheet["B13"].value, "PO-1")
            self.assertEqual(sheet["K13"].value, "KT1000EPS-C300")
            self.assertEqual(sheet["M13"].value, "25801")
            self.assertEqual(sheet["O13"].value, "26BS006-02")
            self.assertEqual(sheet["J13"].value, "PO-2")
            self.assertEqual(sheet["E15"].value, "25804")
            self.assertEqual(sheet["M15"].value, "25805")
            self.assertEqual(sheet["O15"].value, "26BS006-06")
            self.assertIn("C14:D15", {str(r) for r in sheet.merged_cells.ranges})
            self.assertIn("J13:J15", {str(r) for r in sheet.merged_cells.ranges})
            self.assertEqual(sheet["C16"].value, "☑ 有り")
            self.assertEqual(sheet["E16"].value, "ID No.")
            self.assertIn(sheet["G16"].value, (None, ""))
            self.assertEqual(sheet["P16"].value, "☐ 無し")
            self.assertEqual(sheet["C20"].value, "搬入日")
            self.assertEqual(sheet["E20"].value.isoformat(), "2026-08-31T00:00:00")
            self.assertEqual(sheet["I20"].value, "時刻")
            self.assertIn(sheet["J20"].value, (None, ""))
            self.assertIn("C21:G21", {str(r) for r in sheet.merged_cells.ranges})
            self.assertIn("I21:P21", {str(r) for r in sheet.merged_cells.ranges})
            for row in range(21, 25):
                self.assertEqual(sheet[f"C{row}"].border.top.style, "thin")
                self.assertEqual(sheet[f"C{row}"].border.bottom.style, "thin")
                self.assertEqual(sheet[f"H{row}"].border.bottom.style, "thin")
            self.assertEqual(sheet["H21"].value, "t")
            self.assertIn(sheet["H24"].value, (None, ""))
            self.assertNotIn("※１", "\n".join(
                str(cell.value or "")
                for row in sheet.iter_rows()
                for cell in row
            ))
            self.assertEqual(len(sheet.row_breaks.brk), 0)

    def test_workbook_fits_twenty_three_devices_and_notes_on_one_page(self):
        data = sample_data(device_count=23)
        form = move_in_request.shipment_defaults(data, shipment_id="ship-1")
        form.update({"move_in_date": "2026-08-31"})

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "搬入依頼書-23台.xlsx"
            move_in_request.export_xlsx(form, path)
            sheet = load_workbook(path, data_only=False)["搬入依頼書"]
            self.assertEqual(sheet.print_area, "'搬入依頼書'!$A$1:$P$41")
            self.assertEqual(sheet["A24"].value, 23)
            self.assertEqual(sheet["C25"].value, "☑ 有り")
            self.assertEqual(sheet["A34"].value, 8)
            self.assertEqual(len(sheet.row_breaks.brk), 0)

    def test_workbook_starts_a_new_equipment_block_after_thirty_devices(self):
        data = sample_data(device_count=31)
        form = move_in_request.shipment_defaults(data, shipment_id="ship-1")
        form.update({"move_in_date": "2026-08-31"})

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "搬入依頼書-31台.xlsx"
            move_in_request.export_xlsx(form, path)
            sheet = load_workbook(path, data_only=False)["搬入依頼書"]
            self.assertEqual(sheet.print_area, "'搬入依頼書'!$A$1:$P$46")
            self.assertEqual([br.id for br in sheet.row_breaks.brk], [27])
            self.assertEqual(sheet["A28"].value, "No.")
            self.assertEqual(sheet["I28"].value, "No.")
            self.assertEqual(sheet["A29"].value, 31)
            self.assertIn(sheet["I29"].value, (None, ""))
            self.assertEqual(sheet["E29"].value, "258030")
            self.assertEqual(sheet["C30"].value, "☑ 有り")

    def test_workbook_supports_fifty_devices_without_truncation(self):
        data = sample_data(device_count=50)
        form = move_in_request.shipment_defaults(data, shipment_id="ship-1")
        form.update({"move_in_date": "2026-08-31", "notes": "现场需要分两次进场"})

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "搬入依頼書-50台.xlsx"
            move_in_request.export_xlsx(form, path)
            workbook = load_workbook(path, data_only=False)
            sheet = workbook["搬入依頼書"]
            machine_numbers = [
                sheet.cell(row=row, column=col).value
                for row in range(13, 39)
                for col in (5, 13)
                if str(sheet.cell(row=row, column=col).value or "").startswith("2580")
            ]
            self.assertEqual(machine_numbers, [f"2580{i}" for i in range(50)])
            self.assertEqual(sheet["E13"].value, "25800")
            self.assertEqual(sheet["M13"].value, "25801")
            self.assertEqual(sheet["E29"].value, "258030")
            self.assertEqual(sheet["M38"].value, "258049")
            self.assertIsNone(sheet.print_title_rows)
            self.assertEqual([br.id for br in sheet.row_breaks.brk], [27])
            self.assertEqual(sheet["A28"].value, "No.")
            self.assertEqual(sheet["I28"].value, "No.")
            self.assertIn("现场需要分两次进场", "\n".join(
                str(cell.value or "")
                for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row)
                for cell in row
            ))
            self.assertEqual(len(sheet.row_breaks.brk), 1)

    def test_long_special_notes_continue_without_truncation(self):
        data = sample_data(device_count=6)
        form = move_in_request.shipment_defaults(data, shipment_id="ship-1")
        marker = "最后一行特記事項"
        form.update({"move_in_date": "2026-08-31", "notes": "现场要求：" + ("安全确认；" * 45) + marker})

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "搬入依頼書-长特記事項.xlsx"
            move_in_request.export_xlsx(form, path)
            workbook = load_workbook(path, data_only=False)
            sheet = workbook["搬入依頼書"]
            values = "\n".join(
                str(cell.value or "")
                for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row)
                for cell in row
            )
            self.assertIn(marker, values)
            self.assertFalse(any("特記事項（続き）" in str(cell.value or "")
                                 for row in sheet.iter_rows()
                                 for cell in row))

    def test_api_export_reads_current_database_and_writes_installation_folder(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            api = app.Api(str(root / "ledger.db"), operator_name="测试员")
            state = api.load_state()
            data = sample_data()
            api.save_data(data, state["rules"])
            defaults = api.get_move_in_request_defaults("ship-1")
            defaults.update({"move_in_date": "2026-08-31", "contact": "测试联系人"})

            output_dir = root / "搬入依頼書"
            with mock.patch.object(
                app, "_move_in_request_output_dir", return_value=(str(output_dir), "")
            ):
                result = api.export_move_in_request(defaults)

            output = Path(result["path"])
            self.assertTrue(output.is_file())
            self.assertEqual(output.parent, output_dir)
            self.assertEqual(result["device_count"], 2)
            self.assertEqual(load_workbook(output).active["C13"].value, "KT1000FI")

    def test_api_export_rejects_stale_database_revision(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            api = app.Api(str(root / "ledger.db"))
            state = api.load_state()
            api.save_data(sample_data(), state["rules"])
            form = api.get_move_in_request_defaults("ship-1")
            form.update({"move_in_date": "2026-08-31"})
            api.save_data(sample_data(), state["rules"])
            with self.assertRaises(app.database.StaleImportError):
                api.export_move_in_request(form)


if __name__ == "__main__":
    unittest.main()
