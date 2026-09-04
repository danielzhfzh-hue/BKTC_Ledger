import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

import quotation


class QuotationWorkbookTests(unittest.TestCase):
    def test_standard_quotation_workbook_has_formulas_and_dynamic_terms(self):
        quote = {
            "quote_no": "26SHMT0903", "revision_label": "1.0",
            "quote_date": "2026-09-03", "issuer_name": "上海康肯环境技术有限公司",
            "issuer_address": "上海市测试路 1 号", "issuer_contact": "",
            "salesperson": "张三",
            "customer": "测试客户", "customer_address": "客户地址", "contact": "王先生",
            "subject": "KT1000EPS-C300 报价", "currency": "RMB", "tax_rate": 13,
            "incoterm": "DAP", "warranty": "验收后 12 个月",
            "ship_to": "合肥", "delivery_terms": "收到订单后 12 周",
            "description": "真空设备", "payment_terms": "30% 预付，70% 验收后支付",
            "valid_until": "2026-10-03", "notes": "不含安装费",
            "items": [
                {"model": "KT1000EPS-C300", "description": "主机",
                 "quantity": 2, "unit": "台", "unit_price": 1000000,
                 "remark": "标准配置"},
                {"model": "Option-A", "quantity": 1, "unit": "套",
                 "unit_price": 50000, "remark": "选配"},
            ],
        }
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "quotation.xlsx"
            quotation.export_quotation_xlsx(quote, output)
            workbook = load_workbook(output, data_only=False)
            sheet = workbook["Quotation"]

            self.assertEqual(sheet["B4"].value, "26SHMT0903")
            self.assertEqual(sheet["A9"].value, "测试客户")
            self.assertEqual(sheet["A3"].value, "Sales: 张三")
            self.assertEqual(sheet["H18"].value, "=E18*G18")
            self.assertEqual(sheet["H20"].value, "=E20*G20")
            self.assertEqual(sheet["H21"].value, "=H18+H20")
            self.assertIn("30% 预付", sheet["H27"].value)
            self.assertEqual(sheet.print_area, "'Quotation'!$A$1:$I$32")
            self.assertEqual(sheet.page_setup.orientation, "landscape")
            self.assertNotIn("NEXCHIP", " ".join(
                str(cell.value or "") for row in sheet.iter_rows() for cell in row
            ))

    def test_user_text_that_looks_like_a_formula_is_not_executable(self):
        quote = {
            "quote_no": "SAFE-001", "quote_date": "2026-09-04",
            "customer": "=HYPERLINK(\"https://example.invalid\")",
            "items": [{
                "model": "+CMD", "quantity": 1, "unit": "台",
                "unit_price": 1, "remark": "@external",
            }],
        }
        workbook = quotation.build_quotation_workbook(quote)
        sheet = workbook["Quotation"]

        self.assertEqual(sheet["A9"].value, "'=HYPERLINK(\"https://example.invalid\")")
        self.assertEqual(sheet["B18"].value, "'+CMD")
        self.assertEqual(sheet["I18"].value, "'@external")
        self.assertEqual(sheet["H18"].value, "=E18*G18")
        self.assertEqual(sheet["H18"].data_type, "f")


if __name__ == "__main__":
    unittest.main()
